from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


OUTPUT = Path("/Users/zouxiaoli/Documents/ChatGPT/ai询价/造价通AI询价一期需求方案_20260915.docx")
FONT_CN = "PingFang SC"
FONT_EN = "Arial"
HEADER_FILL = "1F4E79"
HEADER_TEXT = "FFFFFF"
ALT_FILL = "EEF4F8"
BORDER = "D9D9D9"
BLACK = RGBColor(0, 0, 0)
WHITE = RGBColor(255, 255, 255)


def set_run_font(run, size=11, bold=False, color=BLACK, name_cn=FONT_CN, name_en=FONT_EN):
    run.font.name = name_en
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name_en)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name_en)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name_cn)


def set_paragraph_font(paragraph, size=11, bold=False, color=BLACK):
    for run in paragraph.runs:
        set_run_font(run, size=size, bold=bold, color=color)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, margin_value in (
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ):
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(margin_value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), BORDER)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_row_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def format_cell(cell, text, header=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.12
    run = paragraph.add_run(str(text))
    set_run_font(run, size=9.5, bold=header, color=WHITE if header else BLACK)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)
    if header:
        shade_cell(cell, HEADER_FILL)


def add_table(doc, headers, rows, widths=None, center_columns=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)

    for idx, header in enumerate(headers):
        format_cell(table.rows[0].cells[idx], header, header=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_repeat_table_header(table.rows[0])
    set_row_cant_split(table.rows[0])

    center_columns = set(center_columns or [])
    for row_idx, row in enumerate(rows):
        table_row = table.add_row()
        set_row_cant_split(table_row)
        cells = table_row.cells
        for col_idx, value in enumerate(row):
            align = WD_ALIGN_PARAGRAPH.CENTER if col_idx in center_columns else WD_ALIGN_PARAGRAPH.LEFT
            format_cell(cells[col_idx], value, align=align)
            if row_idx % 2 == 1:
                shade_cell(cells[col_idx], ALT_FILL)

    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Cm(width)

    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_body(doc, text, bold_lead=None):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.28
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        set_run_font(lead, size=11, bold=True)
        rest = paragraph.add_run(text[len(bold_lead):])
        set_run_font(rest, size=11)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, size=11)
    return paragraph


def add_bullet(doc, text, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    paragraph = doc.add_paragraph(style=style)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.2
    run = paragraph.add_run(text)
    set_run_font(run, size=10.5)
    return paragraph


def add_numbered(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.55)
    paragraph.paragraph_format.first_line_indent = Cm(-0.55)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.2
    run = paragraph.add_run(text)
    set_run_font(run, size=10.5)
    return paragraph


def add_code(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.6)
    paragraph.paragraph_format.right_indent = Cm(0.3)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.12
    run = paragraph.add_run(text)
    set_run_font(run, size=9.5, name_cn="Songti SC", name_en="Courier New")
    return paragraph


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    set_run_font(run, size=9, color=RGBColor(90, 90, 90))


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Cm(1.9)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)

    normal = doc.styles["Normal"]
    normal.font.name = FONT_EN
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.28

    for style_name, size in (("Title", 22), ("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)):
        style = doc.styles[style_name]
        style.font.name = FONT_EN
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
        style.paragraph_format.space_before = Pt(10 if style_name != "Title" else 0)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True

    header = section.header
    header_paragraph = header.paragraphs[0]
    header_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header_paragraph.add_run("造价通 AI 询价一期需求方案")
    set_run_font(header_run, size=8.5, color=RGBColor(90, 90, 90))

    footer = section.footer
    add_page_number(footer.paragraphs[0])


def build_document():
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("造价通 AI 询价一期需求方案")
    set_run_font(title_run, size=22, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run("基于 AI 查价扩展的标准材料历史询价与建议价推荐")
    set_run_font(subtitle_run, size=12, color=RGBColor(70, 70, 70))

    add_body(
        doc,
        "方案结论：一期建议复用现有 AI 查价入口与对话交互，新增 AI 询价能力。系统优先检索近 5 年历史询价数据，无匹配时调用造价通市场价，按系数及价格修正规则形成建议价，推荐前 3 条询价结果，并支持按询价模板导出含盖章要求文件。",
    )
    add_body(
        doc,
        "本文档用于需求评审和跨团队确认。会议内容中已经明确的事项按确认口径记录；样例文件未覆盖、规则冲突或缺少字段的部分统一列入待确认事项，不直接作为开发结论。",
    )

    add_table(
        doc,
        ["项目", "内容"],
        [
            ["文档版本", "V0.1"],
            ["编写日期", "2026-09-15"],
            ["方案阶段", "需求梳理，待业务、数据、算法、研发评审"],
            ["适用范围", "造价通 AI 询价一期，标准数据询价"],
            ["核心资料", "会议记录、AI 查价界面截图、历史询价映射表和系数表、历史询价查询 SQL、价格推荐规则"],
            ["主要依赖", "数据中心、询价部、算法模型团队、研发、产品、法务或合规"],
        ],
        widths=[3.2, 14.2],
    )

    doc.add_page_break()
    doc.add_heading("目录", level=1)
    contents = [
        "1. 项目背景与产品定位",
        "2. 一期范围与目标",
        "3. 核心业务流程",
        "4. 功能需求",
        "5. 价格匹配与推荐规则",
        "6. 数据来源、抽取与治理",
        "7. 页面与交互说明",
        "8. 异常、权限与合规",
        "9. 验收标准",
        "10. 待确认事项",
        "11. 资料来源与术语",
    ]
    for item in contents:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(3)
        run = paragraph.add_run(item)
        set_run_font(run, size=10.5)

    doc.add_heading("1. 项目背景与产品定位", level=1)
    doc.add_heading("1.1 业务现状", level=2)
    add_body(
        doc,
        "造价通首页当前已有 AI 查价入口。用户可在独立对话页用自然语言输入材料名称、规格、关键词和地区，系统返回结构化价格结果，并可在结果不足时展示询价圈数据。",
    )
    add_body(
        doc,
        "现有 AI 查价主要解决价格信息检索和展示。AI 询价需要在此基础上形成从需求识别、历史询价匹配、建议价计算、结果推荐到文件导出的完整闭环，同时保留询价过程历史。",
    )

    doc.add_heading("1.2 产品定位", level=2)
    add_body(
        doc,
        "AI 询价面向需要快速获得标准材料价格参考、询价建议和可导出文件的造价及采购人员。系统通过历史询价优先、造价通市场价兜底的方式形成建议价，并按业务规则推荐 3 条结果。",
    )

    doc.add_heading("1.3 建设目标", level=2)
    for text in [
        "业务目标：将历史询价、市场价和系数规则整合为可追溯的线上询价结果。",
        "用户目标：减少人工查找和整理，通过一次输入获得建议价、询价依据和可导出文件。",
        "系统目标：建立历史询价 ES 数据、系数映射、价格推荐、模板回复和导出留档能力。",
    ]:
        add_bullet(doc, text)

    doc.add_heading("2. 一期范围与目标", level=1)
    add_table(
        doc,
        ["类型", "内容"],
        [
            ["一期已实现", "首页 AI 查价入口；自然语言对话；结构化价格结果；询价圈补充展示。"],
            ["本期新增", "AI 询价入口；历史询价数据检索；系数匹配及建议价；价格推荐前 3 条；询价结果模板；盖章文件导出；上下文和历史记忆；使用次数限制。"],
            ["影响范围", "首页入口、对话与结果页面、历史记录、用户次数、导出模块、历史询价数据、系数配置、价格修正模型及权限日志。"],
            ["不变范围", "现有 AI 查价主流程；标准市场价和信息价既有展示；不改变现有询价圈业务。"],
            ["本期不做", "非标图纸类询价；图纸识别；批量询价；按会议口径默认不展示价格来源和供应商联系方式。"],
            ["主要依赖", "近 5 年历史询价数据及增量；市场价接口；系数数据和计算公式；询价模板；盖章能力；价格修正模型。"],
        ],
        widths=[3.0, 14.4],
    )
    add_body(
        doc,
        "范围冲突提示：会议同时提出支持 Excel 文件解析和不支持批量询价。一期需明确 Excel 是仅导入单条材料参数，还是仅预留技术能力；未确认前不作为验收项。",
    )

    doc.add_heading("3. 核心业务流程", level=1)
    steps = [
        "1. 用户从首页新增的 AI 询价入口进入，输入材料名称、规格、地区、目标时间等信息。",
        "2. 系统识别并标准化询价条件，条件缺失时按既定交互补充，标准化失败时进入兜底链路。",
        "3. 系统优先查询近 5 年历史询价数据。查询和召回规则由数据中心另行提供。",
        "4. 对召回结果执行数据去重、材料名称和特征匹配，形成价格推荐候选集。",
        "5. 按核心时间、材料名称和询价唯一标识依次排序，推荐并锁定前 3 条历史询价结果。",
        "6. 历史询价无结果或按规则需要补充时，调用造价通市场价能力获取结果。",
        "7. 根据映射表确定系数名称，再按待确认的省份、城市和时间规则选择具体系数值。",
        "8. 原始价格通过系数计算和价格修正形成建议价，并保留计算依据供内部追溯。",
        "9. 系统按询价模板生成结果，展示允许披露的供应商信息，默认隐藏联系方式。",
        "10. 用户可导出带造价通水印和盖章要求的文件，系统记录导入、查询、推荐、修正和导出状态。",
    ]
    for step in steps:
        add_numbered(doc, step)

    add_body(
        doc,
        "无结果兜底：历史询价和造价通市场价均无结果时，返回测试阶段、类型不支持、时间过久无法查询等提示，或引导用户进入人工询价。具体文案和触发阈值待确认。",
    )

    doc.add_heading("4. 功能需求", level=1)
    add_table(
        doc,
        ["模块", "功能点", "需求描述", "优先级"],
        [
            ["入口", "AI 询价入口", "在造价通首页新增 AI 询价入口，并支持进入对话式询价页面。", "P0"],
            ["发起询价", "自然语言输入", "识别材料名称、规格、地区、目标时间、品牌等询价条件。", "P0"],
            ["文件解析", "Excel 解析", "支持解析 Excel；是否支持单条或批量需确认，一期不默认支持批量。", "待确认"],
            ["历史检索", "近 5 年数据查询", "优先查询历史询价 ES 数据，支持增量和时间参数化。", "P0"],
            ["价格兜底", "造价通市场价", "历史询价无结果时调用造价通市场价。", "P0"],
            ["系数匹配", "映射与系数选择", "按优先级匹配映射表，确定系数名称及最终系数值。", "P0"],
            ["建议价", "系数计算与修正", "原始价格经系数和价格修正模型转换为建议价。", "P0"],
            ["结果推荐", "推荐前 3 条", "按核心时间、材料名称、询价 ID 依次排序并截取 3 条。", "P0"],
            ["结果模板", "询价回复模板", "按询价部模板组织价格、区域、时间、供应商等字段。", "P0"],
            ["导出", "水印与盖章", "支持导出带造价通水印及盖章要求的文件。", "P0"],
            ["记忆", "上下文与历史", "保留当前会话上下文，并支持查看历史询价与导出记录。", "P1"],
            ["限制", "用户次数", "用户使用次数限制为 10 条，具体计数口径待确认。", "P0"],
            ["隐私", "联系方式和来源", "默认不回复联系方式及价格来源，异常联系方式需提示。", "P0"],
            ["兜底", "人工询价", "无数据或无法满足时提示原因并引导人工询价。", "P1"],
        ],
        widths=[2.3, 3.0, 10.3, 1.8],
        center_columns=[3],
    )

    doc.add_heading("5. 价格匹配与推荐规则", level=1)
    doc.add_heading("5.1 搜索召回与价格推荐的关系", level=2)
    add_body(
        doc,
        "搜索召回解决“哪些记录进入候选集”，价格推荐解决“从候选中推荐哪 3 条”。本方案记录的是价格推荐规则，搜索召回规则仍需要数据中心补充。",
    )
    add_body(
        doc,
        "候选集应优先完成数据去重、材料名称模糊匹配和特征属性匹配，再进入价格推荐排序。去重和召回的先后顺序、模糊匹配阈值及规格和地区的匹配范围均待确认。",
    )

    doc.add_heading("5.2 历史询价前 3 条推荐规则", level=2)
    add_table(
        doc,
        ["优先级", "字段", "排序规则"],
        [
            ["1", "核心时间", "用户填写目标时间时，优先推荐与目标时间最接近的记录；用户未填写时，以当前系统时间为目标，按新内容优先。"],
            ["2", "材料名称", "核心时间相同或均为空时，中文开头优先，其次英文，再其次数字；中文按汉语拼音升序，英文和数字按 ASCII 升序。"],
            ["3", "询价 ID", "前两个排序条件完全相同时，按全局唯一询价 ID 升序。"],
        ],
        widths=[1.8, 2.6, 13.0],
        center_columns=[0],
    )
    add_code(
        doc,
        "候选集 -> 去重和匹配 -> 核心时间排序 -> 材料名称排序 -> 询价 ID 升序 -> 取前 3 条",
    )
    add_body(
        doc,
        "实现要求：三个排序条件必须作为一个整体排序后截取结果，不能按每个优先级分别取前 3 条。若目标时间两侧距离相同，需要业务确认优先过去或未来记录。",
    )

    doc.add_heading("5.3 系数映射规则", level=2)
    add_body(
        doc,
        "映射按优先级从高到低执行，任一规则匹配成功后立即终止。输入包括历史询价的 cid、subcid、stdname 和 features。",
    )
    add_table(
        doc,
        ["顺序", "历史询价条件", "映射表条件", "结果"],
        [
            ["1", "stdname 非空", "stdname 完全相同，且历史询价 features 包含映射表 features", "取对应系数名称"],
            ["2", "stdname 非空", "stdname 完全相同，且映射表 features 为空", "取对应系数名称"],
            ["3", "继续匹配", "subcid 完全相同，且映射表 stdname、features 为空", "取对应系数名称"],
            ["4", "继续匹配", "cid 完全相同，且映射表 subcid、stdname、features 为空", "取对应系数名称"],
            ["5", "前四步均未匹配，含标准化失败", "使用默认 cid=32 对应系数", "取默认系数"],
        ],
        widths=[1.5, 4.2, 8.5, 3.2],
        center_columns=[0],
    )
    add_body(
        doc,
        "feature 包含建议按结构化特征集合处理：忽略顺序、空格和大小写，历史询价需要包含映射表要求的全部特征。例如历史询价为“材质:镀锌;表面处理:热镀锌”时，可匹配映射条件“材质:镀锌”，反向不成立。",
    )

    doc.add_heading("5.4 建议价计算公式", level=2)
    add_code(
        doc,
        "历史询价路径：建议价 = 历史询价 × 对应系数\n市场价路径：建议价 = 造价通市场价 × 对应系数\n价格修正：由价格修正模型进一步处理，模型输入和覆盖规则待定",
    )
    add_body(
        doc,
        "会议中提到的市场价、询价、时间系数是单项使用还是连乘，以及是否叠加历史询价建议折扣中的项目阶段 K 系数，需要在评审前明确。价格取整、含税口径、单位换算和异常价格处理也需要补充。",
    )

    doc.add_heading("5.5 映射表文件核对发现", level=2)
    add_table(
        doc,
        ["问题", "核对结果", "影响"],
        [
            ["默认 cid=32 缺失", "当前工作簿中不存在 cid=32，也没有隐藏数据包含该值。", "默认系数无法按现有文件落地，需要数据中心补充或确认默认系数名称。"],
            ["stdname 重复", "镀锌角钢、镀锌钢管、镀锌槽钢等存在重复映射，且部分映射结果不一致。", "同一条历史询价可能命中不同系数，违反匹配唯一性。"],
            ["cid 重复", "cid=33 同时映射到市政分类和园林分类。", "缺少 subcid 时无法确定应选择哪条映射。"],
            ["系数名称无法确定最终值", "t_coefficient 同一 NAME 会按省份、城市、月份出现多条记录。", "需要补充省份、城市、时间及回退规则后才能计算建议价。"],
        ],
        widths=[3.0, 7.0, 7.4],
    )

    doc.add_heading("6. 数据来源、抽取与治理", level=1)
    doc.add_heading("6.1 历史询价数据来源", level=2)
    add_table(
        doc,
        ["来源表", "用途", "关键条件"],
        [
            ["xunjia.t_ask_price", "询价主体、材料名称、规格、数量、单位、分类、标准名称、特征、品牌、客户备注", "t1.flow_statue=2，已完成询价"],
            ["xunjia.t_ask_offer", "供应商报价、回复时间、供应商、联系方式、税费和运费口径", "t2.flow_statue=1，有效报价"],
            ["xunjia.t_ask_pici", "询价工程师、项目名称、报价区域、询价批次", "按 pici_id 关联"],
        ],
        widths=[4.2, 8.0, 5.2],
    )
    add_body(
        doc,
        "当前 SQL 按回复时间抽取 2025 年数据，可作为字段基础，但产品目标是近 5 年数据，因此时间条件需要参数化并支持存量初始化和增量更新。",
    )

    doc.add_heading("6.2 SQL 字段核对与调整要求", level=2)
    add_table(
        doc,
        ["要求", "说明"],
        [
            ["补充 subcid", "当前 SQL 未返回 t1.subcid，无法执行系数匹配的第 3 优先级。"],
            ["补充唯一标识", "至少增加 ask_id 和 offer_id，用于去重、推荐排序和审计。"],
            ["补充项目阶段", "建议价可能需要项目阶段 K 系数，当前查询未返回该字段。"],
            ["调整时间条件", "避免对 t2.create_on 使用 DATE_FORMAT，改用半开区间以利用索引。"],
            ["明确连接语义", "有报价记录的过滤条件下，LEFT JOIN t2 与 INNER JOIN 的当前结果接近，建议统一语义。"],
            ["修正字段和空值", "has_carriag 需核对拼写，has_carriage、has_tax、联系方式等字段需要处理 NULL。"],
            ["隐私隔离", "供应商联系方式不得进入默认可回复内容，应单独存储并限制访问。"],
        ],
        widths=[4.0, 13.4],
    )

    doc.add_heading("6.3 数据治理规则", level=2)
    for item in [
        "cid、subcid 统一转为字符串并保留前导零，避免 01 与 1 被识别为不同值。",
        "材料名称、规格、标准名称和特征属性进入 ES 前进行去空格、大小写和字段格式统一。",
        "增量数据需要记录抽取批次、更新时间、数据来源和失败重试状态。",
        "历史询价记录至少保留询价 ID、报价 ID、询价批次和来源表主键，便于追溯。",
        "重复数据判定条件仍是待确认项，建议至少覆盖材料、规格、地区、供应商、回复时间和报价。",
    ]:
        add_bullet(doc, item)

    doc.add_heading("7. 页面与交互说明", level=1)
    add_table(
        doc,
        ["页面或状态", "交互要求"],
        [
            ["首页入口", "新增 AI 询价入口，与 AI 查价并列。点击后进入询价对话页。"],
            ["询价输入", "支持自然语言输入材料、规格、地区和时间；Excel 能力待确认。"],
            ["条件确认", "展示系统识别出的核心条件，允许用户修改后重新发起。"],
            ["推荐结果", "展示最多 3 条建议价结果，区分历史询价和市场价兜底结果。价格来源默认不外显。"],
            ["供应商信息", "可展示供应商名称；联系方式默认隐藏，信息异常时提示。"],
            ["导出操作", "提供导出入口，按询价模板生成带水印和盖章要求的文件。"],
            ["历史记录", "展示会话、询价条件、推荐结果、导出状态及重新询价入口。"],
            ["无数据状态", "说明测试阶段、类型不支持、时间过久或转人工询价。"],
            ["次数限制", "展示剩余次数；达到 10 条限制后阻止继续发起，并给出提示。具体口径待确认。"],
        ],
        widths=[3.5, 13.9],
    )
    add_body(
        doc,
        "用户不满意时建议保留原询价记录，提供“修改条件重询”和“按原条件重询”，不直接覆盖历史结果。该交互建议待产品评审。",
    )

    doc.add_heading("7.2 询价盖章报价模板", level=2)
    add_body(
        doc,
        "询价部提供的样例文件为 A4 纵向单页模板。页面顶部左侧展示造价通标识和网址，主体为供应商报价表，页面中部叠加询价专用章，右下角展示二维码，页脚保留页码。",
    )
    add_table(
        doc,
        ["模板区域", "内容", "生成要求"],
        [
            ["品牌区", "造价通 Logo、官网地址", "保持模板版式，不替换业务字段。"],
            ["报价主体", "材料名称、规格型号、数量、单位、品牌、价格、供应商信息、回复备注", "相同材料的名称、规格、数量、单位、品牌跨供应商报价行合并。"],
            ["报价行", "同一材料支持多个供应商报价，每个供应商一条价格记录", "AI 询价推荐 3 条的适用粒度需要确认。"],
            ["供应商信息", "样例包含供应商名称、联系人、手机号", "与会议中“默认不回复联系方式”冲突，需确认导出文件是否脱敏。"],
            ["税费和运费", "回复备注展示是否含税及含运费", "根据报价字段生成，不能由系统猜测。"],
            ["印章区", "询价专用章覆盖在报价表上方", "需确认可使用的章、盖章服务、文件有效性和防伪机制。"],
            ["二维码", "页面右下角展示二维码", "需确认二维码指向、有效性和是否与询价记录一一对应。"],
            ["页码", "页面底部居中显示页码", "模板跨页时页码连续更新。"],
        ],
        widths=[3.0, 7.0, 7.4],
    )
    add_body(
        doc,
        "模板样本支持同一材料对应多个供应商报价，也支持一页展示多个材料。由于一期暂不支持批量询价，建议按“单个材料询价生成一组报价，包含最多 3 条供应商报价”理解，但需业务明确。",
    )

    doc.add_heading("8. 异常、权限与合规", level=1)
    doc.add_heading("8.1 异常与兜底", level=2)
    add_table(
        doc,
        ["场景", "处理要求"],
        [
            ["历史询价无结果", "按规则调用造价通市场价。"],
            ["历史询价和市场价均无结果", "提示测试阶段、类型不支持或时间过久，或引导人工询价。"],
            ["结果不足 3 条", "明确展示实际数量，是否用市场价补足仍需确认。"],
            ["标准化失败", "进入 cid=32 默认系数链路，但当前默认映射缺失。"],
            ["映射重复", "禁止依赖不确定顺序，需补充确定性选择规则。"],
            ["价格修正失败", "保留未修正结果并标记状态，或按待确认规则终止。"],
            ["联系方式异常", "不对外展示联系方式，并按内部数据校验规则提示。"],
            ["导出或盖章失败", "保留询价结果，允许重新导出，并记录失败原因。"],
            ["用户超出次数", "阻止继续发起并提示剩余次数为 0。"],
        ],
        widths=[4.0, 13.4],
    )

    doc.add_heading("8.2 权限、隐私与审计", level=2)
    for item in [
        "普通用户仅可查看自己的会话、询价结果和导出记录。",
        "供应商联系方式属于受限字段，默认不回复；内部查看权限单独控制。",
        "价格来源和模型计算过程默认不对外披露；正式上线前需要法务或合规确认。",
        "询价、系数选择、价格修正、模板生成和导出操作均需记录日志。",
        "盖章文件需要记录文件编号、生成时间、操作者和下载记录，防伪机制待补充。",
    ]:
        add_bullet(doc, item)

    doc.add_heading("9. 验收标准", level=1)
    add_table(
        doc,
        ["编号", "验收项", "验收标准"],
        [
            ["A1", "入口", "首页可进入 AI 询价，用户可发起一次标准数据询价。"],
            ["A2", "条件识别", "系统可识别材料、规格、地区和时间，并允许用户修改。"],
            ["A3", "历史优先", "有符合条件的历史询价时优先返回历史结果。"],
            ["A4", "市场价兜底", "无历史询价时按规则调用造价通市场价。"],
            ["A5", "映射规则", "按五级优先级执行，命中后终止；默认系数在缺失修复后可用。"],
            ["A6", "价格计算", "系数和价格修正计算正确，能够解释原始价与建议价。"],
            ["A7", "前 3 条推荐", "按核心时间、材料名称、询价 ID 整体排序后返回最多 3 条。"],
            ["A8", "去重与召回", "搜索召回规则明确后，重复记录可稳定去除。"],
            ["A9", "结果模板", "结果字段、水印和盖章要求符合询价部模板。"],
            ["A10", "导出", "导出文件内容与页面结果一致，文件可正常打开和下载。"],
            ["A11", "隐私", "供应商联系方式默认不展示，价格来源默认不回复。"],
            ["A12", "历史记忆", "用户可查看历史会话和导出记录，重新询价不会误覆盖原记录。"],
            ["A13", "次数限制", "达到 10 条限制后无法继续发起，提示正确。"],
            ["A14", "无数据异常", "无数据、失败、超时和导出异常均有明确提示。"],
        ],
        widths=[1.3, 3.2, 12.9],
        center_columns=[0],
    )

    doc.add_heading("10. 待确认事项", level=1)
    add_table(
        doc,
        ["编号", "待确认事项", "影响", "优先级"],
        [
            ["Q1", "默认 cid=32 对应哪个系数名称，当前映射文件缺失。", "默认兜底无法实现。", "P0"],
            ["Q2", "重复 stdname 和重复 cid 如何保证唯一匹配。", "同一材料可能得到不同系数。", "P0"],
            ["Q3", "确定系数名称后，如何选择省份、城市和月份的具体 COEF。", "无法计算最终建议价。", "P0"],
            ["Q4", "市场价、询价、时间系数是单项使用还是连乘。", "价格公式不确定。", "P0"],
            ["Q5", "项目阶段 K 系数是否叠加及乘法顺序。", "建议价可能遗漏业务调整。", "P0"],
            ["Q6", "价格修正模型的输入、输出、触发条件和失败兜底。", "修正链路无法开发和验收。", "P0"],
            ["Q7", "搜索召回、模糊匹配和去重规则。", "候选集不稳定。", "P0"],
            ["Q8", "历史数据是否包含同一询价下的多条供应商报价。", "前 3 条结果口径不确定。", "P0"],
            ["Q9", "结果不足 3 条时是否调用市场价补足。", "结果数量和展示逻辑不确定。", "P1"],
            ["Q10", "Excel 是单条导入还是批量，是否纳入一期。", "范围与验收不一致。", "P1"],
            ["Q11", "10 条限制按询价次数、结果条数还是导出次数计算。", "用户可用性和提示逻辑不确定。", "P1"],
            ["Q12", "盖章文件由谁生成、使用何种章、是否具有法律效力。", "导出和合规方案不确定。", "P0"],
            ["Q13", "默认不回复价格来源的边界及合规话术。", "可能影响用户知情和平台合规。", "P0"],
            ["Q14", "联系方式信息异常如何判断和提示。", "提示逻辑无法开发。", "P1"],
            ["Q15", "用户不满意时重新询价是否覆盖结果、是否消耗次数。", "用户流程和计数逻辑不确定。", "P1"],
            ["Q16", "模板样例展示联系人及手机号，与默认不回复联系方式冲突。", "导出文件可能违反既定隐私口径。", "P0"],
            ["Q17", "询价专用章的来源、盖章方式、二维码指向和文件有效期。", "导出文件无法完成模板和合规验收。", "P0"],
            ["Q18", "“锁定 3 条”是每个材料 3 个供应商报价，还是整个询价共 3 条结果。", "结果结构和模板生成逻辑不确定。", "P0"],
        ],
        widths=[1.3, 8.6, 5.7, 1.8],
        center_columns=[0, 3],
    )

    doc.add_heading("11. 资料来源与术语", level=1)
    add_table(
        doc,
        ["资料", "用途"],
        [
            ["AI 查价首页及对话界面截图", "确认现有入口、交互和结果展示基线。"],
            ["会议需求记录", "确认 AI 询价入口、流程、限制、导出和人工兜底要求。"],
            ["历史询价映射表和系数表0910.xlsx", "核对映射表结构、t_coefficient 字段、默认系数及重复数据。"],
            ["历史询价查询 SQL", "确认历史询价来源表、字段、状态和时间口径。"],
            ["价格推荐规则", "确认核心时间、材料名称、询价 ID 的排序和取前 3 条逻辑。"],
            ["询价盖章报价 PDF 模板", "确认输出字段、供应商报价行组织、询价专用章、二维码和页码。"],
        ],
        widths=[6.0, 11.4],
    )

    terms = [
        ("AI 查价", "现有自然语言价格查询能力，主要返回市场或平台价格信息。"),
        ("AI 询价", "基于 AI 查价扩展的历史询价检索、建议价计算、结果推荐和文件导出能力。"),
        ("搜索召回", "根据用户询价条件查询并筛选候选历史询价记录的规则。"),
        ("价格推荐", "对候选记录按核心时间、材料名称和询价 ID 排序并返回前 3 条的规则。"),
        ("建议价", "原始历史询价或市场价经系数及价格修正后形成的结果。"),
        ("t_coefficient", "按省份、城市、系数名称和发布日期存储系数值的表。"),
    ]
    add_table(doc, ["术语", "说明"], terms, widths=[3.4, 14.0])

    doc.add_heading("评审建议", level=1)
    add_body(
        doc,
        "建议先关闭扫描召回、默认系数、重复映射、最终系数选择、价格公式和盖章模板六项依赖，再进入详细原型和开发排期。数据侧可同步完成 ES 字段补充、索引设计和近 5 年存量抽取，产品侧同步确认页面状态和用户计数规则。",
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
