"""Finish naming, cover status, and privacy follow-ups after the main V1.3 patch."""

from __future__ import annotations

import copy
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph


ROOT = Path("/Users/zouxiaoli/Documents/ChatGPT/ai询价")
DOC_PATH = ROOT / "造价通AI询价需求规格说明书（V1.3）.docx"


def set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run._element.getparent().remove(run._element)
    else:
        paragraph.add_run(text)


def replace_exact(doc: Document, old: str, new: str) -> None:
    matches = [p for p in doc.paragraphs if p.text == old]
    if len(matches) != 1:
        raise ValueError(f"段落定位异常：{len(matches)}，{old}")
    set_paragraph_text(matches[0], new)


def replace_all_exact(doc: Document, old: str, new: str, expected: int) -> None:
    matches = [p for p in doc.paragraphs if p.text == old]
    if len(matches) != expected:
        raise ValueError(f"段落定位异常：{len(matches)}，{old}")
    for paragraph in matches:
        set_paragraph_text(paragraph, new)


def insert_bullet_after(anchor: Paragraph, text: str) -> None:
    new_element = copy.deepcopy(anchor._p)
    anchor._p.addnext(new_element)
    paragraph = Paragraph(new_element, anchor._parent)
    set_paragraph_text(paragraph, text)


def main() -> None:
    doc = Document(DOC_PATH)

    replace_exact(
        doc,
        "文档状态：[  ] 草案   [√] 正式发布   [  ] 正在修订",
        "文档状态：[√] 草案（评审中）   [  ] 正式发布   [  ] 正在修订",
    )
    replace_all_exact(doc, "3.4.12 历史询价记录与会话删除", "3.4.12 历史记录与会话删除", 2)
    replace_exact(doc, "图 3-16 历史询价记录列表截图", "图 3-16 历史记录列表截图")
    replace_exact(doc, "历史询价与会话删除：图 3-16、图 3-17。", "历史记录与会话删除：图 3-16、图 3-17。")
    replace_all_exact(doc, "3.5.2 AI 查价通宝交互文案", "3.5.2 AI 查价通宝交互文案（跨项目参考）", 2)

    privacy_anchor = next(
        p for p in doc.paragraphs if p.text.startswith("演示包中的价格、供应商、联系方式均为演示数据")
    )
    insert_bullet_after(
        privacy_anchor,
        "本期不使用用户输入训练模型，正式页面不得展示“输入将被收集用于模型训练”的提示条；"
        "如未来启用训练用途，需单独完成告知、同意、撤回、脱敏、留存与第三方共享评审。",
    )

    reference_table = doc.tables[23]
    for row_index, description in [
        (7, "3.5.2 AI 查价跨项目参考文案、模式切换、查询状态与固定业务回复的依据；不纳入本期验收。"),
        (8, "3.5.2 通宝动作、气泡文案、范围受控及异常场景文案的补充参考；不纳入本期验收。"),
    ]:
        cell = reference_table.rows[row_index].cells[1]
        set_paragraph_text(cell.paragraphs[0], description)

    text = "\n".join(p.text for p in doc.paragraphs)
    text += "\n" + "\n".join(
        cell.text for table in doc.tables for row in table.rows for cell in row.cells
    )
    for expected in [
        "文档状态：[√] 草案（评审中）",
        "3.4.12 历史记录与会话删除",
        "3.5.2 AI 查价通宝交互文案（跨项目参考）",
        "本期不使用用户输入训练模型",
    ]:
        if expected not in text:
            raise ValueError(f"缺少收尾修订：{expected}")
    if "3.4.12 历史询价记录与会话删除" in text:
        raise ValueError("仍存在旧历史记录章节名")

    doc.save(DOC_PATH)
    print(DOC_PATH)


if __name__ == "__main__":
    main()
