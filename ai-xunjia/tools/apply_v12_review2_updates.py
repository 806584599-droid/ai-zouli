"""Apply the second review round (2026-10-15) updates to the V1.2 requirement document."""

from __future__ import annotations

import shutil
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.table import Table, _Row
from docx.text.paragraph import Paragraph


ROOT = Path("/Users/zouxiaoli/Documents/ChatGPT/ai询价")
DOC_PATH = ROOT / "造价通AI询价需求规格说明书（V1.2）.docx"
BACKUP_PATH = ROOT / ".qa/造价通AI询价需求规格说明书（V1.2）_第二次评审更新前备份.docx"

CHANGE_LOG_TEXT = (
    "按 AI 询价第二次评审意见调整：项目阶段提示后置为“询价建议结果 → 项目阶段提示选择”，"
    "输入框占位文案与通宝提示语补充项目阶段并提供复制入口；"
    "输入条件项少于 5 项时在建议价区间旁提示“由于搜索条件不足，建议价区间或存在偏差”；"
    "历史询价列表删除状态列、操作列新增删除入口并明确记录保留一年；"
    "左侧额度区明确 AI 询价 10 条/天、AI 查价 50 条/天且本期不绑定套餐；"
    "AI 询价本期只保留快速思考模式，结果区移除“修改条件”与引用来源展示；"
    "盖章文件采用自动生成或人工审批方式列为待领导批示事项；"
    "导出文件增加造价通水印、固定位置盖章与文件末尾“该文件由 AI 生成，仅供报价参考。”提示语；"
    "重新查询需同步后端并通知第三方模型重新筛选候选供应商、更换候选报价；"
    "补充前期数据治理阶段使用 AI 对同一材料跨月份重复数据做去重的要求，避免与历史价倒推冲突；"
    "技术实现链路统一为：用户输入 → DeepSeek 意图拆解 → 提取材料名称与特征、时间、地区、项目阶段 → "
    "ES 模糊检索 → DeepSeek 候选筛选 → 项目阶段系数计算 → 前端展示。"
)


def set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def set_cell_text(cell, text: str) -> None:
    paragraph = cell.paragraphs[0]
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)
    for extra in list(cell.paragraphs[1:]):
        extra._element.getparent().remove(extra._element)


def replace_paragraph(doc: Document, prefix: str, new_text: str) -> None:
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            set_paragraph_text(paragraph, new_text)
            return
    raise RuntimeError(f"Paragraph not found: {prefix[:50]}")


def append_cell_paragraph(cell, text: str) -> None:
    template = cell.paragraphs[-1]
    new_p = deepcopy(template._p)
    cell._tc.append(new_p)
    paragraph = Paragraph(new_p, cell)
    set_paragraph_text(paragraph, text)


def insert_paragraph_after(paragraph: Paragraph, text: str) -> Paragraph:
    new_xml = deepcopy(paragraph._p)
    paragraph._p.addnext(new_xml)
    new_paragraph = Paragraph(new_xml, paragraph._parent)
    set_paragraph_text(new_paragraph, text)
    return new_paragraph


def find_paragraph(doc: Document, prefix: str) -> Paragraph:
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            return paragraph
    raise RuntimeError(f"Paragraph not found: {prefix[:50]}")


def find_table(doc: Document, header: str | tuple[str, ...]) -> Table:
    expected = (header,) if isinstance(header, str) else header
    for table in doc.tables:
        if not table.rows:
            continue
        cells = [cell.text.strip() for cell in table.rows[0].cells]
        if cells[: len(expected)] == list(expected):
            return table
    raise RuntimeError(f"Table not found: {expected}")


def find_row(table: Table, key: str, column: int = 0) -> _Row:
    for row in table.rows:
        if row.cells[column].text.strip() == key:
            return row
    raise RuntimeError(f"Row not found: {key}")


def insert_row_after(table: Table, anchor_key: str, values: list[str], column: int = 0) -> _Row:
    anchor = find_row(table, anchor_key, column)
    new_tr = deepcopy(anchor._tr)
    anchor._tr.addnext(new_tr)
    new_row = _Row(new_tr, table)
    for index, value in enumerate(values):
        if index < len(new_row.cells):
            set_cell_text(new_row.cells[index], value)
    return new_row


def append_row(table: Table, values: list[str]) -> _Row:
    new_tr = deepcopy(table.rows[-1]._tr)
    table.rows[-1]._tr.addnext(new_tr)
    new_row = _Row(new_tr, table)
    for index, value in enumerate(values):
        if index < len(new_row.cells):
            set_cell_text(new_row.cells[index], value)
    return new_row


def update_row(table: Table, key: str, values: dict[int, str], column: int = 0) -> None:
    row = find_row(table, key, column)
    for index, value in values.items():
        if index < len(row.cells):
            set_cell_text(row.cells[index], value)


def main() -> None:
    BACKUP_PATH.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(DOC_PATH, BACKUP_PATH)
    doc = Document(DOC_PATH)

    # --- 1.3 术语和缩略词 -------------------------------------------------
    terms = find_table(doc, "术语")
    insert_row_after(
        terms,
        "评测 harness",
        ["项目阶段", "询价所处的造价阶段，包含施工图预算、投资估算、设计概算与工程竣工结算；用于确定项目阶段系数，不在前端展示。"],
    )
    insert_row_after(
        terms,
        "项目阶段",
        ["每日额度", "AI 询价每天 10 条、AI 查价每天 50 条的使用次数上限；本期不绑定套餐。"],
    )
    insert_row_after(
        terms,
        "每日额度",
        ["数据治理", "在历史询价进入检索与倒推之前，对同一材料跨月份的重复或冲突记录进行去重与合并的处理过程。"],
    )

    # --- 2.5 技术实现链路 -------------------------------------------------
    replace_paragraph(
        doc,
        "AI 询价采用“大模型意图理解",
        "AI 询价采用“大模型意图理解 + 规则引擎 + ES 检索 + 大模型候选筛选”的组合链路。技术实现链路为："
        "用户输入 → DeepSeek 意图拆解 → 提取材料名称与特征、时间、地区、项目阶段 → ES 模糊检索 → "
        "DeepSeek 候选筛选 → 项目阶段系数计算 → 前端展示。DeepSeek 负责理解用户自然语言并形成结构化材料信息，"
        "规则引擎负责确定系数，ES 负责扩大候选召回，DeepSeek 负责判断候选材料的上下文相关性；"
        "最终推荐顺序仍按本文档确定的确定性规则执行，不由大模型直接决定。下表按该链路逐环节展开处理要求，"
        "其中材料标准化、系数映射与地区召回范围属于链路的内部处理步骤，不对用户展示。",
    )
    chain = find_table(doc, "步骤")
    update_row(
        chain,
        "DeepSeek 意图拆解",
        {
            3: "识别材料、规格特征、地区、目标时间、数量、单位与项目阶段等信息，并输出结构化 JSON；"
            "无法识别的字段留空，不猜测不存在的字段。"
        },
        column=1,
    )
    update_row(
        chain,
        "建议价计算",
        {1: "项目阶段系数计算"},
        column=1,
    )

    # --- 3.2 系统功能表 ---------------------------------------------------
    funcs = find_table(doc, "模块")
    update_row(
        funcs,
        "自然语言输入",
        {
            2: "支持输入材料名称、规格型号、数量、单位、报价城市、目标时间与项目阶段，支持补充修改；"
            "输入框占位文案与通宝提示语均包含“项目阶段”，并提供一键复制询价条件模板。"
        },
        column=1,
    )
    update_row(
        funcs,
        "快速思考模式",
        {
            2: "输入区固定展示“快速思考（推荐）”控件，作为本期 AI 询价唯一思考模式并默认选中，"
            "不提供“深度思考”切换；所选模式写入 reasoningMode，用于后续内容质量分析。"
        },
        column=1,
    )
    update_row(
        funcs,
        "项目阶段采集",
        {
            2: "每次新建询价会话进入正常结果生成流程时，在询价建议结果生成后采集一次项目阶段；"
            "用户选择后按所选阶段系数重新计算并更新建议价。同一会话内修改条件后再次提交或重新查询时沿用已选阶段，"
            "不重复采集；无数据、规格不足或非材料询价场景不采集。选项为施工图预算、投资估算、设计概算与工程竣工结算，"
            "用户未填写或跳过时默认施工图预算。"
        },
        column=1,
    )
    update_row(
        funcs,
        "通宝思考状态",
        {
            2: "首次提交材料询价条件后优先展示通宝思考状态与推理步骤，通宝形象保持轻微上下浮动、左右停顿的动态效果；"
            "思考结束后先生成询价建议结果，随后再展示项目阶段选择，思考结束前不提前弹出阶段选择。"
        },
        column=1,
    )
    update_row(
        funcs,
        "询价建议结果",
        {
            2: "展示建议价区间、材料与规格、数量单位、品牌、建议价、供应商信息与回复备注；"
            "建议价区间及各条建议价按金额档位格式化：低于 1000 元保留两位小数，1000 元（含）至 10000 元（含）保留整数，"
            "高于 10000 元按十位取整。用户输入的条件项少于 5 项时，在建议价区间旁提示“由于搜索条件不足，建议价区间或存在偏差”。"
        },
        column=1,
    )
    update_row(
        funcs,
        "复制与反馈",
        {2: "提供复制、点赞、点踩与重新查询入口；点赞或点踩后图标颜色加深并记录日志事件；结果区不提供“修改条件”入口。"},
        column=1,
    )
    update_row(
        funcs,
        "重新查询",
        {
            2: "点击重新查询会重新生成询价结果并消耗 1 条询价使用次数，操作前需明确提示；"
            "重新生成需同步后端数据，由后端通知第三方模型重新筛选候选供应商并更换候选报价。"
        },
        column=1,
    )
    update_row(
        funcs,
        "盖章报价预览",
        {2: "按询价部模板生成盖章报价预览，包含造价通清晰标识、整页造价通水印、供应商报价表与固定位置询价专用章。"},
        column=1,
    )
    update_row(
        funcs,
        "导出盖章报价文件",
        {
            2: "提供“下载PDF报价文件”入口，导出格式固定为 PDF；文件名按“项目或区域 + 询价盖章报价 + 时间戳”规则生成；"
            "导出文件包含造价通水印与固定位置盖章，并在文件末尾展示“该文件由 AI 生成，仅供报价参考。”提示语。"
        },
        column=1,
    )
    update_row(
        funcs,
        "历史询价记录",
        {
            2: "展示询价时间、询价条件、推荐结果与操作列，不展示状态列；操作列提供查看结果与删除入口。"
            "历史询价记录保留一年，超过一年自动清理。"
        },
        column=1,
    )
    update_row(
        funcs,
        "使用次数",
        {
            2: "每日额度为 AI 询价 10 条、AI 查价 50 条；首次查询、修改后提交、重新询价各扣减 1 次；"
            "查看历史与人工询价不扣次；本期不绑定套餐，不展示赠送次数。"
        },
        column=1,
    )
    insert_row_after(
        funcs,
        "询价建议结果",
        [
            "结果展示",
            "条件不足提示",
            "用户输入的条件项少于 5 项时，在建议价区间旁提示“由于搜索条件不足，建议价区间或存在偏差”，"
            "结果正常展示，不提示用户重新询价。",
            "P0",
        ],
        column=1,
    )
    insert_row_after(
        funcs,
        "导出盖章报价文件",
        [
            "导出",
            "盖章方式待批示",
            "用户申请盖章报价文件采用自动生成还是人工审批方式，需由领导批示后确定；"
            "批示前预览与导出统一标注“盖章生成方式待领导批示”。",
            "P1",
        ],
        column=1,
    )
    append_row(
        funcs,
        [
            "数据治理",
            "历史数据去重",
            "同一材料在不同月份存在重复或冲突历史数据留存时，先在前期数据治理阶段由 AI 完成材料数据去重与合并，"
            "避免与历史价倒推计算冲突；去重规则与结果由数据部门确认。",
            "P0",
        ],
    )
    insert_row_after(
        funcs,
        "价格来源提问",
        [
            "结果展示",
            "引用来源不展示",
            "AI 询价与 AI 查价结果区均不展示引用来源与资料来源列表。",
            "P0",
        ],
        column=1,
    )

    # --- 3.3 全局样式与视觉基线 -------------------------------------------
    replace_paragraph(
        doc,
        "对话页沿用造价通现有会话界面视觉",
        "对话页沿用造价通现有会话界面视觉：左侧为品牌与历史会话栏，右侧为对话区与输入区，"
        "输入区固定展示“AI 查价 / AI 询价”快捷入口、带“推荐”标识的“快速思考”模式控件与通宝形象；"
        "左侧底部明确展示每日额度 AI 询价 10 条 / 天、AI 查价 50 条 / 天。结果卡片使用白底卡片与表格式报价展示，"
        "建议价使用强调色；无法匹配与结果不足状态使用提示色区分，不使用告警弹窗打断用户操作。"
        "结果区不展示引用来源与资料来源列表，AI 询价与 AI 查价保持一致。"
        "首页、对话页与盖章报价预览页右上角均提供“造价通首页”跳转入口，用于从询价流程返回首页。",
    )

    # --- 3.4.2 对话式询价输入与条件识别 ------------------------------------
    replace_paragraph(
        doc,
        "用户在输入区用自然语言描述询价需求",
        "用户在输入区用自然语言描述询价需求，系统识别材料名称、规格型号、数量、单位、报价城市与目标时间，"
        "并尽量识别品牌与含税含运要求。检索条件不使用省份维度，原因是历史询价数据的省份信息不完整或缺失；"
        "材料名称、规格、数量、城市与目标时间不构成对省份字段的过滤条件。当用户未填写地区时，"
        "系统默认按广东省、广西壮族自治区、四川省作为召回范围，并对这些地区的候选结果增加排序权重；"
        "该默认范围不作为硬过滤条件，其他地区存在更匹配结果时仍可进入候选集合。"
        "输入框占位文案为“请输入材料名称、规格型号、数量、单位、报价地区、目标时间与项目阶段”，输入内容与识别结果在提交前可继续修改。"
        "对话页通宝提示语同步提示需要提供项目阶段，并在提示语旁提供复制入口，"
        "点击后可复制询价条件模板“材料名称 + 规格型号 + 数量 + 单位 + 报价地区 + 目标时间 + 项目阶段”，供用户直接粘贴使用。",
    )
    replace_paragraph(
        doc,
        "输入区固定展示思考模式控件",
        "输入区固定展示“快速思考（推荐）”控件，本期 AI 询价仅提供快速思考一种模式并默认选中，不提供“深度思考”切换；"
        "所选模式写入行为日志的 reasoningMode 属性，并用于后续内容质量分析。",
    )

    # --- 3.4.3 通宝思考状态与动态效果 --------------------------------------
    replace_paragraph(
        doc,
        "每次新建询价会话首次提交且进入正常结果生成流程时，通宝按",
        "每次新建询价会话首次提交且进入正常结果生成流程时，通宝按“接收需求 — 理解条件 — 匹配历史询价与市场价”优先展示思考过程；"
        "思考结束后先生成并展示询价建议结果，随后在结果卡片下方通过多轮对话采集一次本次询价所属的项目阶段，"
        "用户选择后按所选阶段系数重新计算并更新建议价。同一会话内修改条件后再次提交或点击重新查询时，沿用已选项目阶段，"
        "不重复采集；仅新建会话、重置会话或用户明确指定其他阶段时重新确定。无数据、规格不足或非材料询价场景不采集项目阶段。"
        "思考态由通宝形象与步骤文案组成，通宝形象保持轻微上下浮动、左右停顿的动画，用于表达思考状态；"
        "动画节奏需平稳，不出现明显跳动或闪烁。",
    )
    replace_paragraph(
        doc,
        "思考结束前不得提前弹出项目阶段选择",
        "项目阶段选择的展示位置固定为询价建议结果之后，不得在思考结束前或结果生成前提前弹出，"
        "避免打断用户对过程反馈与结果的感知。在系统设置减少动态效果（prefers-reduced-motion）时，"
        "动画需降级为静态展示，仅保留步骤文案变化，避免影响可访问性。",
    )

    # --- 3.4.5 系数匹配与建议价计算 ----------------------------------------
    insert_paragraph_after(
        find_paragraph(doc, "倒推约束：参与计算的最新价格与目标材料"),
        "数据治理要求：历史询价中同一材料在不同月份存在重复或冲突数据留存时，会与历史价倒推产生冲突"
        "（同一材料同一期存在多条口径不一致的价格记录，倒推链无法确定基准价格）。此类数据需在前期数据治理阶段由 AI 对材料做去重处理，"
        "按材料标准名称、规格特征、单位与报价地区识别同一材料，保留口径一致的有效记录并合并重复记录；"
        "去重规则、合并优先级与处理结果需由数据部门确认后再进入倒推与建议价计算。",
    )

    # --- 3.4.6 询价建议结果展示 -------------------------------------------
    replace_paragraph(
        doc,
        "结果区展示建议价区间与含税含运口径",
        "结果区展示建议价区间与含税含运口径，并在报价表中按行展示材料名称、规格型号、数量、单位、品牌、建议价、供应商信息与回复备注。"
        "建议价区间及各条建议价按金额档位格式化：低于 1000 元保留两位小数，1000 元（含）至 10000 元（含）保留整数，"
        "高于 10000 元按十位取整。候选结果不足 3 条时不得提示用户重新询价，需直接展示全部候选记录；"
        "建议价之间的差异仅通过价格数值与排序体现，不额外标注价格基础来源。当用户输入的条件项少于 5 项时，"
        "在建议价区间旁展示提示语“由于搜索条件不足，建议价区间或存在偏差”（见 图 3-7），"
        "结果正常展示，不阻断查询、不提示用户重新询价。",
    )
    replace_paragraph(
        doc,
        "项目阶段通过多轮对话采集",
        "项目阶段通过多轮对话采集，展示位置固定在询价建议结果之后：每次新建询价会话进入正常结果生成流程时，"
        "先生成询价建议结果，再在结果卡片下方询问一次本次询价所属的项目阶段，选项为施工图预算、投资估算、设计概算与工程竣工结算；"
        "用户选择后按所选阶段系数重新计算并刷新建议价，未填写或跳过时默认按施工图预算阶段处理。"
        "同一会话内修改条件后再次提交或点击重新查询时，沿用已选项目阶段，不重复弹出阶段采集；"
        "仅新建会话、重置会话或用户明确指定其他阶段时重新确定。无数据、规格不足或非材料询价场景不采集项目阶段。"
        "后端按“价格基础 × 时间系数 × 项目阶段系数”计算建议价，项目阶段系数仅用于后端计算，"
        "对话页与结果区不展示系数、换算公式与提示语。系数表如下：",
    )

    # --- 3.4.7 结果操作 ---------------------------------------------------
    replace_paragraph(
        doc,
        "重新查询图标悬停时展示说明文案",
        "重新查询图标悬停时展示说明文案“点击将重新生成询价结果，并消耗 1 条询价使用次数”，"
        "点击后按当前询价条件重新生成结果并扣减 1 次使用次数。重新查询需同步后端数据，"
        "由后端通知第三方模型重新筛选候选供应商并返回新一轮候选报价，刷新前后的候选供应商不得完全相同；"
        "刷新完成后在结果区提示“已通知后端并触发第三方模型重新筛选候选供应商”。"
        "结果区不提供“修改条件”入口，用户如需调整条件直接在输入框重新输入。所有反馈动作均需写入日志事件，见 3.6。",
    )

    # --- 3.4.11 盖章报价预览与导出 ----------------------------------------
    replace_paragraph(
        doc,
        "点击“导出盖章报价文件”进入盖章报价预览页",
        "点击“导出盖章报价文件”进入盖章报价预览页。预览页按询价部模板呈现：页面顶部展示造价通清晰标识与官网地址，"
        "主体为供应商报价表，整页叠加造价通水印，报价表上方按固定位置叠加询价专用章（盖章位置固定，不随报价行数变化），"
        "表格包含材料名称、规格型号、数量、单位、品牌、价格、供应商信息与回复备注。",
    )
    replace_paragraph(
        doc,
        "预览页提供“显示联系方式”开关",
        "预览页提供“显示联系方式”开关，用于控制预览内容是否展示联系方式字段；提供“下载PDF报价文件”按钮，"
        "导出格式固定为 PDF，文件名规则为“项目或区域_询价盖章报价_时间戳.pdf”，与历史记录中的询价条件保持一致；"
        "导出文件需包含造价通水印与固定位置盖章，并在文件末尾展示“该文件由 AI 生成，仅供报价参考。”提示语。"
        "系统不得导出图片、Word 或其他格式。",
    )
    insert_paragraph_after(
        find_paragraph(doc, "预览页提供“显示联系方式”开关"),
        "用户申请盖章报价文件采用自动生成方式还是人工审批方式，需由领导批示后确定；"
        "批示前预览页与导出文件统一标注“盖章生成方式待领导批示”，并按固定位置盖章展示，不擅自按某一种方式实现。",
    )

    # --- 3.4.12 历史询价记录与会话删除 ------------------------------------
    replace_paragraph(
        doc,
        "历史询价列表展示询价时间、询价条件、推荐结果、状态与操作",
        "历史询价列表展示询价时间、询价条件、推荐结果与操作列，不展示状态列。操作列提供“查看结果／演示对话”与“删除”两个入口，"
        "“删除”用于直接移除该条历史询价记录，需二次确认后生效。历史询价记录保留一年，超过一年自动清理。",
    )

    # --- 3.4.13 使用次数与扣次口径 ----------------------------------------
    replace_paragraph(
        doc,
        "对话页左侧展示使用次数，包括今日剩余与赠送剩余",
        "对话页左下角明确展示每日额度：AI 询价 10 条 / 天、AI 查价 50 条 / 天，并说明扣次规则："
        "首次查询、修改后提交、重新询价各扣 1 次；查看历史与人工询价不扣次。本期不绑定套餐，不展示套餐与赠送次数概念。"
        "次数不足时阻止继续发起查询，提示“查询次数已用完，可转人工询价”，并保留转人工询价入口。",
    )
    replace_paragraph(
        doc,
        "非材料询价输入、规格补充确认过程与查看历史记录不扣减次数",
        "非材料询价输入、规格补充确认过程与查看历史记录不扣减次数；重新查询按 1 次计。"
        "AI 询价与 AI 查价的每日额度独立计算，导出盖章报价文件不额外扣减询价次数。",
    )

    # --- 3.5 提示文案与状态反馈规范 ---------------------------------------
    copy_table = find_table(doc, ("场景", "图标或控件", "文案", "说明"))
    update_row(
        copy_table,
        "项目阶段确认",
        {
            1: "对话快捷选项（结果卡片下方）",
            3: "在询价建议结果生成后、结果卡片下方采集一次；同一会话内修改后再次提交或重新查询沿用已选阶段，不重复弹出。"
            "无数据、规格不足或非材料询价不采集。未填写或跳过默认施工图预算，不展示阶段系数与换算公式。",
        },
    )
    update_row(
        copy_table,
        "重新查询",
        {3: "鼠标悬停展示说明，点击扣 1 次；重新生成时由后端通知第三方模型重新筛选候选供应商。"},
    )
    update_row(
        copy_table,
        "思考模式",
        {2: "快速思考（推荐）", 3: "本期仅提供快速思考一种模式并默认选中，不提供“深度思考”切换；所选模式写入 reasoningMode。"},
    )
    insert_row_after(
        copy_table,
        "非材料询价",
        [
            "输入提示",
            "输入框占位文案",
            "请输入材料名称、规格型号、数量、单位、报价地区、目标时间与项目阶段",
            "占位文案包含项目阶段。",
        ],
    )
    insert_row_after(
        copy_table,
        "输入提示",
        [
            "提示语复制",
            "通宝提示语旁复制图标",
            "复制询价条件模板",
            "点击后复制“材料名称 + 规格型号 + 数量 + 单位 + 报价地区 + 目标时间 + 项目阶段”模板，供用户粘贴使用。",
        ],
    )
    insert_row_after(
        copy_table,
        "候选不足",
        [
            "条件不足",
            "建议价区间旁提示",
            "由于搜索条件不足，建议价区间或存在偏差",
            "输入条件项少于 5 项时展示，结果照常展示，不提示用户重新询价。",
        ],
    )
    insert_row_after(
        copy_table,
        "重新查询",
        [
            "刷新结果",
            "结果区提示",
            "已通知后端并触发第三方模型重新筛选候选供应商",
            "重新查询完成后展示，说明已更换候选报价数据。",
        ],
    )
    insert_row_after(
        copy_table,
        "次数不足",
        [
            "每日额度",
            "左侧额度区",
            "AI 询价 10 条 / 天 · AI 查价 50 条 / 天",
            "明确每日条数；本期不绑定套餐，不展示赠送次数。",
        ],
    )
    insert_row_after(
        copy_table,
        "PDF 导出",
        [
            "导出文件末尾",
            "导出文件页脚",
            "该文件由 AI 生成，仅供报价参考。",
            "PDF 文件末尾固定展示。",
        ],
    )
    insert_row_after(
        copy_table,
        "导出文件末尾",
        [
            "盖章方式",
            "预览页说明",
            "盖章生成方式待领导批示",
            "领导批示前固定展示，不擅自确定自动生成或人工审批。",
        ],
    )
    insert_row_after(
        copy_table,
        "结果数量",
        [
            "引用来源",
            "结果区",
            "不展示引用与资料来源",
            "AI 询价与 AI 查价结果区均不展示引用来源。",
        ],
    )

    # --- 3.6 日志埋点 -----------------------------------------------------
    logs = find_table(doc, "事件名")
    update_row(
        logs,
        "ai_inquiry_requery",
        {
            2: "previousRequestId、newRequestId、query、refinementRound、previousSupplierIds、newSupplierIds、"
            "thirdPartyRefined、reasoningMode"
        },
    )
    update_row(
        logs,
        "ai_inquiry_quote_export",
        {
            2: "fileName、fileType（PDF）、recommendationCount、includeContact、projectPhase、requestId、"
            "fileWatermark、aiGeneratedNotice"
        },
    )
    insert_row_after(
        logs,
        "ai_inquiry_query_result",
        [
            "ai_inquiry_project_phase_select",
            "用户在询价建议结果后选择或跳过项目阶段",
            "projectPhase、requestId、conversationId、selectionSource（user/default）、resultUpdated",
        ],
    )
    insert_row_after(
        logs,
        "ai_inquiry_requery",
        [
            "ai_inquiry_candidate_refresh",
            "重新查询触发后端通知第三方模型重新筛选候选供应商",
            "requestId、previousRequestId、refinementRound、previousSupplierIds、newSupplierIds、thirdPartyModel",
        ],
    )
    insert_row_after(
        logs,
        "ai_inquiry_query_submit",
        [
            "ai_inquiry_insufficient_condition_hint",
            "条件不足提示展示",
            "requestId、query、conditionCount、threshold、handling（result_kept）",
        ],
    )

    # --- 4.1 安全与隐私 ---------------------------------------------------
    privacy_anchor = find_paragraph(doc, "盖章报价文件需记录文件编号")
    insert_paragraph_after(
        privacy_anchor,
        "导出文件需包含造价通水印、固定位置盖章与“该文件由 AI 生成，仅供报价参考。”标识，"
        "避免文件被二次修改后冒充正式盖章报价文件。",
    )

    # --- 4.2 日志与审计 ---------------------------------------------------
    audit_anchor = find_paragraph(doc, "日志不得记录与合规要求冲突的用户敏感信息")
    insert_paragraph_after(
        audit_anchor,
        "项目阶段选择、重新查询触发的候选供应商更换、条件不足提示与导出文件水印及 AI 生成标识需写入日志事件，"
        "便于按 requestId 还原用户操作顺序与文件生成口径。",
    )

    # --- 4.5 风险边界 -----------------------------------------------------
    risks = find_table(doc, "风险")
    insert_row_after(
        risks,
        "系数映射重复数据",
        [
            "历史数据重复与倒推冲突",
            "同一材料不同月份存在重复或冲突历史数据时，倒推基准价格不确定，建议价会出现偏差；"
            "需在前期数据治理阶段由 AI 完成材料数据去重与合并，去重规则由数据部门确认。",
        ],
    )
    insert_row_after(
        risks,
        "导出文件合规",
        [
            "盖章生成方式未定",
            "盖章报价文件采用自动生成还是人工审批方式尚未批示，影响预览与导出流程实现方式；"
            "批示前统一标注“盖章生成方式待领导批示”，不按任一方式先行实现。",
        ],
    )

    # --- 4.6 评测与回归 ---------------------------------------------------
    harness = find_table(doc, "评测层级")
    update_row(
        harness,
        "前端交互",
        {
            2: "思考前置、结果后阶段采集、同会话阶段复用、结果操作、PDF 导出、转人工与异常提示。",
            3: "P0 Playwright 流程全部通过；项目阶段提示位于询价建议结果之后，重新查询不得重复采集项目阶段。",
        },
    )
    update_row(
        harness,
        "日志与审计",
        {
            2: "查询、结果、项目阶段选择、重新查询与候选更换、扣次、点赞、点踩、导出、转人工、价格来源提问与会话删除。"
        },
    )

    # --- 5.1 验收标准 -----------------------------------------------------
    acceptance = find_table(doc, ("编号", "验收项", "验收标准"))
    update_row(
        acceptance,
        "A6",
        {
            2: "首次提交后优先展示通宝思考状态与步骤文案，思考结束后先生成询价建议结果，项目阶段选择位于结果之后；"
            "输入区仅展示“快速思考（推荐）”一种模式并默认选中，所选模式正确写入 reasoningMode。"
            "通宝形象动态平稳，减少动态效果时降级为静态。"
        },
        column=0,
    )
    update_row(
        acceptance,
        "A12",
        {
            2: "在询价建议结果生成后采集一次项目阶段，用户选择后按对应阶段系数更新建议价；未填写或跳过时默认施工图预算；"
            "同一会话内修改后再次提交或重新查询沿用已选阶段，不重复采集；无数据、规格不足或非材料询价不采集。"
            "后端按对应阶段系数计算建议价，对话页与结果区不展示系数、换算公式与提示语。"
        },
        column=0,
    )
    update_row(
        acceptance,
        "A15",
        {
            2: "重新查询前提示消耗 1 次使用次数，点击后重新生成结果并扣减 1 次；重新生成需同步后端数据，"
            "由后端通知第三方模型重新筛选候选供应商，刷新后候选供应商与刷新前不同。"
        },
        column=0,
    )
    update_row(
        acceptance,
        "A18",
        {
            2: "预览页展示造价通清晰标识、整页造价通水印、报价表与固定位置询价专用章，可切换联系方式显示；"
            "点击“下载PDF报价文件”后导出格式为 PDF，文件名以 .pdf 结尾，"
            "文件末尾包含“该文件由 AI 生成，仅供报价参考。”提示语。"
        },
        column=0,
    )
    update_row(
        acceptance,
        "A19",
        {
            2: "历史询价列表展示询价时间、询价条件、推荐结果与操作列，不展示状态列；操作列提供删除入口，"
            "删除需二次确认，删除后条目移除并记录日志；历史记录保留一年，超过一年自动清理。"
        },
        column=0,
    )
    update_row(
        acceptance,
        "A20",
        {
            2: "每日额度为 AI 询价 10 条、AI 查价 50 条；首次查询、修改后提交、重新查询各扣 1 次；"
            "查看历史与人工询价不扣次，次数不足时有明确提示；本期不绑定套餐，不展示赠送次数。"
        },
        column=0,
    )
    update_row(
        acceptance,
        "A21",
        {
            2: "点赞、点踩、取消反馈、通宝互动、会话删除、项目阶段选择、重新查询触发的候选更换与条件不足提示均写入日志事件，"
            "属性完整可追溯。"
        },
        column=0,
    )
    update_row(
        acceptance,
        "A28",
        {2: "输入区默认选中“快速思考（推荐）”，本期不提供“深度思考”切换；模式上报与 reasoningMode 一致。"},
        column=0,
    )
    insert_row_after(
        acceptance,
        "A34",
        [
            "A35",
            "条件不足提示",
            "输入条件项少于 5 项时，建议价区间旁展示“由于搜索条件不足，建议价区间或存在偏差”，结果照常展示，不提示重新询价。",
        ],
        column=0,
    )
    insert_row_after(
        acceptance,
        "A35",
        [
            "A36",
            "技术链路口径",
            "技术实现链路与 2.5 一致：用户输入 → DeepSeek 意图拆解 → 提取材料名称与特征、时间、地区、项目阶段 → "
            "ES 模糊检索 → DeepSeek 候选筛选 → 项目阶段系数计算 → 前端展示。",
        ],
        column=0,
    )
    insert_row_after(
        acceptance,
        "A36",
        [
            "A37",
            "数据治理去重",
            "历史询价同一材料跨月份存在重复或冲突数据时，前期数据治理阶段可由 AI 完成去重与合并，去重规则可追溯，"
            "不与历史价倒推冲突。",
        ],
        column=0,
    )
    insert_row_after(
        acceptance,
        "A37",
        [
            "A38",
            "盖章生成方式",
            "盖章报价文件采用自动生成或人工审批方式在批示前不作实现约定；预览页与导出文件统一标注“盖章生成方式待领导批示”。",
        ],
        column=0,
    )
    insert_row_after(
        acceptance,
        "A38",
        [
            "A39",
            "输入与提示语含项目阶段",
            "输入框占位文案与通宝提示语均包含“项目阶段”，并提供复制入口，复制内容为"
            "“材料名称 + 规格型号 + 数量 + 单位 + 报价地区 + 目标时间 + 项目阶段”模板。",
        ],
        column=0,
    )
    insert_row_after(
        acceptance,
        "A39",
        ["A40", "引用来源不展示", "AI 询价与 AI 查价结果区均不展示引用来源与资料来源列表。"],
        column=0,
    )

    # --- 5.2 待确认事项 ---------------------------------------------------
    pending = find_table(doc, ("编号", "待确认事项", "影响", "优先级"))
    update_row(
        pending,
        "Q4",
        {
            1: "使用次数按查询次数统计已确认（AI 询价 10 条/天、AI 查价 50 条/天，本期不绑定套餐）；导出次数是否单独限频待确认。",
            2: "额度展示与限频逻辑的边界口径。",
            3: "P1",
        },
        column=0,
    )
    insert_row_after(
        pending,
        "Q8",
        [
            "Q9",
            "盖章报价文件采用自动生成还是人工审批方式，需领导批示。",
            "盖章预览与导出流程的实现方式无法定稿。",
            "P0",
        ],
        column=0,
    )
    insert_row_after(
        pending,
        "Q9",
        [
            "Q10",
            "历史询价同一材料跨月份重复数据的去重与合并规则，以及去重结果的确认方式。",
            "影响历史价倒推基准与建议价准确性。",
            "P0",
        ],
        column=0,
    )

    # --- 6.1 资料来源 -----------------------------------------------------
    sources = find_table(doc, "资料")
    insert_row_after(
        sources,
        "需求确认与评审记录",
        [
            "AI 询价第二次评审意见（2026-10-15）",
            "项目阶段后置、输入提示与复制、条件不足提示、历史记录与删除、每日额度、思考模式与引用展示、盖章方式待批示、"
            "导出水印与 AI 生成提示、重新查询更换候选、数据治理去重与技术链路口径的更新依据。",
        ],
    )

    # --- 变更履历 ---------------------------------------------------------
    for table in doc.tables[:3]:
        for row in table.rows:
            cells = row.cells
            if len(cells) < 4:
                continue
            if cells[1].text.strip() != "V1.2":
                continue
            description = cells[2]
            if "第二次评审" in description.text:
                break
            append_cell_paragraph(description, CHANGE_LOG_TEXT)
            break
        else:
            continue
        break

    doc.save(DOC_PATH)
    print(f"Saved: {DOC_PATH}")
    print(f"Backup: {BACKUP_PATH}")


if __name__ == "__main__":
    main()
