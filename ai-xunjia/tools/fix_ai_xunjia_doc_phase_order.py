"""Align the AI inquiry specification with the current result-then-phase flow."""

from __future__ import annotations

import os
import shutil
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from PIL import Image

ROOT = Path("/Users/zouxiaoli/Documents/ChatGPT/ai询价")
DOC_PATH = ROOT / "造价通AI询价需求规格说明书（V1.3）.docx"
BACKUP_PATH = ROOT / ".qa/造价通AI询价需求规格说明书（V1.3）_阶段顺序修正前备份.docx"
DESIGN_IMAGE = ROOT / ".qa/ai_xunjia_doc/pc/23_thinking_design.png"


def replace_paragraph_text(paragraph, old: str, new: str) -> None:
    if old not in paragraph.text:
        raise ValueError(f"未找到待替换段落：{old[:42]}")
    if len(paragraph.runs) != 1:
        raise ValueError(f"待替换段落格式超出预期：runs={len(paragraph.runs)}")
    paragraph.runs[0].text = paragraph.runs[0].text.replace(old, new)


def replace_cell_text(cell, old: str, new: str) -> None:
    for paragraph in cell.paragraphs:
        if old in paragraph.text:
            replace_paragraph_text(paragraph, old, new)
            return
    raise ValueError(f"未找到待替换表格文本：{old[:42]}")


def replace_zip_member(source: Path, destination: Path, member: str, payload: bytes) -> None:
    with zipfile.ZipFile(source, "r") as source_zip, zipfile.ZipFile(
        destination, "w", zipfile.ZIP_DEFLATED
    ) as destination_zip:
        replaced = False
        for item in source_zip.infolist():
            data = payload if item.filename == member else source_zip.read(item.filename)
            destination_zip.writestr(item, data)
            replaced = replaced or item.filename == member
        if not replaced:
            raise ValueError(f"DOCX 中未找到图片：{member}")


def patch_document() -> None:
    doc = Document(DOC_PATH)

    paragraph_replacements = [
        (
            "通宝思考状态按“准备、思考、项目阶段确认、结果承接”四个阶段设计。",
            "通宝思考状态按“准备、思考、结果承接、项目阶段确认”四个阶段设计。",
        ),
        (
            "思考结束前不得提前弹出项目阶段选择，避免打断用户对过程反馈的感知。"
            "思考结束后展示项目阶段选项，用户选择后再承接询价建议、缺少规格或无可用数据等结果。",
            "思考结束前不得提前弹出项目阶段选择，避免打断用户对过程反馈的感知。"
            "思考结束后先承接询价建议、缺少规格或无可用数据等结果；"
            "对于正常询价建议，结果承接后再展示项目阶段选项，用户确认后按所选阶段系数刷新建议价。"
            "缺少规格或无可用数据等兜底结果直接结束，不再追加阶段确认。",
        ),
        (
            "首次会话中，通宝思考结束后再询问本次询价所属的项目阶段",
            "首次会话中，通宝思考结束并展示询价建议结果后，再询问本次询价所属的项目阶段",
        ),
    ]
    for old, new in paragraph_replacements:
        matching = [paragraph for paragraph in doc.paragraphs if old in paragraph.text]
        if len(matching) != 1:
            raise ValueError(f"段落匹配数量异常：{len(matching)}，{old[:42]}")
        replace_paragraph_text(matching[0], old, new)

    table_replacements = [
        (
            0,
            4,
            2,
            "明确准备、思考、项目阶段确认与结果承接四个阶段",
            "明确准备、思考、结果承接与项目阶段确认四个阶段",
        ),
        (
            8,
            14,
            2,
            "首次会话先展示通宝思考状态；思考结束后，再通过多轮对话询问本次询价所属项目阶段。",
            "首次会话先展示通宝思考状态；思考结束后先承接询价结果，正常询价结果承接后再通过多轮对话询问本次询价所属项目阶段。",
        ),
        (
            8,
            15,
            2,
            "思考结束前不提前弹出项目阶段选择。",
            "思考结束后先承接询价结果，正常结果承接后再展示项目阶段选择。",
        ),
        (
            15,
            10,
            1,
            "思考结束，等待用户选择项目阶段",
            "结果承接后，等待用户选择项目阶段",
        ),
        (
            15,
            10,
            4,
            "项目阶段选项展示后提示，不在思考结束前出现。",
            "项目阶段选项在正常询价结果承接后展示，不在思考结束前或结果生成前出现。",
        ),
        (
            21,
            6,
            2,
            "首次提交后优先展示通宝思考状态与步骤文案，思考结束前不提前弹出项目阶段选择；",
            "首次提交后优先展示通宝思考状态与步骤文案，思考结束前不提前弹出项目阶段选择，正常询价结果承接后再展示阶段选择；",
        ),
        (
            21,
            12,
            2,
            "材料询价提交后通过多轮对话确认项目阶段",
            "材料询价提交后先展示通宝思考并承接询价结果，正常结果承接后再通过多轮对话确认项目阶段",
        ),
        (
            21,
            34,
            2,
            "准备、思考、项目阶段确认与结果承接按设计图顺序出现；",
            "准备、思考、结果承接与项目阶段确认按设计图顺序出现；",
        ),
    ]
    for table_index, row_index, cell_index, old, new in table_replacements:
        replace_cell_text(doc.tables[table_index].rows[row_index].cells[cell_index], old, new)

    temp_doc = DOC_PATH.with_suffix(".phase-order.tmp.docx")
    doc.save(temp_doc)

    if Image.open(DESIGN_IMAGE).size != (3200, 2000):
        raise ValueError("图 3-5 设计图尺寸不是 3200×2000")
    final_temp = DOC_PATH.with_suffix(".phase-order.final.docx")
    replace_zip_member(
        temp_doc,
        final_temp,
        "word/media/image6.png",
        DESIGN_IMAGE.read_bytes(),
    )
    os.replace(final_temp, DOC_PATH)
    temp_doc.unlink(missing_ok=True)


def main() -> None:
    if not BACKUP_PATH.exists():
        shutil.copy2(DOC_PATH, BACKUP_PATH)
    patch_document()
    print(BACKUP_PATH)
    print(DOC_PATH)


if __name__ == "__main__":
    main()
