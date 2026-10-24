"""Apply the V1.2 audit updates to the existing requirement document."""

from __future__ import annotations

import shutil
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


ROOT = Path("/Users/zouxiaoli/Documents/ChatGPT/ai询价")
DOC_PATH = ROOT / "造价通AI询价需求规格说明书（V1.2）.docx"
BACKUP_PATH = ROOT / ".qa/造价通AI询价需求规格说明书（V1.2）_审计更新前备份.docx"
PREVIEW_SHOT = ROOT / ".qa/ai_xunjia_doc/pc/13_preview_seal.png"


def set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def set_cell_text(cell, text: str) -> None:
    paragraph = cell.paragraphs[0]
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)
    for extra in list(cell.paragraphs[1:]):
        extra._element.getparent().remove(extra._element)


def replace_paragraph(doc: Document, old: str, new: str) -> None:
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == old:
            set_paragraph_text(paragraph, new)
            return
    raise RuntimeError(f"Paragraph not found: {old[:60]}")


def insert_paragraph_after(paragraph: Paragraph, text: str) -> Paragraph:
    new_xml = deepcopy(paragraph._p)
    paragraph._p.addnext(new_xml)
    new_paragraph = Paragraph(new_xml, paragraph._parent)
    set_paragraph_text(new_paragraph, text)
    return new_paragraph


def find_table(doc: Document, first_header: str) -> Table:
    for table in doc.tables:
        if table.rows and table.rows[0].cells[0].text.strip() == first_header:
            return table
    raise RuntimeError(f"Table not found: {first_header}")


def find_table_by_headers(doc: Document, headers: list[str]) -> Table:
    for table in doc.tables:
        if not table.rows:
            continue
        actual = [cell.text.strip() for cell in table.rows[0].cells]
        if actual[: len(headers)] == headers:
            return table
    raise RuntimeError(f"Table not found by headers: {headers}")


def insert_row_after(table: Table, row_index: int, values: list[str]) -> None:
    source_tr = table.rows[row_index]._tr
    new_tr = deepcopy(source_tr)
    source_tr.addnext(new_tr)
    new_row = table.rows[row_index + 1]
    for cell, value in zip(new_row.cells, values):
        set_cell_text(cell, value)


def rebuild_event_table(table: Table, rows: list[list[str]]) -> None:
    body_rows = table.rows[1:]
    template = deepcopy(body_rows[0]._tr)
    for row in body_rows:
        row._tr.getparent().remove(row._tr)
    last_tr = table.rows[0]._tr
    for values in rows:
        new_tr = deepcopy(template)
        last_tr.addnext(new_tr)
        last_tr = new_tr
    for row, values in zip(table.rows[1:], rows):
        for cell, value in zip(row.cells, values):
            set_cell_text(cell, value)


def replace_preview_image(doc: Document) -> None:
    caption = next(
        paragraph
        for paragraph in doc.paragraphs
        if paragraph.text.strip().startswith("图 3-14 盖章报价预览")
    )
    images_before = sum(
        1
        for paragraph in doc.paragraphs
        if paragraph._p is not caption._p
        and paragraph._p.getprevious() is not None
        and paragraph._p.getroottree().getpath(paragraph._p)
        < caption._p.getroottree().getpath(caption._p)
        and paragraph._p.findall(".//" + qn("w:drawing"))
    )
    if images_before < 1:
        raise RuntimeError("Preview figure image not found")
    shape = doc.inline_shapes[images_before - 1]
    blip = shape._inline.find(".//" + qn("a:blip"))
    rid = blip.get(qn("r:embed"))
    image_part = doc.part.related_parts[rid]
    image_part._blob = PREVIEW_SHOT.read_bytes()


def main() -> None:
    if not PREVIEW_SHOT.exists():
        raise FileNotFoundError(PREVIEW_SHOT)

    BACKUP_PATH.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(DOC_PATH, BACKUP_PATH)
    doc = Document(DOC_PATH)

    replace_paragraph(
        doc,
        "对话页沿用造价通现有会话界面视觉：左侧为品牌与历史会话栏，右侧为对话区与输入区，输入区固定展示“AI 查价 / AI 询价”快捷入口与通宝形象。结果卡片使用白底卡片与表格式报价展示，建议价使用强调色；无法匹配与结果不足状态使用提示色区分，不使用告警弹窗打断用户操作。首页、对话页与盖章报价预览页右上角均提供“造价通首页”跳转入口，用于从询价流程返回首页。",
        "对话页沿用造价通现有会话界面视觉：左侧为品牌与历史会话栏，右侧为对话区与输入区，输入区固定展示“AI 查价 / AI 询价”快捷入口、带“推荐”标识的“快速思考”模式控件与通宝形象。结果卡片使用白底卡片与表格式报价展示，建议价使用强调色；无法匹配与结果不足状态使用提示色区分，不使用告警弹窗打断用户操作。首页、对话页与盖章报价预览页右上角均提供“造价通首页”跳转入口，用于从询价流程返回首页。",
    )

    input_paragraph = next(
        paragraph
        for paragraph in doc.paragraphs
        if paragraph.text.strip().startswith("用户在输入区用自然语言描述询价需求")
    )
    insert_paragraph_after(
        input_paragraph,
        "输入区固定展示思考模式控件，默认选中“快速思考”并显示“推荐”标识；点击后可切换为“深度思考”。快速思考用于更快返回询价结果，深度思考增加地区差异与价格合理性复核。用户选择结果需写入行为日志的 reasoningMode 属性，并用于后续内容质量分析。",
    )

    replace_paragraph(
        doc,
        "结果区展示建议价区间与含税含运口径，并在报价表中按行展示材料名称、规格型号、数量、单位、品牌、建议价、供应商信息与回复备注。候选结果不足 3 条时不得提示用户重新询价，需直接展示全部候选记录；建议价之间的差异仅通过价格数值与排序体现，不额外标注价格基础来源。",
        "结果区展示建议价区间与含税含运口径，并在报价表中按行展示材料名称、规格型号、数量、单位、品牌、建议价、供应商信息与回复备注。建议价区间及各条建议价按金额档位格式化：低于 1000 元保留两位小数，1000 元（含）至 10000 元（含）保留整数，高于 10000 元按十位取整。候选结果不足 3 条时不得提示用户重新询价，需直接展示全部候选记录；建议价之间的差异仅通过价格数值与排序体现，不额外标注价格基础来源。",
    )

    replace_paragraph(
        doc,
        "首次会话中，用户提交材料询价条件后，通宝按“接收需求 — 理解条件 — 匹配历史询价与市场价”优先展示思考过程；思考结束后，再通过多轮对话采集项目阶段，最后进入询价结果展示。思考态由通宝形象与步骤文案组成，通宝形象保持轻微上下浮动、左右停顿的动画，用于表达思考状态；动画节奏需平稳，不出现明显跳动或闪烁。",
        "每次新建询价会话首次提交且进入正常结果生成流程时，通宝按“接收需求 — 理解条件 — 匹配历史询价与市场价”优先展示思考过程；思考结束后、结果生成前，通过多轮对话采集一次本次询价所属的项目阶段，随后进入询价结果展示。同一会话内修改条件后再次提交或点击重新查询时，沿用已选项目阶段，不重复采集；仅新建会话、重置会话或用户明确指定其他阶段时重新确定。无数据、规格不足或非材料询价场景不采集项目阶段。思考态由通宝形象与步骤文案组成，通宝形象保持轻微上下浮动、左右停顿的动画，用于表达思考状态；动画节奏需平稳，不出现明显跳动或闪烁。",
    )

    replace_paragraph(
        doc,
        "项目阶段通过多轮对话采集：首次会话中，通宝思考结束后再询问本次询价所属的项目阶段，选项为施工图预算、投资估算、设计概算与工程竣工结算；用户未填写或跳过时默认按施工图预算阶段处理。后端按“价格基础 × 时间系数 × 项目阶段系数”计算建议价，项目阶段系数仅用于后端计算，对话页与结果区不展示系数、换算公式与提示语。系数表如下：",
        "项目阶段通过多轮对话采集：每次新建询价会话首次提交且进入正常结果生成流程时，在通宝思考结束、结果生成前询问一次本次询价所属的项目阶段，选项为施工图预算、投资估算、设计概算与工程竣工结算；用户未填写或跳过时默认按施工图预算阶段处理。同一会话内修改条件后再次提交或点击重新查询时，沿用已选项目阶段，不重复弹出阶段采集；仅新建会话、重置会话或用户明确指定其他阶段时重新确定。无数据、规格不足或非材料询价场景不采集项目阶段。后端按“价格基础 × 时间系数 × 项目阶段系数”计算建议价，项目阶段系数仅用于后端计算，对话页与结果区不展示系数、换算公式与提示语。系数表如下：",
    )

    replace_paragraph(
        doc,
        "当用户询问价格来源、建议价依据时，按业务口径默认不回复，系统不展示价格基础、系数名称、系数值、市场价切换过程与内部计算规则。该场景需保留用户提问记录，行为日志需记录用户提问，用于后续评估是否需要补充对外说明话术。",
        "当用户询问价格来源、建议价依据时，按业务口径默认不回复，系统不展示价格基础、系数名称、系数值、市场价切换过程与内部计算规则。该场景需保留用户提问记录，并写入 ai_inquiry_price_source_question 事件，用于后续评估是否需要补充对外说明话术。",
    )

    replace_paragraph(
        doc,
        "当历史询价与造价通市场价均无可用数据时，结果区展示“暂无法自动查询到可用数据”，说明可能是目标时间过久或数据库暂未收录，并提供调整目标时间、扩大报价区域与转人工询价三个出口。系统不得在无数据场景生成建议价与供应商报价。",
        "当历史询价与造价通市场价均无可用数据时，结果区展示“暂无法自动查询到可用数据”，说明可能是目标时间过久或数据库暂未收录，并提供调整目标时间、扩大报价区域与转人工询价三个出口。无数据对话页顶部需展示“您的输入将被收集用于模型训练”提示条及“关闭”入口，用户关闭后本会话不再展示。系统不得在无数据场景生成建议价与供应商报价。",
    )

    replace_paragraph(
        doc,
        "预览页提供“显示联系方式”开关，用于控制预览内容是否展示联系方式字段；提供“下载演示文件”按钮导出文件，导出文件名规则为“项目或区域_询价盖章报价_时间戳”，与历史记录中的询价条件保持一致。",
        "预览页提供“显示联系方式”开关，用于控制预览内容是否展示联系方式字段；提供“下载PDF报价文件”按钮，导出格式固定为 PDF，文件名规则为“项目或区域_询价盖章报价_时间戳.pdf”，与历史记录中的询价条件保持一致。系统不得导出图片、Word 或其他格式。",
    )

    replace_paragraph(
        doc,
        "点赞、点踩、取消反馈、通宝互动点击、会话删除需全部落库，字段包含事件名、事件时间、会话标识与业务上下文。",
        "查询提交、结果生成、重新查询、次数扣减、点赞、点踩、取消反馈、通宝互动点击、PDF 导出、转人工询价、价格来源提问与会话删除需全部落库，字段包含事件名、事件时间、会话标识与业务上下文。",
    )
    replace_paragraph(
        doc,
        "查询、重新查询、次数扣减、导出与转人工询价需可追溯，能够按会话还原用户操作顺序。",
        "查询、结果生成、重新查询、次数扣减、PDF 导出与转人工询价需通过 §3.6 定义的事件名可追溯，并可通过 conversationId、requestId 按会话还原用户操作顺序。",
    )

    replace_paragraph(
        doc,
        "单次询价从提交到结果展示的等待时间需与思考状态动画节奏匹配，避免长时间无反馈。",
        "用户提交后 300ms 内必须出现明确的处理反馈。",
    )
    replace_paragraph(
        doc,
        "结果区优先渲染建议价与报价行，导出与历史记录加载不得阻塞对话主流程。",
        "正常服务状态下，材料询价从提交到结果展示的端到端耗时要求 P90 ≤ 10 秒、P95 ≤ 15 秒；无数据或规格不足结论 P95 ≤ 5 秒。上述指标按连续 20 次查询统计。",
    )
    replace_paragraph(
        doc,
        "思考动画与结果渲染需避免明显卡顿，页面滚动与结果渲染不出现高频重排。",
        "结果数据返回后，结果区首屏渲染 ≤ 1 秒；报价行不超过 10 行时新增渲染 ≤ 500ms。PDF 报价文件预览 ≤ 3 秒，PDF 文件生成 ≤ 5 秒（报价行不超过 10 行、常规企业网络环境）。思考动画帧率不低于 30 FPS，单个主线程长任务不超过 500ms，页面滚动与结果操作不得出现可感知卡顿。",
    )

    function_table = find_table_by_headers(doc, ["模块", "功能点", "需求描述", "优先级"])
    for row in function_table.rows:
        if row.cells[1].text.strip() == "项目阶段采集":
            set_cell_text(
                row.cells[2],
                "每次新建询价会话首次提交且进入正常结果生成流程时，在通宝思考结束、结果生成前采集一次项目阶段；同一会话内修改条件后再次提交或重新查询时沿用已选阶段，不重复采集。无数据、规格不足或非材料询价场景不采集。选项为施工图预算、投资估算、设计概算与工程竣工结算，用户未填写或跳过时默认施工图预算。",
            )
        if row.cells[1].text.strip() == "询价建议结果":
            set_cell_text(
                row.cells[2],
                "展示建议价区间、材料与规格、数量单位、品牌、建议价、供应商信息与回复备注；建议价区间及各条建议价按金额档位格式化：低于 1000 元保留两位小数，1000 元（含）至 10000 元（含）保留整数，高于 10000 元按十位取整。",
            )
    insert_row_after(
        function_table,
        2,
        [
            "输入与识别",
            "快速思考模式",
            "输入区固定展示“快速思考”控件并带“推荐”标识，默认选中；点击可切换“快速思考 / 深度思考”。快速思考更快返回结果，深度思考增加地区差异与价格合理性复核，所选模式写入 reasoningMode。",
            "P0",
        ],
    )
    insert_row_after(
        function_table,
        3,
        [
            "安全与隐私",
            "模型训练提示条",
            "对话页顶部展示“您的输入将被收集用于模型训练”提示条，并提供“关闭”入口；无数据结果等对话状态保持一致展示。",
            "P0",
        ],
    )
    for row in function_table.rows:
        if row.cells[1].text.strip() == "导出盖章报价文件":
            set_cell_text(row.cells[2], "提供“下载PDF报价文件”入口，导出格式固定为 PDF；文件名按“项目或区域 + 询价盖章报价 + 时间戳”规则生成。")

    prompt_table = find_table_by_headers(doc, ["场景", "图标或控件", "文案", "说明"])
    for row in prompt_table.rows:
        label = row.cells[0].text.strip()
        if label == "结果数量":
            set_cell_text(row.cells[3], "按实际候选数量展示；建议价低于 1000 元保留两位小数，1000 元至 10000 元保留整数，高于 10000 元按十位取整。")
        if label == "复制":
            set_cell_text(row.cells[3], "复制建议价区间与供应商价格；金额按低于 1000 元保留两位小数、1000 元至 10000 元保留整数、高于 10000 元按十位取整的规则输出。")
        if label == "项目阶段确认":
            set_cell_text(row.cells[3], "每次新建会话首次提交且进入正常结果流程后采集一次；同一会话内修改后再次提交或重新查询沿用已选阶段，不重复弹出。无数据、规格不足或非材料询价不采集。未填写或跳过默认施工图预算，不展示阶段系数与换算公式。")
    for values in [
        ["模型训练提示", "顶部提示条", "您的输入将被收集用于模型训练。关闭", "用户可主动关闭；无数据页及对话状态一致展示。"],
        ["思考模式", "输入区控件", "快速思考（推荐） / 深度思考", "默认快速思考；点击可切换，所选模式写入 reasoningMode。"],
        ["PDF 导出", "下载按钮", "下载PDF报价文件", "导出格式固定为 PDF，文件名以 .pdf 结尾。"],
    ]:
        insert_row_after(prompt_table, len(prompt_table.rows) - 1, values)

    event_table = find_table_by_headers(doc, ["事件名", "触发时机", "关键属性"])
    rebuild_event_table(
        event_table,
        [
            ["ai_inquiry_query_submit", "用户提交首次查询或修改后的询价条件", "conversationId、requestId、query、materialName、spec、quantity、unit、region、targetTime、projectPhase、reasoningMode、actionType"],
            ["ai_inquiry_query_result", "询价结果生成完成", "requestId、resultStatus、recommendationCount、suggestedPriceMin、suggestedPriceMax、durationMs"],
            ["ai_inquiry_requery", "用户点击重新查询", "previousRequestId、newRequestId、query、reasoningMode"],
            ["ai_inquiry_credit_deduct", "查询次数扣减成功", "scene（query/modify/requery）、amount、beforeRemaining、afterRemaining、bonusRemaining、requestId"],
            ["ai_inquiry_like", "用户点击点赞并选中", "feedbackType、query、materialName、spec、recommendationCount、recommendationIds、reasoningMode"],
            ["ai_inquiry_dislike", "用户点击点踩并选中", "同上，feedbackType 为 dislike"],
            ["ai_inquiry_feedback_cancel", "用户取消已选中的点赞或点踩", "previousFeedback、query、materialName、recommendationCount"],
            ["ai_inquiry_tongbao_click", "用户点击通宝形象或结果区通宝头像", "入口位置、随机话术标识、当前会话与询价条件"],
            ["ai_inquiry_quote_export", "用户下载或导出盖章报价文件", "fileName、fileType（PDF）、recommendationCount、includeContact、projectPhase、requestId"],
            ["ai_inquiry_manual_submit", "用户提交转人工询价", "source（result/no_data/incomplete/supplier_contact）、contactName、contactPhoneMasked、expectedResponseTime、quoteBasis、requirements、requestId"],
            ["ai_inquiry_price_source_question", "用户询问价格来源或建议价依据", "query、requestId、conversationId、handling（default_no_reply）"],
            ["ai_inquiry_conversation_delete", "用户在历史会话中确认删除", "conversationTitle、query、wasActive"],
        ],
    )

    acceptance_table = find_table_by_headers(doc, ["编号", "验收项", "验收标准"])
    for row in acceptance_table.rows:
        label = row.cells[0].text.strip()
        if label == "A6":
            set_cell_text(row.cells[2], "首次提交后优先展示通宝思考状态与步骤文案，思考结束前不提前弹出项目阶段选择；输入区默认展示“快速思考（推荐）”，可切换“深度思考”，所选模式正确写入 reasoningMode。通宝形象动态平稳，减少动态效果时降级为静态。")
        if label == "A11":
            set_cell_text(row.cells[2], "结果卡片展示建议价区间、材料与规格、数量单位、品牌、建议价、供应商信息与回复备注；建议价低于 1000 元保留两位小数，1000 元（含）至 10000 元（含）保留整数，高于 10000 元按十位取整。")
        if label == "A12":
            set_cell_text(row.cells[2], "每次新建会话首次提交且进入正常结果生成流程时采集一次项目阶段，未填写或跳过时默认施工图预算；同一会话内修改后再次提交或重新查询沿用已选阶段，不重复采集；无数据、规格不足或非材料询价不采集。后端按对应阶段系数计算建议价，对话页与结果区不展示系数、换算公式与提示语。")
        if label == "A18":
            set_cell_text(row.cells[2], "预览页展示造价通清晰标识、报价表与询价专用章，可切换联系方式显示；点击“下载PDF报价文件”后导出格式为 PDF，文件名以 .pdf 结尾。")
    for values in [
        ["A27", "模型训练提示条", "无数据页及对话状态顶部展示“您的输入将被收集用于模型训练”提示条与“关闭”入口，关闭状态与展示状态可验证。"],
        ["A28", "思考模式", "输入区默认选中“快速思考（推荐）”，可切换“深度思考”；两种模式的步骤文案与 reasoningMode 上报一致。"],
        ["A29", "性能门槛", "连续执行20次：提交反馈≤300ms；结果P90≤10秒、P95≤15秒；无数据或规格不足结论P95≤5秒；结果返回后首屏渲染≤1秒；PDF预览≤3秒、生成≤5秒；动画≥30FPS且主线程长任务≤500ms。"],
    ]:
        insert_row_after(acceptance_table, len(acceptance_table.rows) - 1, values)

    replace_preview_image(doc)
    doc.save(DOC_PATH)
    doc_after = Document(DOC_PATH)
    print(f"Saved: {DOC_PATH}")
    print(f"Backup: {BACKUP_PATH}")
    print(f"Inline images: {len(doc_after.inline_shapes)}")


if __name__ == "__main__":
    main()
