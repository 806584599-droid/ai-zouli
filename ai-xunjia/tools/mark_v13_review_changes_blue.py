"""Mark V1.3 review edits in blue while preserving all other run formatting."""

from __future__ import annotations

import argparse
import difflib
import os
import re
import shutil
from dataclasses import dataclass
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.document import Document as DocumentObject
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.shared import RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph


ROOT = Path("/Users/zouxiaoli/Documents/ChatGPT/ai询价")
DOC_PATH = ROOT / "造价通AI询价需求规格说明书（V1.3）.docx"
BASELINE_PATH = (
    ROOT / ".qa/造价通AI询价需求规格说明书（V1.3）_评审修订前备份_20261020.docx"
)
PRE_MARK_BACKUP_PATH = (
    ROOT / ".qa/造价通AI询价需求规格说明书（V1.3）_蓝色标注前备份_20261020.docx"
)
BLUE = RGBColor(0x00, 0x00, 0xFF)


@dataclass(frozen=True)
class BodyBlock:
    kind: str
    signature: str
    element: Paragraph | Table
    paragraph_index: int | None


def normalize(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def table_header_key(table: Table) -> str:
    return "\x1f".join(normalize(cell.text) for cell in table.rows[0].cells)


def table_row_key(row) -> str:
    return "\x1e".join(normalize(cell.text) for cell in row.cells)


def iter_body_blocks(doc: DocumentObject) -> list[BodyBlock]:
    blocks: list[BodyBlock] = []
    paragraph_index = 0
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            paragraph = Paragraph(child, doc)
            text = normalize(paragraph.text)
            if text:
                blocks.append(
                    BodyBlock(
                        "paragraph",
                        f"p:{text}",
                        paragraph,
                        paragraph_index,
                    )
                )
            paragraph_index += 1
        elif child.tag == qn("w:tbl"):
            table = Table(child, doc)
            blocks.append(
                BodyBlock("table", f"t:{table_header_key(table)}", table, None)
            )
    return blocks


def changed_paragraphs(
    baseline: DocumentObject, current: DocumentObject
) -> list[Paragraph]:
    baseline_blocks = iter_body_blocks(baseline)
    current_blocks = iter_body_blocks(current)
    matcher = difflib.SequenceMatcher(
        a=[block.signature for block in baseline_blocks],
        b=[block.signature for block in current_blocks],
        autojunk=False,
    )

    targets: list[Paragraph] = []
    for tag, _a_start, _a_end, b_start, b_end in matcher.get_opcodes():
        if tag == "equal":
            continue
        for block in current_blocks[b_start:b_end]:
            if block.kind == "paragraph":
                targets.append(block.element)
    return targets


def changed_paragraph_indices(
    baseline: DocumentObject, current: DocumentObject
) -> list[int]:
    baseline_blocks = iter_body_blocks(baseline)
    current_blocks = iter_body_blocks(current)
    matcher = difflib.SequenceMatcher(
        a=[block.signature for block in baseline_blocks],
        b=[block.signature for block in current_blocks],
        autojunk=False,
    )

    indices: list[int] = []
    for tag, _a_start, _a_end, b_start, b_end in matcher.get_opcodes():
        if tag == "equal":
            continue
        for block in current_blocks[b_start:b_end]:
            if block.kind == "paragraph":
                if block.paragraph_index is None:
                    raise ValueError("变更段落缺少正文索引")
                indices.append(block.paragraph_index)
    return indices


def pair_replaced_rows(
    baseline_rows: list, current_rows: list
) -> tuple[list[tuple[int, int]], list[int]]:
    """Pair changed rows by content similarity; return unmatched current indices."""
    candidates: list[tuple[float, int, int]] = []
    baseline_keys = [table_row_key(row) for row in baseline_rows]
    current_keys = [table_row_key(row) for row in current_rows]
    for baseline_index, baseline_key in enumerate(baseline_keys):
        for current_index, current_key in enumerate(current_keys):
            ratio = difflib.SequenceMatcher(
                a=baseline_key, b=current_key, autojunk=False
            ).ratio()
            candidates.append((ratio, baseline_index, current_index))

    candidates.sort(key=lambda item: (-item[0], item[1], item[2]))
    used_baseline: set[int] = set()
    used_current: set[int] = set()
    pairs: list[tuple[int, int]] = []
    for _ratio, baseline_index, current_index in candidates:
        if baseline_index in used_baseline or current_index in used_current:
            continue
        used_baseline.add(baseline_index)
        used_current.add(current_index)
        pairs.append((baseline_index, current_index))

    pairs.sort()
    unmatched_current = [
        index for index in range(len(current_rows)) if index not in used_current
    ]
    return pairs, unmatched_current


def changed_table_cells(
    baseline: DocumentObject, current: DocumentObject
) -> list[tuple[int, int, int]]:
    """Return current-document table, row and column indices that changed."""
    if len(current.tables) < 1:
        raise ValueError("当前文档缺少变更履历表")

    baseline_tables = {table_header_key(table): table for table in baseline.tables}
    baseline_headers = set(baseline_tables)
    targets: list[tuple[int, int, int]] = []

    for current_index, current_table in enumerate(current.tables):
        # The change-history table was explicitly out of scope for this revision.
        if current_index == 0:
            continue

        header = table_header_key(current_table)
        baseline_table = baseline_tables.get(header)
        if baseline_table is None:
            for row_index in range(len(current_table.rows)):
                for column_index in range(len(current_table.rows[row_index].cells)):
                    targets.append((current_index, row_index, column_index))
            continue

        baseline_rows = list(baseline_table.rows)
        current_rows = list(current_table.rows)
        baseline_keys = [table_row_key(row) for row in baseline_rows]
        current_keys = [table_row_key(row) for row in current_rows]
        matcher = difflib.SequenceMatcher(
            a=baseline_keys, b=current_keys, autojunk=False
        )

        for tag, a_start, a_end, b_start, b_end in matcher.get_opcodes():
            if tag == "equal":
                continue
            if tag == "insert":
                for current_row_index in range(b_start, b_end):
                    for column_index in range(
                        len(current_rows[current_row_index].cells)
                    ):
                        targets.append(
                            (current_index, current_row_index, column_index)
                        )
                continue
            if tag == "delete":
                continue

            baseline_slice = baseline_rows[a_start:a_end]
            current_slice = current_rows[b_start:b_end]
            pairs, unmatched_current = pair_replaced_rows(
                baseline_slice, current_slice
            )
            for baseline_row_index, local_current_index in pairs:
                baseline_cells = baseline_slice[baseline_row_index].cells
                current_cells = current_slice[local_current_index].cells
                if len(baseline_cells) != len(current_cells):
                    raise ValueError(
                        f"表格 {current_index} 替换行单元格数量不一致："
                        f"{len(baseline_cells)} != {len(current_cells)}"
                    )
                for column_index, (baseline_cell, current_cell) in enumerate(
                    zip(baseline_cells, current_cells)
                ):
                    if normalize(baseline_cell.text) != normalize(current_cell.text):
                        targets.append(
                            (
                                current_index,
                                b_start + local_current_index,
                                column_index,
                            )
                        )
            for local_current_index in unmatched_current:
                current_row_index = b_start + local_current_index
                for column_index in range(
                    len(current_rows[current_row_index].cells)
                ):
                    targets.append(
                        (current_index, current_row_index, column_index)
                    )

    if not baseline_headers:
        raise ValueError("评审前备份缺少可匹配表格")
    return sorted(set(targets))


def paint_xml_runs(container) -> int:
    painted = 0
    for run_element in container.iter(qn("w:r")):
        run_properties = run_element.get_or_add_rPr()
        color = run_properties.get_or_add_color()
        color.val = BLUE
        for attribute in ("w:themeColor", "w:themeTint", "w:themeShade"):
            color.attrib.pop(qn(attribute), None)
        painted += 1
    return painted


def paint_xml(
    document_xml: bytes,
    paragraph_indices: list[int],
    cells: list[tuple[int, int, int]],
) -> tuple[bytes, int]:
    root = parse_xml(document_xml)
    body_paragraphs = root.body.findall(qn("w:p"))
    body_tables = root.body.findall(qn("w:tbl"))

    painted = 0
    for paragraph_index in paragraph_indices:
        painted += paint_xml_runs(body_paragraphs[paragraph_index])

    for table_index, row_index, column_index in cells:
        table_rows = body_tables[table_index].findall(qn("w:tr"))
        table_cells = table_rows[row_index].findall(qn("w:tc"))
        painted += paint_xml_runs(table_cells[column_index])

    xml = root.xml.encode("utf-8")
    return xml, painted


def write_package_preserving_fonts(
    source_path: Path,
    output_path: Path,
    document_xml: bytes,
) -> None:
    temporary_path = output_path.with_name(f".{output_path.name}.blue.tmp")
    try:
        with ZipFile(source_path, "r") as source_archive, ZipFile(
            temporary_path, "w"
        ) as output_archive:
            for entry in source_archive.infolist():
                data = source_archive.read(entry.filename)
                if entry.filename == "word/document.xml":
                    data = document_xml
                output_archive.writestr(entry, data)
        os.replace(temporary_path, output_path)
    finally:
        if temporary_path.exists():
            temporary_path.unlink()


def run_color(run_element) -> str | None:
    run_properties = run_element.find(qn("w:rPr"))
    if run_properties is None:
        return None
    color = run_properties.find(qn("w:color"))
    if color is None:
        return None
    return color.get(qn("w:val"))


def assert_blue(paragraphs: list[Paragraph], label: str) -> None:
    for paragraph in paragraphs:
        text_runs = [
            run
            for run in paragraph._p.iter(qn("w:r"))
            if "".join(node.text or "" for node in run.iter(qn("w:t")))
        ]
        if text_runs and any(run_color(run) != "0000FF" for run in text_runs):
            raise ValueError(f"{label}仍有未标蓝文本：{paragraph.text[:80]}")


def history_table_text(doc: DocumentObject) -> list[str]:
    return [table_row_key(row) for row in doc.tables[0].rows]


def summary_text(
    paragraphs: list[Paragraph], cells: list[tuple[int, int, int]]
) -> str:
    table_counts: dict[int, int] = {}
    for table_index, _row_index, _column_index in cells:
        table_counts[table_index] = table_counts.get(table_index, 0) + 1
    table_summary = "，".join(
        f"T{table_index}: {count} 个单元格"
        for table_index, count in sorted(table_counts.items())
    )
    return (
        f"目标段落 {len(paragraphs)} 个；目标表格单元格 {len(cells)} 个"
        + (f"（{table_summary}）" if table_summary else "")
    )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--apply",
        action="store_true",
        help="将蓝色标注写入文档；省略时仅试运行并输出统计。",
    )
    args = parser.parse_args()

    if PRE_MARK_BACKUP_PATH.exists():
        source_path = PRE_MARK_BACKUP_PATH
    else:
        source_path = DOC_PATH
        shutil.copy2(DOC_PATH, PRE_MARK_BACKUP_PATH)

    baseline = Document(BASELINE_PATH)
    current = Document(source_path)
    paragraphs = changed_paragraphs(baseline, current)
    paragraph_indices = changed_paragraph_indices(baseline, current)
    cells = changed_table_cells(baseline, current)

    if history_table_text(baseline) != history_table_text(current):
        raise ValueError("变更履历表与评审前备份不一致，停止标注")

    print(summary_text(paragraphs, cells))
    print("变更履历表：未纳入蓝色标注")

    if not args.apply:
        print("试运行完成，未修改文档。")
        return

    with ZipFile(source_path, "r") as archive:
        source_document_xml = archive.read("word/document.xml")
    output_document_xml, painted_runs = paint_xml(
        source_document_xml, paragraph_indices, cells
    )
    write_package_preserving_fonts(
        source_path, DOC_PATH, output_document_xml
    )

    marked = Document(DOC_PATH)
    assert_blue(
        [marked.paragraphs[index] for index in paragraph_indices], "变更段落"
    )
    for table_index, row_index, column_index in cells:
        cell = marked.tables[table_index].rows[row_index].cells[column_index]
        assert_blue(cell.paragraphs, f"T{table_index} R{row_index} C{column_index}")

    if history_table_text(baseline) != history_table_text(marked):
        raise ValueError("蓝色标注后变更履历表内容发生变化")

    print(f"已标蓝文本 run：{painted_runs} 个")
    print(DOC_PATH)


if __name__ == "__main__":
    main()
