"""Insert the project-phase confirmation screenshot into the V1.2 DOCX."""

from __future__ import annotations

import re
import shutil
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph


ROOT = Path("/Users/zouxiaoli/Documents/ChatGPT/ai询价")
DOC_PATH = ROOT / "造价通AI询价需求规格说明书（V1.2）.docx"
SHOT_PATH = ROOT / ".qa/ai_xunjia_doc/pc/18_phase_question.png"
BACKUP_PATH = ROOT / ".qa/造价通AI询价需求规格说明书（V1.2）_补图前备份.docx"

NEW_CAPTION = "图 3-8 项目阶段多轮问答确认截图"
TARGET_TEXT_PREFIX = "项目阶段系数与换算关系由数据部门维护"
LIST_ANCHOR = "询价建议结果与候选不足：图 3-6、图 3-7；"
NEW_LIST_ITEM = "项目阶段多轮问答确认：图 3-8；"


def set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def set_caption_run_format(run) -> None:
    run.font.size = Pt(9.5)
    run.font.bold = False
    run.font.color.rgb = RGBColor(0x5A, 0x5A, 0x5A)
    run.font.name = "Times New Roman"
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")
    fonts.set(qn("w:eastAsia"), "宋体")


def find_paragraph(doc: Document, predicate) -> Paragraph:
    for paragraph in doc.paragraphs:
        if predicate(paragraph.text.strip()):
            return paragraph
    raise RuntimeError("Target paragraph not found")


def renumber_caption(paragraph: Paragraph) -> None:
    match = re.match(r"^图 3-(\d+)\s", paragraph.text.strip())
    if not match:
        return
    number = int(match.group(1))
    if 8 <= number <= 15:
        updated = re.sub(
            r"^图 3-\d+",
            f"图 3-{number + 1}",
            paragraph.text.strip(),
            count=1,
        )
        set_paragraph_text(paragraph, updated)


def renumber_screenshot_list_item(paragraph: Paragraph) -> None:
    text = paragraph.text.strip()
    mappings = {
        "结果操作与反馈：图 3-8；": "结果操作与反馈：图 3-9；",
        "联系方式追问与不可用兜底：图 3-9、图 3-10；": "联系方式追问与不可用兜底：图 3-10、图 3-11；",
        "无可用数据与人工询价：图 3-11、图 3-12；": "无可用数据与人工询价：图 3-12、图 3-13；",
        "盖章报价预览：图 3-13；": "盖章报价预览：图 3-14；",
        "历史询价与会话删除：图 3-14、图 3-15。": "历史询价与会话删除：图 3-15、图 3-16。",
    }
    if text in mappings:
        set_paragraph_text(paragraph, mappings[text])


def main() -> None:
    if not SHOT_PATH.exists():
        raise FileNotFoundError(SHOT_PATH)

    doc = Document(DOC_PATH)
    if any(paragraph.text.strip() == NEW_CAPTION for paragraph in doc.paragraphs):
        print("Figure already present; no change made.")
        return

    BACKUP_PATH.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(DOC_PATH, BACKUP_PATH)

    # Renumber existing captions and list references before inserting the new figure.
    for paragraph in list(doc.paragraphs):
        renumber_caption(paragraph)
        renumber_screenshot_list_item(paragraph)

    target = find_paragraph(
        doc, lambda text: text.startswith(TARGET_TEXT_PREFIX)
    )

    image_paragraph = doc.add_paragraph()
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.paragraph_format.space_before = Pt(6)
    image_paragraph.paragraph_format.space_after = Pt(2)
    image_paragraph.paragraph_format.line_spacing = 1.0
    image_paragraph.add_run().add_picture(str(SHOT_PATH), width=Cm(16.0))

    caption_paragraph = doc.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.space_before = Pt(3)
    caption_paragraph.paragraph_format.space_after = Pt(8)
    caption_paragraph.paragraph_format.line_spacing = 1.15
    caption_run = caption_paragraph.add_run(NEW_CAPTION)
    set_caption_run_format(caption_run)

    target._p.addnext(image_paragraph._p)
    image_paragraph._p.addnext(caption_paragraph._p)

    list_anchor = find_paragraph(doc, lambda text: text == LIST_ANCHOR)
    new_list_xml = deepcopy(list_anchor._p)
    list_anchor._p.addnext(new_list_xml)
    new_list_paragraph = Paragraph(new_list_xml, list_anchor._parent)
    set_paragraph_text(new_list_paragraph, NEW_LIST_ITEM)

    doc.save(DOC_PATH)
    doc_after = Document(DOC_PATH)
    print(f"Saved: {DOC_PATH}")
    print(f"Backup: {BACKUP_PATH}")
    print(f"Inline images: {len(doc_after.inline_shapes)}")


if __name__ == "__main__":
    main()
