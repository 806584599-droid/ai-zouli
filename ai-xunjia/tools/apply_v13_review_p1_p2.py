"""Apply the V1.3 review changes for P1/P2 and P0-06 through P0-13."""

from __future__ import annotations

import copy
import shutil
from pathlib import Path

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph


ROOT = Path("/Users/zouxiaoli/Documents/ChatGPT/ai询价")
DOC_PATH = ROOT / "造价通AI询价需求规格说明书（V1.3）.docx"
BACKUP_PATH = ROOT / ".qa/造价通AI询价需求规格说明书（V1.3）_评审修订前备份_20261020.docx"


def set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run._element.getparent().remove(run._element)
    else:
        paragraph.add_run(text)


def replace_paragraph_contains(doc: Document, old: str, new: str) -> None:
    matches = [paragraph for paragraph in doc.paragraphs if old in paragraph.text]
    if len(matches) != 1:
        raise ValueError(f"段落匹配数量异常：{len(matches)}，{old[:50]}")
    set_paragraph_text(matches[0], new)


def find_paragraph(doc: Document, contains: str) -> Paragraph:
    matches = [paragraph for paragraph in doc.paragraphs if contains in paragraph.text]
    if len(matches) != 1:
        raise ValueError(f"定位段落失败：{len(matches)}，{contains[:50]}")
    return matches[0]


def set_cell_text(cell, text: str) -> None:
    paragraph = cell.paragraphs[0]
    set_paragraph_text(paragraph, text)
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)


def set_table_row(table: Table, row_index: int, values: list[str]) -> None:
    row = table.rows[row_index]
    if len(row.cells) != len(values):
        raise ValueError(f"表格列数不匹配：row={row_index}, cells={len(row.cells)}")
    for cell, value in zip(row.cells, values):
        set_cell_text(cell, value)


def append_table_row(table: Table, values: list[str], template_index: int = -1) -> None:
    template = table.rows[template_index]._tr
    new_tr = copy.deepcopy(template)
    table._tbl.append(new_tr)
    set_table_row(table, len(table.rows) - 1, values)


def insert_table_row_before_last(table: Table, values: list[str]) -> None:
    template = table.rows[-2]._tr
    new_tr = copy.deepcopy(template)
    table.rows[-1]._tr.addprevious(new_tr)
    set_table_row(table, len(table.rows) - 2, values)


def insert_paragraph_after(
    anchor: Paragraph,
    text: str,
    source: Paragraph | None = None,
) -> Paragraph:
    source = source or anchor
    new_element = copy.deepcopy(source._p)
    anchor._p.addnext(new_element)
    paragraph = Paragraph(new_element, anchor._parent)
    set_paragraph_text(paragraph, text)
    return paragraph


def insert_heading_after(doc: Document, anchor: Paragraph, text: str, level: int) -> Paragraph:
    paragraph = doc.add_paragraph(text, style=f"Heading {level}")
    anchor._p.addnext(paragraph._p)
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def add_toc_entry(doc: Document) -> None:
    candidates = [p for p in doc.paragraphs[:70] if p.text.strip() == "5.1 验收标准"]
    if len(candidates) != 1:
        raise ValueError(f"目录中的 5.1 条目数量异常：{len(candidates)}")
    insert_paragraph_after(candidates[0], "5.2 待确认事项", candidates[0])


def build_phase_pending_table(doc: Document) -> Table:
    source_table = doc.tables[13]
    table_element = copy.deepcopy(source_table._tbl)
    table = Table(table_element, doc)

    header_template = copy.deepcopy(table.rows[0]._tr)
    body_template = copy.deepcopy(table.rows[1]._tr)
    for row in list(table.rows):
        table._tbl.remove(row._tr)

    table._tbl.append(header_template)
    headers = ["编号", "事项", "当前状态与影响", "责任方与优先级"]
    for cell, value in zip(table.rows[0].cells, headers):
        set_cell_text(cell, value)

    rows = [
        [
            "Q1",
            "默认系数 cid=32 不存在",
            "待确认。映射表与系数表均无法定位该默认项，标准化失败或规则未命中时建议价链路仍可能中断。",
            "数据部门 / P0",
        ],
        [
            "Q2",
            "系数表缺少 2025–2026 数据",
            "待确认。需补齐目标年份，或明确最近可用月份的兜底规则；补齐前不得将当期建议价作为生产验收结论。",
            "数据部门 / P0",
        ],
        [
            "Q3",
            "城市名称未统一",
            "待确认。需提供省市标准字典，将广州/广州市、南宁/南宁市、成都/成都市、佛山/佛山市等写法统一后再检索。",
            "数据部门 / P0",
        ],
        [
            "Q4",
            "映射表重复与并列命中",
            "待确认。需清理完全重复行，并明确特征更具体、版本更新、映射 ID 升序等裁决规则。",
            "数据部门 / P0",
        ],
        [
            "Q5",
            "系数表随机/未采记录",
            "待确认。含“随机/未采/无数据源”的记录只能用于演示，生产验收前需清理或替换。",
            "数据部门 / P0",
        ],
        [
            "Q6",
            "项目阶段与结果展示顺序",
            "已确认。按“准备 → 思考 → 结果承接 → 阶段确认”执行，阶段确认后原地刷新同一张结果卡。",
            "产品 / P0",
        ],
        [
            "Q7",
            "AI 查价文案范围",
            "已确认。3.5.2 作为跨项目参考文案，不属于本期 AI 询价功能与发布门禁。",
            "产品 / 参考项",
        ],
        [
            "Q8",
            "输入是否用于模型训练",
            "已确认。本期不使用用户输入训练模型，正式页面不展示“用于模型训练”的提示条。",
            "产品与安全 / P0",
        ],
        [
            "Q9",
            "文档发布状态",
            "已确认。评审修订阶段按“草案（评审中）”管理，不影响既有变更履历。",
            "产品 / 文档管理",
        ],
        [
            "Q10",
            "长期记忆用户入口",
            "已确认。本期仅保留后台与审计能力，不展示查看、修改、清除入口，也不支持对话自动写入。",
            "产品 / P1",
        ],
    ]

    for values in rows:
        new_tr = copy.deepcopy(body_template)
        table._tbl.append(new_tr)
        set_table_row(table, len(table.rows) - 1, values)
    return table


def insert_phase_pending_section(doc: Document) -> None:
    acceptance_table = doc.tables[21]
    heading = doc.add_paragraph("5.2 待确认事项", style="Heading 2")
    intro = doc.add_paragraph(
        "以下事项单独登记，未关闭项不得作为开发完成或生产验收依据。"
        "已确认事项作为当前文档口径；变更履历不在本轮修订范围内。"
    )
    table = build_phase_pending_table(doc)

    acceptance_table._tbl.addnext(heading._p)
    heading._p.addnext(intro._p)
    intro._p.addnext(table._tbl)
    heading.paragraph_format.keep_with_next = True
    intro.paragraph_format.keep_with_next = True


def apply_paragraph_changes(doc: Document) -> None:
    replace_paragraph_contains(
        doc,
        "本文档用于说明造价通 AI 询价一期",
        "本文档用于说明造价通 AI 询价一期（标准数据询价）的功能范围、页面交互、业务规则、数据口径与日志要求，"
        "作为产品、交互、研发、数据、测试与运营的共同依据。文档中的规则以本期已确认口径为准；"
        "尚未确认的事项统一列入 5.2 待确认事项，在关闭前不得作为开发完成或生产验收依据。",
    )
    replace_paragraph_contains(
        doc,
        "业务侧同时明确了三项关键口径：建议价按价格基础 × 时间系数计算",
        "业务侧同时明确了三项关键口径：建议价按“价格基础 × 时间系数 × 项目阶段系数”计算；"
        "历史询价无结果时取造价通市场价继续按同一公式计算；用户追问供应商联系方式时在对话内给出联系方式，"
        "追问联系方式不可用时按记录较久或供应商更换电话答复并可转人工询价；用户问起价格来源时默认不回复。"
        "以上口径需要落到产品规则、页面文案与日志中。",
    )
    replace_paragraph_contains(
        doc,
        "结构化询价状态至少包含 material、spec、quantity、unit、region、targetTime、projectPhase 七类槽位",
        "结构化询价状态包含 material、spec、quantity、unit、region、targetTime、projectPhase、brand、"
        "taxIncluded、freightIncluded 十类业务槽位，并单独保存 resultRef。每个字段保存 value、status、source、"
        "confidence 与 updatedAt。status 取 confirmed、inferred、missing、conflicted、not_required；"
        "source 取 user、memory、system、inference。用户明确输入或确认的字段为 confirmed；"
        "模型推断但未经确认的字段为 inferred；关键字段缺失为 missing；前后输入冲突为 conflicted；"
        "当前场景不需要该字段时为 not_required。",
    )
    insert_paragraph_after(
        find_paragraph(doc, "status 取 confirmed、inferred、missing、conflicted、not_required"),
        "confidence 取 0 至 1。材料、规格、数量、单位、地区、目标时间与项目阶段等关键槽位，"
        "只有 confidence 不低于 0.8 且无冲突时才能参与计算；低于 0.8 时保持 inferred 并按缺失字段追问。"
        "出现 conflicted 时，以用户本轮明确输入为准，必要时追问一次，并在内部日志记录冲突字段、来源与解决结果。"
        "模型未返回 confidence 时不得默认置为高置信度。",
    )
    replace_paragraph_contains(
        doc,
        "上下文合并优先级为：用户本轮明确输入 > 会话已确认状态",
        "上下文合并优先级为：用户本轮明确输入 > 会话已确认状态 > 后台偏好 > 系统默认值。"
        "当用户修改材料、规格、数量、单位、地区、目标时间、品牌、含税或含运条件时，系统必须清空旧候选、"
        "旧建议价、旧供应商结果与导出引用，按合并后的新条件重新进入检索与计算；项目阶段作为会话级条件保留，"
        "避免重复追问。阶段变化只触发同一张结果卡的原地重算，不重新检索候选，也不扣减使用次数。",
    )
    replace_paragraph_contains(
        doc,
        "每次请求保存上下文版本快照与 requestId",
        "每次请求保存上下文版本快照与 requestId，记录 prompt 版本、模型版本、检索版本和系数版本。"
        "审计原文与模型上下文分离存储：审计原文加密保存 180 天，仅合规审计角色可访问；"
        "模型每轮只读取当前结构化状态、最近对话摘要和本轮询价结果，不直接读取全部历史原文。",
    )
    replace_paragraph_contains(
        doc,
        "对话页沿用造价通现有会话界面视觉",
        "对话页沿用造价通现有会话界面视觉：左侧为品牌与历史记录栏，右侧为对话区与输入区，"
        "输入区固定展示“AI 查价 / 通宝 AI”快捷入口与通宝形象。结果卡片使用白底卡片与表格式报价展示，"
        "建议价使用强调色；无法匹配与结果不足状态使用提示色区分，不使用告警弹窗打断用户操作。"
        "首页、对话页与盖章报价预览页右上角均提供“造价通首页”跳转入口，用于从询价流程返回首页。"
        "对外演示版隐藏“内部演示”按钮、规则抽屉与待批示文案；内部规则如展示，公式必须为三因子口径。",
    )
    replace_paragraph_contains(
        doc,
        "首页搜索工作台保留既有“AI 查价”按钮，并新增“AI 询价”按钮",
        "首页搜索工作台保留既有“AI 查价”按钮，并新增“通宝 AI”入口（功能名为 AI 询价），入口使用通宝头像标识。"
        "点击“通宝 AI”进入对话式询价页，进入后默认展示空状态与通宝问候语。对话页与盖章报价预览页右上角展示"
        "“造价通首页”入口，点击后返回造价通首页视图，不改变当前会话内容与使用次数。",
    )
    replace_paragraph_contains(
        doc,
        "用户在输入区用自然语言描述询价需求",
        "用户在输入区用自然语言描述询价需求，系统识别并结构化材料名称、规格型号、数量、单位、报价城市、"
        "目标时间、品牌、含税条件与含运条件，项目阶段可选填。检索条件不使用省份维度，原因是历史询价数据的"
        "省份信息不完整或缺失；材料名称、规格、数量、城市与目标时间不构成对省份字段的过滤条件。"
        "当用户未填写地区时，系统默认按广东省、广西壮族自治区、四川省作为召回范围，并对这些地区的候选结果"
        "增加排序权重；该默认范围不作为硬过滤条件，其他地区存在更匹配结果时仍可进入候选集合。"
        "输入框占位文案统一为“请输入材料名称、规格、数量、单位、地区、目标时间和项目阶段”，"
        "输入内容与识别结果在提交前可继续修改。",
    )
    replace_paragraph_contains(
        doc,
        "思考结束前不得提前弹出项目阶段选择",
        "思考结束前不得提前弹出项目阶段选择，避免打断用户对过程反馈的感知。思考结束后先承接询价建议、"
        "缺少规格或无可用数据等结果；对于正常询价建议，结果承接后再展示项目阶段选项。用户选择阶段后，"
        "系统原地刷新同一张结果卡，不得追加第二张结果卡；阶段变化不扣减次数。缺少规格或无可用数据等兜底结果"
        "直接结束，不再追加阶段确认。阶段问题待选时，只有施工图预算、投资估算、设计概算与工程竣工结算四个选项"
        "被视为阶段选择；其他文本必须退出或保留阶段待选并进入正常条件修改或问答流程，不得静默按施工图预算处理。"
        "通宝形象在准备与思考阶段保持轻微上下浮动、左右停顿的动画，节奏需平稳，不出现明显跳动或闪烁；"
        "在系统设置减少动态效果（prefers-reduced-motion）时，动画需降级为静态展示，仅保留状态文案与进度变化，"
        "避免影响可访问性。",
    )
    replace_paragraph_contains(
        doc,
        "历史价倒推规则：当目标期缺少可用历史价格时",
        "历史价倒推规则：当目标期缺少可用历史价格时，允许从最新一期有效价格开始，按“上一期价格 = 本期价格 ÷ "
        "本期系数”逐期向前倒推。每次倒推使用已知期，即被除数期对应的系数；待推算目标期的系数不参与该步倒推，"
        "不得连续复用同一个最新系数。",
    )
    insert_paragraph_after(
        find_paragraph(doc, "features 包含关系按结构化特征集合处理"),
        "映射表字段 cid、subcid、stdname 与 features 统一按字符串处理，入库与比较前去除首尾空格，"
        "空值统一归一为空字符串；不得因 32 与“32”、空值与 null 等类型差异产生不同匹配结果。"
        "映射表版本变更时，以新版本 ID 与生效时间记录内部日志。",
    )
    replace_paragraph_contains(
        doc,
        "结果卡片标题为“已识别为标准材料或设备，并完成历史询价匹配”",
        "结果卡片标题统一为“已完成标准材料识别与报价匹配”，不因使用历史询价或市场价兜底而改变用户可见标题。"
        "副标题按实际候选数量展示：满 3 条时展示“通宝给您推荐 3 条最优供应商报价”，"
        "不足 3 条时展示“通宝给您展示全部 N 条候选供应商报价”。",
    )
    replace_paragraph_contains(
        doc,
        "结果区展示建议价区间与含税含运口径",
        "结果区展示建议价区间与含税含运口径，并在报价表中按行展示材料名称、规格型号、数量、单位、品牌、"
        "建议价、供应商信息与回复备注。候选结果不足 3 条时不得提示用户重新询价，需直接展示全部候选记录；"
        "建议价之间的差异仅通过价格数值与排序体现，不额外标注价格基础来源。",
    )
    result_display_anchor = find_paragraph(doc, "结果区展示建议价区间与含税含运口径")
    price_range_paragraph = insert_paragraph_after(
        result_display_anchor,
        "建议价区间取结果卡中全部展示候选建议价的最小值至最大值；只有 1 条候选时上下限相同。"
        "区间与逐条金额统一保留 2 位小数，采用四舍五入；计算过程使用全精度，仅在最终展示时舍入。"
        "金额为空、小于等于 0 或无法转为有限数值时，该候选不得进入结果卡并转入无数据兜底；"
        "极大值需标记异常并阻止导出。",
    )
    insert_paragraph_after(
        price_range_paragraph,
        "查价条件完整度按 material、spec、quantity、unit、region、targetTime 六项槽位判定，每项只计一次。"
        "少于 5 项时展示“由于查价条件不足，建议价区间可能存在偏差”，但不得阻断查询；"
        "相同条件词重复出现不得重复计数。",
    )
    replace_paragraph_contains(
        doc,
        "项目阶段通过多轮对话采集：首次会话中",
        "项目阶段允许用户在首轮输入中主动填写；若已明确填写，系统直接使用该阶段，结果承接后不再重复询问。"
        "未填写时，通宝在结果承接后通过多轮对话询问，选项为施工图预算、投资估算、设计概算与工程竣工结算；"
        "用户跳过时默认施工图预算。用户后续明确选择其他阶段时，保留 requestId，原地重算同一结果卡，"
        "不扣减次数。后端按“价格基础 × 时间系数 × 项目阶段系数”计算建议价，项目阶段系数仅用于后端计算，"
        "对话页与结果区不展示系数、换算公式与提示语。系数表如下：",
    )
    replace_paragraph_contains(
        doc,
        "重新查询图标悬停时展示说明文案",
        "重新查询图标悬停时展示说明文案“点击将重新生成询价结果，并消耗 1 条询价使用次数”，"
        "点击后按当前询价条件重新生成结果并扣减 1 次使用次数。重新查询只要求重新生成结果，"
        "不承诺必须更换候选供应商；用户可见文案不得出现“保证更换候选”或“第三方模型重新筛选”等表述。"
        "所有反馈动作均需写入日志事件，见 3.6。",
    )
    replace_paragraph_contains(
        doc,
        "缺少关键规格时先走规格补充流程",
        "缺少关键规格时先走规格补充流程，用户确认后再进入查询；用户放弃补充时可选择转人工询价。"
        "人工询价不再使用 demo 内建表单，结果区、无数据、缺少规格与联系方式不可用等场景统一跳转 "
        "https://xunjia.zjtcn.com/release.html，由造价通页面承接后续登录、联系方式和询价要求填写；"
        "demo 不收集联系人、电话，不生成工单、受理状态或回流状态，本 PRD 也不承诺外链后的处理时效。",
    )
    replace_paragraph_contains(
        doc,
        "点击“导出盖章报价文件”进入盖章报价预览页",
        "结果区按钮统一为“导出盖章报价文件”，点击后进入同页侧边盖章报价预览。预览页按询价部模板呈现："
        "页面顶部展示造价通清晰标识、官网地址与防伪水印，主体为供应商报价表，表格包含材料名称、规格型号、"
        "数量、单位、品牌、价格、供应商信息与回复备注。报价表在加盖询价专用章前显示“未盖章”，"
        "未盖章时不得下载；盖章后提供“撤销盖章”，撤销后恢复未盖章状态。",
    )
    replace_paragraph_contains(
        doc,
        "预览页提供“显示联系方式”开关",
        "预览页提供“显示联系方式”开关，默认关闭，仅控制预览与导出中的联系方式字段。"
        "下载按钮统一为“下载PDF报价文件”，用户完成人工盖章确认后可生成并下载真实 PDF。"
        "文件名规则为“项目或区域_询价盖章报价_YYYYMMDDHHmmss.pdf”，由系统生成，不支持自定义覆盖。"
        "每次生成和下载需记录文件编号、requestId、生成时间、操作者、联系人与联系方式是否显示、下载时间与结果。"
        "页面不得展示“盖章生成方式待领导批示”等内部待定文案。",
    )
    insert_paragraph_after(
        find_paragraph(doc, "下载按钮统一为“下载PDF报价文件”"),
        "盖章报价预览以同页侧边预览为唯一验收实现。独立页面 quote-preview.html 仅保留为历史演示参考，"
        "不属于本期产品页面，也不作为 A18 的验收对象。",
    )
    replace_paragraph_contains(
        doc,
        "历史查价记录列表展示查询时间",
        "历史记录按“历史查询记录”展示，列表默认按查询时间倒序，最新记录在前；每页 5 条并展示总条数，"
        "右上角提供“新建对话”。列表展示查询时间、查价条件与操作三列，不展示推荐结果列。"
        "查询时间记录本次查询的提交时间，查价条件展示当轮合并后用于检索的材料、规格、数量、地区与目标时间。"
        "有正常询价结果的记录显示“查看结果”，无数据或缺少规格的记录显示“演示对话”；两个入口均回到对应会话。",
    )
    replace_paragraph_contains(
        doc,
        "会话条目右侧提供更多操作入口",
        "历史记录支持单条记录行内删除与会话删除两种粒度。单条删除仅移除该条历史查询记录并写入"
        "ai_inquiry_history_record_delete；会话删除需二次确认，确认框标题为“确定删除对话？”，"
        "确认后级联删除该会话下全部历史查询记录、会话状态与结果引用，并写入 ai_inquiry_conversation_delete。"
        "用户侧历史记录默认保留近一年，按查询时间每日清理超期数据；行为日志与审计记录不随用户删除动作删除，"
        "按独立留存策略保留。删除成功提示“已删除”，且不影响其他会话、历史统计或其他用户数据。",
    )
    replace_paragraph_contains(
        doc,
        "对话页左侧展示今日额度",
        "对话页左侧统一展示“今日额度：AI 询价 X 条 / AI 查价 50 条”，X 为当前剩余次数，不展示赠送剩余；"
        "每日额度按北京时间 00:00 重置。AI 查价额度由既有权益系统提供，本 PRD 只展示，不改造其查询主流程。"
        "查询、修改检索条件后重新提交、重新查询各扣减 1 次；查看历史、人工询价、信息型追问与项目阶段变更不扣次。"
        "次数不足时阻止继续发起查询，提示“查询次数已用完，可转人工询价”，并保留转人工询价入口。",
    )
    replace_paragraph_contains(
        doc,
        "非材料询价输入、规格补充确认过程与查看历史记录不扣减次数",
        "非材料询价输入、规格补充确认、纯信息型追问、项目阶段确认或变更、查看历史记录不扣减次数；"
        "重新查询按 1 次计。每次扣次成功后，在输入框下方展示“本次查询已消耗 1 次查询次数”，"
        "持续 3 秒后淡出，不再重复弹出 toast；失败或兜底拦截不展示该提示。",
    )
    replace_paragraph_contains(
        doc,
        "询价上下文作为后台处理能力实现",
        "询价上下文作为后台处理能力实现，不在 demo 对话页、结果卡或输入区新增上下文摘要、记忆抽屉或“上下文”按钮。"
        "后台保存 material、spec、quantity、unit、region、targetTime、projectPhase、brand、taxIncluded、"
        "freightIncluded 与 resultRef 的槽位、来源和确认状态，用于多轮修改与请求回放。"
        "前端不得保留无法触达的上下文抽屉、记忆设置或相关死代码。",
    )
    replace_paragraph_contains(
        doc,
        "系统每轮在意图拆解后执行状态合并",
        "系统每轮在意图拆解后执行状态合并：用户新输入的条件覆盖同名字段；用户仅补充缺失字段时保留其他已确认条件；"
        "模型推断出的字段先标记为 inferred，涉及规格、地区、目标时间等关键字段时需追问。"
        "“第二个供应商”“这个价格”等指代只允许解析当前最新有效结果中的具体记录 ID；条件变更、重新查询或结果失效后，"
        "旧 resultRef 与 supplierId 不得复用。存在多结果、多候选或指代不明时，先追问澄清，不猜测对象。",
    )
    replace_paragraph_contains(
        doc,
        "用户完成一轮询价后，后续输入只要包含材料",
        "用户完成一轮询价后，仅当本轮输入存在明确修改意图，且材料、规格、数量、单位、地区、目标时间、品牌、"
        "含税或含运槽位实际发生变化时，才按“上一轮已确认条件 + 本轮明确覆盖字段”合并并形成新的有效询价条件。"
        "“数量为什么是 4600”“这个价格含税吗”“第二个供应商是谁”等信息型追问不改变槽位、不扣次、不重新检索。"
        "仅修改地区、数量或时间等条件也属于有效询价，不得被“非材料询价”规则拦截；未在本轮提及的字段继续沿用"
        "上一轮已确认值。",
    )
    replace_paragraph_contains(
        doc,
        "材料发生变化时，新一轮必须使用本轮提供的规格",
        "材料发生变化时，新一轮必须使用本轮提供的规格；若本轮只更换材料而未提供规格，旧材料规格不得沿用，"
        "系统按缺少关键规格处理并引导用户补充。地区、目标时间、数量、单位、品牌、含税或含运条件发生变化时，"
        "新材料与规格沿用上一轮确认值；品牌、含税与含运未在本轮提及时沿用当前会话已确认值。",
    )
    replace_paragraph_contains(
        doc,
        "上下文失效规则：材料、规格、数量、地区或目标时间发生变化时",
        "上下文失效规则：材料、规格、数量、单位、地区、目标时间、品牌、含税或含运条件发生变化时，"
        "旧候选、建议价、供应商结果与导出引用全部失效，系统重新检索并生成新的 requestId；"
        "合并后的条件形成一次有效二次询价并扣减 1 次使用次数。项目阶段在同一会话内沿用；"
        "用户明确选择其他阶段时保留 requestId，只重算建议价与结果版本，不重新检索、不扣次、不追加结果卡。"
        "上下文变更原因与失效范围写入内部日志，不新增用户可见的上下文面板。",
    )
    replace_paragraph_contains(
        doc,
        "上下文窗口采用有界组装",
        "上下文窗口采用有界组装：优先保留当前结构化状态、输出约束与本轮询价结果；其次保留最近 3 轮对话；"
        "更早历史由规则摘要器压缩为“已确认条件、变更记录、未解决问题、结果引用”摘要。"
        "摘要生成失败时回退为“当前结构化状态 + 最近 3 轮原话”，不得丢失已确认槽位或价格计算规则。"
        "当上下文超出预算时，先删除无关原话和重复结果，不得删除当前确认字段或价格计算规则。",
    )
    replace_paragraph_contains(
        doc,
        "长期记忆仅作为后台能力保存低风险",
        "长期记忆仅作为后台能力保存低风险、稳定且可解释的用户偏好，例如默认报价地区、常用单位与材料别名。"
        "价格、供应商、联系方式、项目报价结果以及模型推测出的用户意图不得写入长期记忆。"
        "每次询价的价格与供应商信息必须重新检索，不得从长期记忆中复用。",
    )
    replace_paragraph_contains(
        doc,
        "长期记忆写入必须显式发生",
        "本期不开放长期记忆的查看、修改、清除或对话写入入口，仅保留后台读取、审计与人工配置能力。"
        "单次询价输入不得自动写入长期记忆；后续版本上线用户入口前，需补充授权、删除、跨设备同步与审计验收。",
    )
    replace_paragraph_contains(
        doc,
        "发生冲突时，本轮明确输入始终优先",
        "发生冲突时，本轮明确输入始终优先。例如后台默认地区为广东广州，但用户本次输入广东深圳，"
        "系统按广东深圳检索，并在上下文中标记来源为用户输入；用户未填写地区时，系统才沿用后台偏好。"
        "记忆来源写入内部日志，不在 demo 结果区增加“记忆”标签。",
    )
    replace_paragraph_contains(
        doc,
        "本节依据“ai查价模型切换”项目的",
        "本节依据“ai查价模型切换”项目的《造价通AI查价双模式需求规格说明书（V1.1）》、"
        "《AI查价双模式设计方案》及交互原型整理，作为 AI 查价跨项目参考文案，"
        "不属于本期 AI 询价功能、Demo 比对或发布验收范围。C06、C08、C13 涉及的计时、等待确认、"
        "取消与完成通知按 AI 查价项目自身交互稿实现和验收，本 PRD 不重复承诺其控件能力。"
        "AI 查价的快速思考与深度思考独立于 AI 询价流程，不得直接复用 AI 询价的 T01-T32 文案。",
    )
    replace_paragraph_contains(
        doc,
        "AI 查价的全部通宝气泡、模式提示、错误提示",
        "AI 询价仅使用“快速思考”这一处理方式；AI 查价可使用“快速思考”“深度思考”“线上快速”“本地深度”等业务名称。"
        "用户可见文案不得出现快速模型、深度模型、第三方模型、底层模型名称、模型版本、供应商、接口地址或系统提示词。"
        "重新筛选只可表述为“重新筛选候选报价”，不得向用户暴露模型或后端调用细节。",
    )
    replace_paragraph_contains(
        doc,
        "长期记忆仅保存默认地区、常用单位与材料别名等低风险偏好",
        "长期记忆仅保存默认地区、常用单位与材料别名等低风险偏好，不得保存联系方式、价格或供应商事实。"
        "本期不提供用户查看、修改与清除入口；后台访问需权限控制并记录审计日志。",
    )
    replace_paragraph_contains(
        doc,
        "查询、重新查询、次数扣减、导出与转人工询价需可追溯",
        "查询、重新查询、次数扣减、PDF 生成与下载、历史单条删除、会话删除、转人工询价与价格来源提问需可追溯，"
        "能够按 conversationId、requestId 还原用户操作顺序。",
    )
    replace_paragraph_contains(
        doc,
        "上下文更新、长期记忆保存与清除、记忆沿用与结果兜底动作",
        "上下文更新、长期记忆保存与清除、记忆沿用与结果兜底动作需按 conversationId、requestId 串联，"
        "可回放每次结果使用的状态版本与结果引用。审计原文与行为日志分开存储：审计原文加密保存 180 天，"
        "仅合规审计角色可访问；行为日志只记录结构化字段，不写入原始联系方式、完整对话原文或其他敏感信息。",
    )
    replace_paragraph_contains(
        doc,
        "单次询价从提交到结果展示的等待时间",
        "提交后 1 秒内必须出现首个可见反馈，准备状态启动时间 P95 不高于 2 秒；"
        "快速思考从提交到结果展示 P95 不高于 10 秒。",
    )
    replace_paragraph_contains(
        doc,
        "结果区优先渲染建议价与报价行",
        "快速思考超过 20 秒时展示超时提示，并提供重试或转人工入口；"
        "用户取消后停止前端等待并记录取消事件。",
    )
    replace_paragraph_contains(
        doc,
        "思考动画与结果渲染需避免明显卡顿",
        "PDF 生成与下载 P95 不高于 5 秒，历史记录首屏加载 P95 不高于 2 秒；"
        "以上操作不得阻塞对话主流程，页面滚动与结果渲染不出现高频重排。",
    )
    replace_paragraph_contains(
        doc,
        "上下文与记忆需纳入 AI 询价发布门禁",
        "上下文与记忆需纳入 AI 询价发布门禁。使用固定冻结评测集，样本量不少于 200 条，"
        "覆盖材料别名、缺失规格、条件修改、项目阶段复用、长期记忆冲突、无数据、联系方式追问、"
        "价格来源提问、品牌/税费/运费修改与候选供应商刷新等场景。评测集、标注规范、样本版本与失败判定"
        "由产品、数据与测试共同维护；样本标签采用双人标注，不一致项由产品负责人裁决。"
        "每次发布只使用冻结集计算门禁指标，固定输入、模型输出、检索结果、状态快照、系数版本与最终报价均需可回放。",
    )
    insert_paragraph_after(
        find_paragraph(doc, "上下文与记忆需纳入 AI 询价发布门禁"),
        "模型、Prompt、检索策略、上下文模板或记忆规则变更后，必须在同一冻结集上重新执行全量回放；"
        "不得使用线上点踩、重新查询或临时抽样替代固定评测集。",
    )


def apply_table_changes(doc: Document) -> None:
    tables = doc.tables

    set_table_row(
        tables[1],
        13,
        [
            "询价上下文",
            "当前会话中已确认的材料、规格、数量、单位、地区、目标时间、项目阶段、品牌、含税与含运条件，"
            "以及本轮结果引用；是模型每轮处理的事实源。",
        ],
    )
    set_table_row(
        tables[1],
        14,
        [
            "长期记忆",
            "后台保存的低风险偏好，例如默认报价地区、常用单位与材料别名；不等同于价格事实。"
            "本期不向用户展示查看、修改或清除入口。",
        ],
    )
    for values in [
        [
            "历史记录",
            "面向用户的历史查询列表总称；列表中的一条记录称为“历史查询记录”，AI 询价产生的业务结果沿用“历史询价记录”。",
        ],
        [
            "建议价区间",
            "当前结果卡中全部展示候选建议价的最小值至最大值；只有 1 条候选时上下限相同。",
        ],
        [
            "今日额度",
            "对话页展示的当日可用查询次数，统一按“今日额度”表述，不使用未定义的赠送剩余口径。",
        ],
        [
            "快速思考/深度思考",
            "用户可见的思考模式名称；不得改写为快速模型、深度模型、第三方模型或底层模型名称。",
        ],
    ]:
        append_table_row(tables[1], values)

    append_table_row(
        tables[2],
        [
            "造价通 AI 询价需求评审清单（V1.3，2026-10-19）",
            "本轮 P1/P2 与 P0-06 至 P0-13 修订的问题来源和处置依据。",
        ],
    )

    set_table_row(
        tables[3],
        1,
        [
            "本期解决",
            "首页新增“通宝 AI”入口；对话式询价输入与非材料询价拦截；历史询价检索与推荐前 3 条；"
            "系数匹配与建议价计算；询价建议结果展示；复制、点赞、点踩、重新查询等结果操作；"
            "供应商联系方式与价格来源话术；盖章报价预览与 PDF 导出；历史记录、单条删除与会话删除；"
            "结构化上下文与后台偏好；使用次数与日志埋点。",
        ],
    )
    set_table_row(
        tables[3],
        2,
        [
            "本期不解决",
            "不覆盖非标图纸类询价与图纸识别；不提供批量询价；不改造 AI 查价主流程，"
            "3.5.2 仅作为跨项目参考文案；不开放长期记忆用户入口；不开放语音输入；"
            "不对用户开放系数管理后台的操作权限；不对外展示价格来源与内部计算过程。",
        ],
    )

    set_table_row(
        tables[5],
        3,
        [
            "对话上下文",
            "最近 3 轮原话、用户修改动作、对上一轮结果的指代",
            "理解“第二个供应商”“改成 DN110”等指代",
            "短期保留；更早历史压缩为状态变更摘要，摘要失败时仅保留当前状态与最近 3 轮",
        ],
    )

    set_table_row(
        tables[6],
        5,
        [
            "项目阶段",
            "施工图预算、投资估算、设计概算、工程竣工结算",
            "属于会话级字段；同一会话修改条件后沿用。用户明确修改时保留 requestId，只重算建议价，不扣次、不追加结果卡；"
            "兜底结果场景可标记为 not_required",
        ],
    )
    set_table_row(
        tables[6],
        6,
        [
            "结果引用",
            "本轮 requestId、候选记录 ID 与建议价版本",
            "换材料、规格、数量、单位、地区、目标时间、品牌、含税或含运条件后，旧候选、建议价与导出引用失效；"
            "旧 supplierId/recordId 不得复用",
        ],
    )
    insert_table_row_before_last(
        tables[6],
        [
            "品牌",
            "用户指定或识别到的品牌要求",
            "未指定时不作为硬过滤条件；用户修改品牌后旧候选与建议价失效",
        ],
    )
    insert_table_row_before_last(
        tables[6],
        [
            "含税条件",
            "含税或不含税口径",
            "按本轮明确条件覆盖；未提及时沿用当前会话已确认值，修改后旧结果失效",
        ],
    )
    insert_table_row_before_last(
        tables[6],
        [
            "含运条件",
            "含运费或不含运费口径",
            "按本轮明确条件覆盖；未提及时沿用当前会话已确认值，修改后旧结果失效",
        ],
    )

    set_table_row(
        tables[7],
        4,
        [
            "询价部与运营",
            "提供盖章报价模板与报价字段口径，通过造价通人工询价页面承接线下人工询价。",
        ],
    )

    feature_rows = {
        1: ["入口与跳转", "通宝 AI 入口", "在造价通首页提供“通宝 AI”入口（功能名 AI 询价），点击进入对话式询价页；对话页与盖章报价预览页右上角提供“造价通首页”跳转。", "P0"],
        2: ["输入与识别", "自然语言输入", "支持输入材料名称、规格型号、数量、单位、报价城市、目标时间、项目阶段、品牌、含税与含运条件，支持补充修改。", "P0"],
        8: ["上下文", "结构化询价状态", "将 material/spec/quantity/unit/region/targetTime/projectPhase/brand/taxIncluded/freightIncluded 与 resultRef 保存为带状态、来源和置信度的会话字段；条件变化时使旧候选与旧建议价失效。", "P0"],
        9: ["上下文", "条件覆盖与二次询价", "仅在本轮有明确修改意图且检索槽位实际变化时覆盖上一轮同名字段，未提及字段沿用；合并后生成新 requestId、重新检索并按 1 次查询扣次。纯信息型追问不触发。", "P0"],
        10: ["上下文", "后台上下文状态", "上下文用于状态合并、结果失效与日志回放，不在 demo 对话页、结果卡或输入区新增展示入口，也不保留不可达的抽屉或设置代码。", "P0"],
        11: ["记忆", "默认地区偏好", "仅保留后台偏好、审计与人工配置能力；本期不展示用户入口，不支持单次询价自动写入。", "P1"],
        14: ["输入与识别", "项目阶段采集", "项目阶段可在首轮输入中主动填写；未填写时在正常结果承接后询问。阶段选择后原地刷新同一张结果卡，不扣次。", "P0"],
        29: ["导出", "盖章报价预览", "按询价部模板生成同页侧边预览，包含造价通清晰标识、官网地址、防伪水印、供应商报价表、盖章与撤销盖章；显示联系方式默认关闭。", "P0"],
        30: ["导出", "下载 PDF 报价文件", "人工盖章确认后生成真实 PDF，按钮为“下载PDF报价文件”；文件名按“项目或区域_询价盖章报价_YYYYMMDDHHmmss.pdf”生成，并记录生成与下载日志。", "P0"],
        31: ["记录", "历史记录", "历史查询记录按查询时间倒序，每页 5 条并展示总数；展示查询时间、查价条件与操作，不展示推荐结果列。", "P0"],
        32: ["记录", "历史记录删除", "支持单条记录删除与会话级删除；会话删除需二次确认并级联清理，两种操作分别记录日志。", "P0"],
        33: ["次数", "今日额度与扣次", "展示“今日额度：AI 询价 X 条 / AI 查价 50 条”；查询、条件变更后重新检索、重新查询各扣 1 次，阶段变更与信息型追问不扣次。", "P0"],
        34: ["通宝文案", "AI 查价双模式话术（参考）", "3.5.2 的 C01-C15 仅作为 AI 查价跨项目参考，不属于本期 AI 询价功能、Demo 比对或发布验收范围。", "参考"],
        35: ["日志", "行为日志", "记录查询、重新查询、扣次、导出、价格来源提问、规格补充、非材料拦截、联系方式追问、历史删除等事件，字段见 3.6。", "P0"],
    }
    for row_index, values in feature_rows.items():
        set_table_row(tables[8], row_index, values)

    set_table_row(
        tables[14],
        1,
        [
            "用户偏好",
            "默认报价地区、常用单位、材料别名、常用展示偏好",
            "价格、供应商、联系人、电话、未经确认的规格",
            "仅当本次输入缺失对应字段时作为内部默认值，内部来源标记为 memory；本期不展示用户入口",
        ],
    )

    set_table_row(
        tables[15],
        7,
        [
            "T07",
            "思考步骤推进",
            "思考步骤卡片",
            "正在理解材料、规格与报价地区... / 正在匹配历史询价与市场价... / 正在核对字段、供应商与报价记录...",
            "按顺序逐条切换，仅保留 3 个步骤；AI 询价不提供深度思考入口。",
        ],
    )
    set_table_row(
        tables[15],
        16,
        [
            "T16",
            "询价结果生成",
            "通宝气泡",
            "已匹配到 {N} 条候选报价；点击刷新会重新筛选候选报价～",
            "{N}替换为实际候选数量；不足 3 条时不提示重新询价。不得出现“第三方模型”“更换候选供应商”等承诺。",
        ],
    )
    set_table_row(
        tables[15],
        30,
        [
            "T30",
            "语音输入入口",
            "不展示可点击入口",
            "语音输入未纳入本期范围，请使用文字描述材料、规格和地区。",
            "本期正式页面隐藏麦克风入口；旧版截图中的语音按钮仅作历史素材，不纳入验收。",
        ],
    )
    for values in [
        [
            "T31",
            "项目阶段待选时输入非阶段文本",
            "通宝气泡",
            "未选择项目阶段，我先按新的查价条件继续识别～",
            "不得将任意输入静默当作默认阶段；应退出或保留阶段待选并进入正常条件修改或问答流程。",
        ],
        [
            "T32",
            "扣次成功",
            "输入框下方提示",
            "本次查询已消耗 1 次查询次数",
            "展示 3 秒后淡出；失败、拦截与不扣次场景不展示，不重复弹 toast。",
        ],
        [
            "T33",
            "查价条件不足",
            "结果卡提示",
            "由于查价条件不足，建议价区间可能存在偏差～",
            "material、spec、quantity、unit、region、targetTime 六项槽位少于 5 项时展示；相同条件词不重复计数。",
        ],
    ]:
        append_table_row(tables[15], values)

    set_table_row(
        tables[17],
        1,
        ["结果标题", "结果卡片标题", "已完成标准材料识别与报价匹配", "固定文案，不因历史价或市场价兜底改变。"],
    )

    set_table_row(
        tables[18],
        5,
        [
            "ai_inquiry_conversation_delete",
            "用户确认删除当前会话",
            "conversationTitle、query、cascadeCount、wasActive、deletedAt",
        ],
    )
    for values in [
        [
            "ai_inquiry_query",
            "用户提交首轮有效询价",
            "conversationId、requestId、slotsSnapshot、stateVersion、retrievalVersion、coefficientVersion",
        ],
        [
            "ai_inquiry_requery",
            "用户点击重新查询或修改条件后重新检索",
            "conversationId、requestId、previousRequestId、changedFields、deducted",
        ],
        [
            "ai_inquiry_quota_deduct",
            "实际扣减查询次数成功",
            "quotaType、beforeCount、afterCount、reason、requestId、deductedAt",
        ],
        [
            "ai_inquiry_export",
            "用户生成或下载 PDF 报价文件",
            "requestId、fileId、fileName、generationStatus、downloadStatus、sealed、showContact",
        ],
        [
            "ai_inquiry_price_source_ask",
            "用户询问价格来源或建议价依据",
            "conversationId、requestId、query、responseMode",
        ],
        [
            "ai_inquiry_spec_supplement",
            "用户提交或关闭规格补充",
            "requestId、selectedFields、source、submitted、deducted",
        ],
        [
            "ai_inquiry_non_material_intercept",
            "非材料询价输入被拦截",
            "query、detectionReason、deducted=false",
        ],
        [
            "ai_inquiry_contact_ask",
            "用户追问供应商联系方式",
            "requestId、supplierId、available、fallbackUsed",
        ],
        [
            "ai_inquiry_history_record_delete",
            "用户删除单条历史查询记录",
            "recordId、conversationId、cascadeScope、deletedAt",
        ],
    ]:
        append_table_row(tables[18], values)

    for values in [
        [
            "用户输入训练与提示",
            "本期不使用用户输入训练模型；正式页面不得展示“输入将被收集用于模型训练”的提示条。"
            "如未来启用训练用途，需单独完成授权、撤回、脱敏、留存与第三方共享评审。",
        ],
        [
            "人工询价外链边界",
            "外链仅跳转造价通人工询价页面，不生成工单、不回传受理状态；"
            "产品文案不得承诺外链后的处理时效。",
        ],
        [
            "导出口径与内部文案",
            "PDF 生成、盖章与下载记录必须可追溯；内部演示按钮、规则抽屉、待领导批示等信息不得进入对外演示版。",
        ],
    ]:
        append_table_row(tables[19], values)

    set_table_row(
        tables[20],
        1,
        [
            "关键槽位准确率",
            "冻结评测集（不少于 200 条）中材料、规格、数量、单位、地区、目标时间、项目阶段识别正确数 / 总样本数",
            "正式发布不低于 95%",
        ],
    )
    set_table_row(
        tables[20],
        2,
        [
            "上下文污染率",
            "冻结评测集中修改条件后仍引用旧候选、旧价格或旧供应商的样本数 / 总样本数",
            "0",
        ],
    )

    acceptance_rows = {
        1: ["A1", "首页入口", "造价通首页可进入“通宝 AI”对话式询价页，入口与既有 AI 查价入口并存且可区分。"],
        3: ["A3", "输入识别", "可识别材料名称、规格、数量、单位、报价城市、目标时间、品牌、含税与含运条件，并支持补充修改；省份不作为检索条件参与过滤。"],
        6: ["A6", "思考与阶段顺序", "首次提交后展示通宝思考状态；结果承接后再展示项目阶段选择，阶段变更原地刷新同一张结果卡。阶段待选时非阶段文本不得被静默按默认阶段处理。"],
        11: ["A11", "结果展示", "结果卡片展示建议价区间、材料与规格、数量单位、品牌、建议价、供应商信息与回复备注；区间取展示候选最小值至最大值，金额统一 2 位小数四舍五入；六项关键条件少于 5 项时展示条件不足提示。"],
        12: ["A12", "项目阶段采集与系数", "项目阶段可在首轮填写；未填写时在结果承接后询问。阶段变更保留 requestId、只重算同一结果卡且不扣次；后端按对应阶段系数计算建议价，页面不展示系数、公式与提示语。"],
        15: ["A15", "重新查询", "重新查询前提示消耗 1 次使用次数，点击后按当前条件重新生成结果并扣减 1 次；不承诺必须更换候选供应商。"],
        18: ["A18", "盖章报价与 PDF", "同页侧边预览展示标识、官网地址、防伪水印和报价表；联系方式默认关闭；必须人工盖章后才能撤销或下载 PDF。PDF 文件名符合规则，并记录文件编号、生成、下载与操作者。"],
        19: ["A19", "历史记录与删除", "历史查询记录按查询时间倒序，每页 5 条并展示总数；有结果记录显示“查看结果”，无数据或缺规格记录显示“演示对话”。支持单条删除和会话级删除，默认保留近一年，日志独立留存。"],
        20: ["A20", "使用次数", "展示“今日额度：AI 询价 X 条 / AI 查价 50 条”；查询、条件变更后重新检索、重新查询各扣 1 次；阶段变更、信息型追问、查看历史与人工询价不扣次；扣次成功后展示 3 秒提示。"],
        21: ["A21", "日志埋点", "查询、重新查询、扣次、PDF 生成与下载、价格来源提问、规格补充、非材料拦截、联系方式追问、历史单条删除与会话删除均写入日志，属性完整可追溯。"],
        27: ["A27", "上下文状态", "后台保存十类业务槽位与 resultRef 的值、状态、来源、置信度和更新时间；关键槽位 confidence 低于 0.8 时追问。demo 不展示上下文摘要、记忆抽屉或相关入口。"],
        28: ["A28", "上下文失效", "修改材料、规格、数量、单位、地区、目标时间、品牌、含税或含运条件后，旧候选、建议价、供应商与导出引用失效并重新生成；阶段变更走同卡重算且不扣次。"],
        29: ["A29", "长期记忆", "本期仅保留后台偏好、审计与人工配置能力，不展示查看、修改、清除或对话写入入口；单次询价输入不自动保存。"],
        31: ["A31", "上下文与记忆日志", "上下文更新、记忆沿用、审计访问与结果兜底均写入日志，可按 conversationId、requestId 还原；审计原文与行为日志分开存储。"],
        32: ["A32", "上下文窗口与回放", "每轮保存状态版本、prompt 版本、模型版本、检索版本与系数版本；上下文固定保留最近 3 轮，摘要失败回退为当前状态加最近 3 轮，不得丢失确认字段。"],
        33: ["A33", "转人工统一跳转", "结果区、无数据、缺少规格与联系方式不可用等场景均在新页面打开同一造价通人工询价地址；不弹出自建表单，不收集联系人、电话，不生成工单或受理状态。"],
        34: ["A34", "通宝思考与文案一致性", "准备、思考、结果承接与项目阶段确认按设计图顺序出现；AI 询价仅使用快速思考 3 个步骤，不展示深度入口；用户可见文案不出现快速模型、深度模型或第三方模型。"],
        36: ["A36", "用户可见文案与内部信息隔离", "AI 询价页面不展示快速模型、深度模型、第三方模型、内部规则、待领导批示或“用于模型训练”提示；规则抽屉仅内部版可用且公式为三因子口径。"],
    }
    for row_index, values in acceptance_rows.items():
        set_table_row(tables[21], row_index, values)


def insert_supporting_notes(doc: Document) -> None:
    diagrams_anchor = next(
        p for p in doc.paragraphs if p.text.startswith("演示包为静态页面，解压后直接打开")
    )
    insert_paragraph_after(
        diagrams_anchor,
        "配套交互说明图清单为：00_AI询价业务流程图、01_PC端_基于demo截图_AI询价交互说明、"
        "02_PC端_项目阶段多轮问答_交互说明、04_PC端_AI查价思考加载态_交互说明。"
        "03 为保留编号，当前未使用；后续新增说明图时不得重排既有编号。",
    )


def verify(doc: Document) -> None:
    required = [
        "准备、思考、结果承接、项目阶段确认",
        "建议价 = 价格基础 × 时间系数 × 项目阶段系数",
        "5.2 待确认事项",
        "今日额度：AI 询价 X 条 / AI 查价 50 条",
        "本次查询已消耗 1 次查询次数",
        "ai_inquiry_export",
    ]
    all_text = "\n".join(p.text for p in doc.paragraphs)
    all_text += "\n" + "\n".join(
        cell.text for table in doc.tables for row in table.rows for cell in row.cells
    )
    for term in required:
        if term not in all_text:
            raise ValueError(f"缺少预期修订内容：{term}")

    forbidden = [
        "建议价按价格基础 × 时间系数计算",
        "建议价等于价格基础乘时间系数",
        "只要包含材料、规格、数量、单位、地区、目标时间或项目阶段中的任一条件",
        "用户可查看、修改与清除",
        "下载演示文件",
        "取消反馈、通宝互动与会话删除均写入日志事件",
    ]
    for term in forbidden:
        if term in all_text:
            raise ValueError(f"仍存在旧口径：{term}")


def main() -> None:
    if not BACKUP_PATH.exists():
        shutil.copy2(DOC_PATH, BACKUP_PATH)
    doc = Document(DOC_PATH)
    apply_paragraph_changes(doc)
    apply_table_changes(doc)
    add_toc_entry(doc)
    insert_phase_pending_section(doc)
    insert_supporting_notes(doc)
    verify(doc)
    doc.save(DOC_PATH)
    print(DOC_PATH)


if __name__ == "__main__":
    main()
