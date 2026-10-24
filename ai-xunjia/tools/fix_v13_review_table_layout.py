"""Normalize widths for the two four-column review tables and the new heading."""

from __future__ import annotations

import copy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path("/Users/zouxiaoli/Documents/ChatGPT/ai询价")
DOC_PATH = ROOT / "造价通AI询价需求规格说明书（V1.3）.docx"
TABLE_WIDTH = 9412


def set_table_width(table, widths: list[int]) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    grid_columns = table._tbl.tblGrid.gridCol_lst
    if len(grid_columns) != len(widths):
        raise ValueError(f"表格列数异常：{len(grid_columns)}")
    for grid, width in zip(grid_columns, widths):
        grid.set(qn("w:w"), str(width))

    for row in table.rows:
        cells = row._tr.tc_lst
        if len(cells) != len(widths):
            raise ValueError(f"单元格数量异常：{len(cells)}")
        for cell, width in zip(cells, widths):
            tc_pr = cell.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.insert(0, tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")


def find_table(doc: Document, header: str, row_count: int):
    matches = [
        table
        for table in doc.tables
        if len(table.rows) == row_count
        and len(table.columns) == 4
        and table.rows[0].cells[0].text == header
    ]
    if len(matches) != 1:
        raise ValueError(f"四列表定位异常：{len(matches)}，header={header}, rows={row_count}")
    return matches[0]


def set_paragraph_text(paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run._element.getparent().remove(run._element)
    else:
        paragraph.add_run(text)


def replace_paragraph(doc: Document, old_startswith: str, new: str) -> None:
    matches = [p for p in doc.paragraphs if p.text.startswith(old_startswith)]
    if len(matches) != 1:
        raise ValueError(f"段落定位异常：{len(matches)}，{old_startswith}")
    set_paragraph_text(matches[0], new)


def set_cell_text(cell, text: str) -> None:
    set_paragraph_text(cell.paragraphs[0], text)
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)


def main() -> None:
    doc = Document(DOC_PATH)

    context_example = find_table(doc, "场景", 4)
    pending_questions = find_table(doc, "编号", 11)
    set_table_width(context_example, [1300, 2850, 2300, 2962])
    set_table_width(pending_questions, [700, 2300, 3900, 2512])

    headings = [
        p for p in doc.paragraphs if p.style.name == "Heading 2" and p.text == "5.1 验收标准"
    ]
    new_headings = [
        p for p in doc.paragraphs if p.style.name == "Heading 2" and p.text == "5.2 待确认事项"
    ]
    if len(headings) != 1 or len(new_headings) != 1:
        raise ValueError("5.1 或 5.2 标题定位异常")

    source = headings[0]
    target = new_headings[0]
    if target._p.pPr is not None:
        target._p.remove(target._p.pPr)
    target._p.insert(0, copy.deepcopy(source._p.pPr))
    if target.runs and source.runs:
        target_run = target.runs[0]._r
        if target_run.rPr is not None:
            target_run.remove(target_run.rPr)
        target_run.insert(0, copy.deepcopy(source.runs[0]._r.rPr))

    replace_paragraph(
        doc,
        "技术链路中的“历史价格 × 项目阶段系数”",
        "建议价计算以 3.4.5 的公式为唯一口径：建议价 = 价格基础 × 时间系数 × 项目阶段系数。"
        "技术链路、页面文案与内部说明均不得使用两因子简化公式，也不得省略项目阶段系数。",
    )
    set_cell_text(
        doc.tables[1].rows[6].cells[1],
        "在价格基础上按时间系数与项目阶段系数换算后的参考价格，"
        "公式为：价格基础 × 时间系数 × 项目阶段系数。",
    )
    set_cell_text(
        doc.tables[15].rows[1].cells[4],
        "进入新对话时展示，作为首次引导；项目阶段为可选填项，未填写时在结果承接后询问。",
    )
    set_cell_text(
        doc.tables[15].rows[2].cells[4],
        "不阻断输入，仅提示必填要素；项目阶段可提前填写，未填写时在结果承接后补充。",
    )

    doc.save(DOC_PATH)
    print(DOC_PATH)


if __name__ == "__main__":
    main()
