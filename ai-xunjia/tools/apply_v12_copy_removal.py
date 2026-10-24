"""Remove the copy entry of the 通宝提示语 after the 2026-10-16 review follow-up.

Scope:
  - 3.4.2 body text no longer mentions the copy entry beside the 通宝提示语.
  - 3.5 提示文案与状态反馈规范 drops the "提示语复制" row.
  - A39 acceptance item keeps 项目阶段 wording without the copy entry.
  - 评审记录 adds a follow-up row and 变更履历 records the adjustment.
"""

from __future__ import annotations

import shutil
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph


ROOT = Path("/Users/zouxiaoli/Documents/ChatGPT/ai询价")
DOC_PATH = ROOT / "造价通AI询价需求规格说明书（V1.2）.docx"
BACKUP_PATH = ROOT / ".qa/造价通AI询价需求规格说明书（V1.2）_复制入口删除前备份.docx"

OLD_BODY_SENTENCE = (
    "对话页通宝提示语同步提示需要提供项目阶段，并在提示语旁提供复制入口，"
    "点击后可复制询价条件模板“材料名称 + 规格型号 + 数量 + 单位 + 报价地区 + 目标时间 + 项目阶段”，"
    "供用户直接粘贴使用。"
)
NEW_BODY_SENTENCE = "对话页通宝提示语同步提示需要提供项目阶段，仅作要素引导，不提供复制入口。"

A39_TEXT = (
    "输入框占位文案与通宝提示语均包含“项目阶段”。"
    "通宝提示语旁不提供复制入口与询价条件模板复制能力；结果区的复制操作不受影响。"
)

CHANGE_LOG_TEXT = (
    "按 AI 询价第二次评审补充确认调整：删除通宝提示语旁的复制入口及其询价条件模板复制能力，"
    "输入框占位文案与通宝提示语仍保留“项目阶段”要素提示；"
    "结果区的复制、点赞、点踩与重新查询操作保持不变。"
)

REVIEW_ROW = (
    "AI 询价第二次评审补充确认（2026-10-16）",
    "删除通宝提示语旁复制入口与询价条件模板复制，保留输入框与提示语中的项目阶段要素；"
    "结果区复制、点赞、点踩与重新查询操作口径不变。",
)


def set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def drop_row(table, index: int) -> None:
    row = table.rows[index]
    row._element.getparent().remove(row._element)


def append_cell_paragraph(cell, text: str) -> None:
    template = cell.paragraphs[-1]
    new_p = deepcopy(template._p)
    cell._tc.append(new_p)
    set_paragraph_text(Paragraph(new_p, cell), text)


def append_table_row(table, values) -> None:
    new_row = deepcopy(table.rows[-1]._tr)
    table._tbl.append(new_row)
    row = table.rows[-1]
    for cell, value in zip(row.cells, values):
        set_paragraph_text(cell.paragraphs[0], value)


def find_row_by_first_cell(table, prefix: str):
    for row in table.rows:
        if row.cells[0].text.strip().startswith(prefix):
            return row
    raise RuntimeError(f"row not found: {prefix}")


def main() -> None:
    BACKUP_PATH.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(DOC_PATH, BACKUP_PATH)
    doc = Document(DOC_PATH)

    body_updated = False
    for paragraph in doc.paragraphs:
        if OLD_BODY_SENTENCE in paragraph.text:
            set_paragraph_text(
                paragraph, paragraph.text.replace(OLD_BODY_SENTENCE, NEW_BODY_SENTENCE)
            )
            body_updated = True
    if not body_updated:
        raise RuntimeError("3.4.2 copy sentence not found")
    print("3.4.2 body updated")

    text_table = doc.tables[12]
    for index, row in enumerate(text_table.rows):
        if row.cells[0].text.strip() == "提示语复制":
            drop_row(text_table, index)
            print("3.5 提示语复制 row removed")
            break
    else:
        raise RuntimeError("提示语复制 row not found")

    a39_row = find_row_by_first_cell(doc.tables[16], "A39")
    set_paragraph_text(a39_row.cells[2].paragraphs[0], A39_TEXT)
    print("A39 updated")

    review_table = doc.tables[18]
    append_table_row(review_table, REVIEW_ROW)
    print("review record row appended")

    for table in doc.tables[:2]:
        for row in table.rows:
            cells = row.cells
            if len(cells) >= 4 and cells[1].text.strip() == "V1.2":
                append_cell_paragraph(cells[2], CHANGE_LOG_TEXT)
                print("change log updated")
                break

    doc.save(DOC_PATH)
    print(f"Saved: {DOC_PATH}")
    print(f"Backup: {BACKUP_PATH}")


if __name__ == "__main__":
    main()
