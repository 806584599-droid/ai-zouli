"""Generate 造价通AI询价_免责协议（草案V1）.docx.

Styling follows tools/build_ai_xunjia_spec_doc.py: A4 page, 宋体 / Times New Roman,
H1 16pt, H2 14pt, H3 12pt, body 10.5pt with 2-character first-line indent,
1.25 line spacing, bordered tables with #D9EAF7 header fill.
"""

from __future__ import annotations

import os
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path("/Users/zouxiaoli/Documents/ChatGPT/ai询价")
QA_MODE = os.environ.get("DISCLAIMER_DOC_QA") == "1"
OUTPUT = (
    ROOT / ".qa/ai_xunjia_disclaimer_qa/造价通AI询价_免责协议（草案V1）.docx"
    if QA_MODE
    else ROOT / "造价通AI询价_免责协议（草案V1）.docx"
)

FONT_CN = "宋体"
FONT_EN = "Times New Roman"
BLACK = RGBColor(0x00, 0x00, 0x00)
CAPTION_GRAY = RGBColor(0x5A, 0x5A, 0x5A)
HEADER_FILL = "D9EAF7"

BODY_SIZE = 10.5
H1_SIZE = 16
H2_SIZE = 14
H3_SIZE = 12
CAPTION_SIZE = 9.5

doc = Document()


# ---------------------------------------------------------------- base styles
def set_run_font(run, size=BODY_SIZE, bold=False, color=BLACK):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT_EN
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    fonts.set(qn("w:ascii"), FONT_EN)
    fonts.set(qn("w:hAnsi"), FONT_EN)
    fonts.set(qn("w:eastAsia"), FONT_CN)


def set_spacing(paragraph, before=0, after=6, line=1.25):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def add_heading(text, level, page_break=False):
    size = {1: H1_SIZE, 2: H2_SIZE, 3: H3_SIZE}[level]
    before = {1: 12, 2: 8, 3: 6}[level]
    after = {1: 8, 2: 6, 3: 4}[level]
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    set_spacing(paragraph, before=before, after=after, line=1.25)
    if page_break:
        paragraph.paragraph_format.page_break_before = True
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=True)
    return paragraph


def add_body(text, indent=True, bold=False):
    paragraph = doc.add_paragraph()
    set_spacing(paragraph)
    if indent:
        paragraph.paragraph_format.first_line_indent = Pt(BODY_SIZE * 2)
    run = paragraph.add_run(text)
    set_run_font(run, bold=bold)
    return paragraph


def add_bullet(text):
    paragraph = doc.add_paragraph(style="List Bullet")
    set_spacing(paragraph, after=4)
    run = paragraph.add_run(text)
    set_run_font(run, bold=False)
    return paragraph


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = OxmlElement(f"w:{name}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tc_pr.append(mar)


def set_cell_vcenter(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    node = OxmlElement("w:vAlign")
    node.set(qn("w:val"), "center")
    tc_pr.append(node)


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), "000000")
        borders.append(node)
    tbl_pr.append(borders)


def style_cell(cell, text, bold=False, center=False, size=BODY_SIZE):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    set_spacing(paragraph, before=0, after=0, line=1.15)
    lines = str(text).split("\n")
    for index, line in enumerate(lines):
        if index:
            paragraph = cell.add_paragraph()
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
            )
            set_spacing(paragraph, before=0, after=0, line=1.15)
        run = paragraph.add_run(line)
        set_run_font(run, size=size, bold=bold)
    set_cell_margins(cell)
    set_cell_vcenter(cell)


CM_TO_TWIPS = 567


def apply_table_widths(table, widths):
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(int(sum(widths) * CM_TO_TWIPS)))
    tbl_w.set(qn("w:type"), "dxa")
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    grid = table._tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for column, width in zip(grid.findall(qn("w:gridCol")), widths):
            column.set(qn("w:w"), str(int(width * CM_TO_TWIPS)))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Cm(width)


def add_table(headers, rows, widths=None, sizes=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Normal Table"
    set_table_borders(table)
    header_cells = table.rows[0].cells
    header_row_pr = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    header_row_pr.append(repeat)
    for index, header in enumerate(headers):
        style_cell(header_cells[index], header, bold=True, center=True, size=9.5)
        shade_cell(header_cells[index], HEADER_FILL)
    for row in rows:
        table_row = table.add_row()
        cant_split = OxmlElement("w:cantSplit")
        table_row._tr.get_or_add_trPr().append(cant_split)
        cells = table_row.cells
        for index, value in enumerate(row):
            size = 9.5 if sizes is None else sizes[index]
            style_cell(cells[index], value, size=size)
    if widths:
        apply_table_widths(table, widths)
    spacer = doc.add_paragraph()
    set_spacing(spacer, before=0, after=6, line=1.0)
    return table


# ------------------------------------------------------------- page settings
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.left_margin = Cm(2.49)
section.right_margin = Cm(1.5)
section.top_margin = Cm(2)
section.bottom_margin = Cm(2.49)

normal = doc.styles["Normal"]
normal.font.name = FONT_EN
normal.font.size = Pt(BODY_SIZE)
normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
normal.paragraph_format.line_spacing = 1.25


def centered_line(text, size, bold, after=4, before=0, page_break=False):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_spacing(paragraph, before=before, after=after, line=1.15)
    if page_break:
        paragraph.paragraph_format.page_break_before = True
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return paragraph


# -------------------------------------------------------------------- 封面
for _ in range(4):
    centered_line("", BODY_SIZE, False, after=0)
centered_line("造价通 AI 询价", 22, True, after=6)
centered_line("免责协议与免责声明", 26, True, after=20)
for _ in range(3):
    centered_line("", BODY_SIZE, False, after=0)
centered_line("草案 V1.0", H2_SIZE, False, after=6)
centered_line("2026-10-16", H2_SIZE, False, after=6)
centered_line("面向：产品 / 法务 / 前端 / 客服", H2_SIZE, False, after=6)

add_heading("编制说明", 2, page_break=True)
add_body(
    "本文件为产品侧建议稿，用于并入《造价通用户协议》并落地为页面文案，"
    "所有条款需经法务审定后方可正式生效。"
)

# ------------------------------------------------------------------ 一、目标
add_heading("一、设计目标与适用边界", 1)
add_body(
    "覆盖范围：造价通 AI 询价服务，含 AI 查价、AI 询价、历史询价匹配、建议价与建议价区间生成、"
    "候选供应商报价推荐、供应商联系方式展示、盖章报价文件预览与导出。"
)
add_body("设计上围绕三个目标展开：")
add_bullet("讲清楚结果来源与性质：由造价通 AI 基于 building DS 模型自动生成，属于辅助参考信息。")
add_bullet(
    "写明白责任边界：用户使用结果数据作出的决策与后果由用户自行承担，造价通网站不对结果负责。"
)
add_bullet(
    "在不影响转化体验的前提下完成有效告知：首次使用弹窗勾选、结果页常驻提示、导出文件页脚三处留痕。"
)
add_body(
    "边界：免责条款只覆盖“AI 生成与推荐”相关内容，不豁免造价通自身产品的重大过失、数据安全义务"
    "与付费服务合同义务。"
)

# ------------------------------------------------------------------ 二、速览
add_heading("二、免责口径速览（一页版）", 1)
for index, item in enumerate(
    [
        "询价结果由 building DS 模型自动生成，属于人工智能生成内容，不是人工造价咨询成果。",
        "结果仅供参考，不构成要约、承诺、报价担保，也不构成合同附件或专业意见。",
        "不得作为投标报价、结算、审计、司法鉴定或对外报批的唯一依据。",
        "历史询价数据存在滞后、缺失与口径差异，市场价格随时波动，不保证与实际成交价一致。",
        "候选供应商、联系方式与回复备注来自第三方或历史记录，未经造价通核实，不构成资质与履约背书。",
        "用户与供应商之间的后续交易与纠纷，由双方自行处理，与造价通无关。",
        "用户使用结果数据产生的任何损失，造价通网站及其关联方不承担责任。",
        "用户使用本服务即视为已知悉并接受上述说明。",
    ],
    start=1,
):
    add_body(f"{index}. {item}", indent=True)

# ------------------------------------------------------ 三、用户协议新增章节
add_heading("三、用户协议新增章节（正式条款建议稿）", 1, page_break=True)
add_body(
    "建议作为《造价通用户协议》独立一章，标题为“AI 询价服务说明与免责条款”，"
    "条款编号按用户协议现有顺序顺延。"
)

add_heading("第 X 条 AI 询价服务说明与免责条款", 2)

add_heading("X.1 定义", 3)
add_body(
    "（1）“AI 询价服务”指造价通网站向用户提供的，基于人工智能模型自动完成材料设备价格检索、"
    "历史询价匹配、建议价生成、候选供应商推荐及报价文件生成的服务，包含 AI 查价与 AI 询价功能。"
)
add_body(
    "（2）“AI 生成内容”指 AI 询价服务运行过程中自动输出的全部内容，包括但不限于建议价、"
    "建议价区间、询价结论、候选供应商报价、供应商联系方式、回复备注、价格来源说明及导出的报价文件。"
)
add_body(
    "（3）“造价通 AI 询价模型”指造价通网站用于生成上述内容的模型与算法，"
    "现阶段基于 building DS 模型及相关数据与规则构建。"
)

add_heading("X.2 生成方式与内容性质", 3)
add_body(
    "（1）用户知悉并确认，AI 查价与 AI 询价的全部结果均由造价通 AI 询价模型自动生成，"
    "生成过程由算法完成，未经人工逐条审核。"
)
add_body(
    "（2）AI 生成内容属于人工智能生成内容，可能因模型能力边界、训练数据、检索条件、输入完整度等因素"
    "出现不准确、不完整、不适用或与实际情况不符的情形，也可能出现模型幻觉。"
)
add_body("（3）AI 生成内容仅代表系统基于现有数据计算得出的参考信息，不代表造价通网站的立场、观点或专业意见。")
add_body(
    "（4）模型版本、算法与规则可能因技术迭代调整，造价通网站无需就调整事先征得用户同意，"
    "但会以合理方式提示。"
)

add_heading("X.3 仅供参考，不构成任何承诺", 3)
add_body("（1）AI 生成内容仅供用户内部参考，不构成要约、要约邀请、承诺或价格担保，不构成任何合同或协议的组成部分。")
add_body("（2）AI 生成内容不构成工程造价、法律、财税、审计、司法鉴定等任何专业意见，不能替代人工造价咨询与专业判断。")
add_body(
    "（3）用户不得将 AI 生成内容作为投标报价、合同签订、工程结算、审计核对、司法鉴定、"
    "对外报批或公开发布的唯一依据或直接依据。"
)
add_body("（4）实际交易价格、供应商资质与履约能力，以用户与供应商另行签署的正式文件及实际履行为准。")

add_heading("X.4 数据来源、时效性与匹配范围", 3)
add_body(
    "（1）AI 生成内容依据造价通网站历史询价数据、公开市场信息及造价通数据库生成，"
    "不保证覆盖全部材料、品牌、规格型号、地区与时间范围。"
)
add_body(
    "（2）历史数据可能存在采集口径差异、时间滞后、字段缺失或与实际市场情况不一致的情形，"
    "市场价格亦会随供需、政策与原材料价格波动而变化，故 AI 生成内容不保证与实际市场价或实际成交价一致。"
)
add_body(
    "（3）当用户提供的材料名称、规格型号、数量、单位、地区、目标时间或项目阶段等信息不完整或不明确时，"
    "结果偏差风险显著提高，用户应自行核实关键参数后再使用。"
)
add_body(
    "（4）建议价系依据用户所选项目阶段对应系数对历史询价数据计算所得，不同项目阶段的建议价口径不同；"
    "页面未展示的内部计算参数不作为用户主张结果的依据，用户应以页面实际展示结果为准。"
)
add_body("（5）含税、含运费等价格口径以页面标注为准；页面未标注的部分，用户应自行向供应商确认。")

add_heading("X.5 第三方供应商信息", 3)
add_body(
    "（1）服务中展示的供应商名称、联系方式、报价及回复备注，来自第三方提供或历史询价记录，"
    "造价通网站未对其真实性、准确性、完整性和时效性进行核实。"
)
add_body("（2）展示第三方供应商不构成造价通网站对该供应商资质、信用、履约能力、商品质量或价格的担保、推荐或背书。")
add_body(
    "（3）用户与第三方供应商之间的询价沟通、交易、合同履行及由此产生的任何争议与损失，"
    "由用户与供应商自行协商解决并自行承担，与造价通网站无关。"
)
add_body(
    "（4）第三方供应商的联系方式可能因变更、失效等原因无法联系，造价通网站不因此承担责任，"
    "用户可改用人工询价渠道。"
)

add_heading("X.6 导出文件与盖章说明", 3)
add_body(
    "（1）用户通过本服务导出的报价文件由系统根据 AI 生成内容自动排版生成，"
    "页面与文件均标注“本结果由造价通 AI 生成，仅供参考”等标识。"
)
add_body(
    "（2）导出文件中的盖章为系统依据平台规则生成，仅用于用户内部流转与参考，"
    "不代表造价通网站对外出具的价格证明、报价承诺或担保函，也不构成对任何第三方的意思表示。"
)
add_body("（3）用户应在导出文件约定位置完整保留 AI 生成标识、水印与免责提示，不得删除、遮盖或以其他方式使其无法辨识。")

add_heading("X.7 不作担保", 3)
add_body("在法律允许的最大范围内，造价通网站不对 AI 询价服务及 AI 生成内容作任何明示或默示的担保，包括但不限于：")
add_body("（1）准确性、完整性、及时性、可靠性担保；")
add_body("（2）适用于特定用途或满足用户特定期望的担保；")
add_body("（3）服务不中断、无错误、无延迟、无数据丢失的担保；")
add_body("（4）与任何特定供应商实际报价或成交价一致的担保。")

add_heading("X.8 责任限制与免责", 3)
add_body(
    "（1）用户知悉并同意，AI 生成内容不是价格结论，用户基于 AI 生成内容作出的任何判断、决策与行为"
    "（包括报价、投标、采购、签约、结算等），由用户自行负责。"
)
add_body(
    "（2）在法律允许的最大范围内，造价通网站及其关联公司、管理人员、员工与合作方，"
    "不对用户因使用或无法使用 AI 询价服务及 AI 生成内容而产生的任何直接或间接损失承担责任，"
    "包括但不限于报价差额损失、利润损失、投标失败、合同违约、工期延误、返工损失、商誉损失及第三方索赔。"
)
add_body(
    "（3）上述责任限制适用于合同、侵权、严格责任或其他任何法律理论下的主张，"
    "即使造价通网站已被事先告知该等损失的可能性。"
)
add_body(
    "（4）若依适用法律必须承担赔偿责任，造价通网站的累计赔偿责任以用户就该争议服务实际支付的费用总额为上限；"
    "使用免费或赠送次数产生的争议，赔偿上限为【待法务确认，建议人民币 0 元或象征性金额】。"
)
add_body("（5）本条不适用于依法不得免除或限制的责任，例如因造价通网站故意或重大过失造成的人身损害。")

add_heading("X.9 用户义务", 3)
add_body("（1）用户应自行核实材料名称、规格型号、数量、单位、地区、目标时间、项目阶段等输入信息的准确性与完整性。")
add_body("（2）用户应在使用 AI 生成内容前通过供应商正式报价、市场询价、人工询价等渠道交叉验证。")
add_body("（3）用户不得利用本服务从事任何违反法律法规、行业规范或损害第三方权益的行为，不得批量抓取、反向工程、恶意刷取使用次数。")
add_body("（4）用户因违反本条规定或不当使用 AI 生成内容而产生的一切后果由用户自行承担；给造价通网站或第三方造成损失的，用户应予以赔偿。")

add_heading("X.10 知识产权与使用许可", 3)
add_body("（1）AI 询价服务、模型、算法、界面设计及平台数据权益归造价通网站或其权利人所有。")
add_body(
    "（2）在用户遵守本协议的前提下，造价通网站授予用户一项非独占、不可转让、不可转授权、"
    "仅限内部业务使用的许可，用于查阅与使用 AI 生成内容。"
)
add_body("（3）未经书面许可，用户不得将 AI 生成内容用于对外商业分发、转售、训练其他模型或以其他方式向第三方提供。")

add_heading("X.11 条款更新", 3)
add_body(
    "造价通网站可根据法律法规、监管要求、模型与业务调整更新本章条款，并通过站内公告、弹窗等合理方式提示。"
    "更新后的条款自公布之日起生效，用户继续使用本服务即视为接受更新后的条款；"
    "用户不接受的，应停止使用本服务。"
)

add_heading("X.12 生效与接受", 3)
add_body("（1）用户首次使用 AI 查价或 AI 询价服务时，应阅读并勾选同意《造价通用户协议》及本章条款。")
add_body("（2）用户勾选同意或实际使用本服务，即视为已充分阅读、理解并接受本章全部内容，特别是显著提示的免责与责任限制条款。")

# ---------------------------------------------------------------- 四、触点表
add_heading("四、产品内展示触点与文案", 1, page_break=True)
add_table(
    ["触点", "位置", "建议文案", "交互与优先级"],
    [
        [
            "首次使用弹窗",
            "AI 查价 / AI 询价首次进入",
            "标题“AI 询价服务说明与免责声明”。正文：1）本服务由造价通 AI 基于 building DS 模型自动生成结果，"
            "可能出现不准确或不完整；2）结果仅供参考，不构成报价承诺、合同依据或专业意见，请以供应商正式报价为准；"
            "3）请勿将结果作为投标、结算、审计、司法鉴定的唯一依据；4）因使用结果数据产生的决策与后果，"
            "由用户自行承担，造价通网站不对结果负责。",
            "勾选框“我已阅读并同意《用户协议》与《AI 询价服务免责声明》”＋“同意并继续”/“暂不使用”。P0",
        ],
        [
            "结果页结论区",
            "询价建议结果标题下方",
            "本结果由造价通 AI（building DS 模型）自动生成，仅供参考，不构成报价承诺或缔约邀请；实际价格以供应商正式报价为准。",
            "常驻小字，全部结果通用。P0",
        ],
        ["结果页展开项", "与上条同行", "查看 AI 生成说明", "展开显示第三部分条款精简版。P1"],
        [
            "条件不足提示",
            "结果页质量提示",
            "现有“由于搜索条件不足，建议价区间或存在偏差”后追加“结果仅供参考，请补充规格、数量、地区等信息后重新查询。”",
            "复用现有提示组件。P1",
        ],
        [
            "供应商联系方式弹窗",
            "展示第三方联系方式时",
            "以下联系方式来自第三方或历史询价记录，造价通未对其真实性、有效性进行核实，请自行判断并承担交易风险。",
            "弹窗内说明条。P1",
        ],
        [
            "导出文件页脚",
            "报价文件末尾",
            "现有“该文件由 AI 生成，仅供报价参考。”扩展为两行：本文件由造价通 AI（building DS 模型）自动生成，"
            "内容仅供参考，不构成报价承诺、价格证明或合同依据。用户使用本文件数据产生的后果由用户自行承担，"
            "造价通网站不对结果负责。",
            "随文件导出，不得删除。P0",
        ],
        ["盖章区说明", "盖章附近", "本盖章由系统按平台规则自动生成，仅限内部参考使用。", "小字标注。P1"],
        [
            "AI 查价结果",
            "查价结果卡片底部",
            "以上价格由造价通 AI 基于公开与历史数据生成，仅供参考，不作为定价与结算依据。",
            "常驻小字。P0",
        ],
        ["输入框下方", "对话输入区", "AI 结果仅供参考，请核实后使用", "常驻，不干扰输入。P2"],
        ["帮助中心", "AI 询价 FAQ", "《AI 询价结果是怎么来的？可以用于投标吗？》", "沉淀完整条款与常见问题。P2"],
    ],
    widths=[2.2, 3.0, 7.4, 4.4],
)

# ---------------------------------------------------------------- 五、留痕
add_heading("五、合规留痕要求", 1)
add_body("1. 同意记录：留存用户 ID、协议与免责声明版本号、同意时间、渠道；协议版本升级后需再次确认或站内提示。")
add_body("2. 结果标识：每条 AI 生成结果保留可追溯的生成标识（生成时间、模型版本、询价参数），便于事后核查。")
add_body("3. 导出留痕：导出文件中固定保留 AI 生成标识、水印与免责提示。")
add_body("4. 免责声明不可绕过：前端不得提供隐藏、关闭后永久失效等使免责提示无法再次触达的交互；用户关闭后再次使用应可重新查看。")

# ---------------------------------------------------------------- 六、待确认
add_heading("六、待确认事项", 1)
add_body(
    "1. 模型名称对外口径：协议与页面是否统一披露为“building DS 模型”，还是改用统一对外名称"
    "（如“造价通造价大模型”），建议以品牌口径为准并保留模型变更说明。"
)
add_body("2. 责任上限金额：X.8（4）中免费与赠送次数场景的赔偿上限取值，需法务确认。")
add_body(
    "3. 盖章效力口径：X.6（2）关于系统盖章效力的表述，需与造价通正式对外文件规则及业务口径对齐；"
    "若盖章方式为待领导批示状态，条款需同步保留弹性表述。"
)
add_body("4. 条款挂载位置：本章作为《用户协议》独立章节，还是单列《AI 询价服务特别条款》，需法务确认衔接方式。")
add_body("5. 争议解决与法律适用：沿用《用户协议》现有约定，如有特别安排需单独确认。")
add_body("6. 弹窗频次：首次使用弹窗是否按账户一次性展示，或按协议版本升级后再次展示，需产品确认。")

# ------------------------------------------------------- 七、AI 生成文件文案
add_heading("七、AI 生成文件免责文案", 1, page_break=True)

add_heading("7.1 推荐定稿（导出文件末尾，完整版）", 2)
add_body(
    "本文件由造价通 AI 基于 building DS 模型自动生成，内容仅供参考，不构成报价承诺、价格证明、"
    "合同依据或任何专业意见，亦不得作为投标报价、工程结算、审计核对或司法鉴定的依据。"
    "请以供应商正式报价及双方签署的书面文件为准。用户使用本文件数据所作出的任何判断与决策，"
    "以及由此产生的后果，由用户自行承担，造价通网站不对本文件内容及其使用结果承担任何责任。"
    "本文件的 AI 生成标识、水印与免责提示不得删除、遮盖或以其他方式使其无法辨识。"
)

add_heading("7.2 短版（页脚或结果区常驻，一行）", 2)
add_body(
    "本文件由造价通 AI（building DS 模型）自动生成，仅供参考，不作为报价或结算依据，"
    "造价通不对使用结果负责。"
)

add_heading("7.3 超短版（角标或水印旁）", 2)
add_body("AI 生成内容 · 仅供参考")

add_heading("7.4 下载前提示（弹窗或 Toast）", 2)
add_body(
    "该文件由造价通 AI 自动生成，仅供内部参考，不构成对外报价或价格证明。"
    "下载并使用即视为您已知悉并接受上述说明。"
)

add_heading("7.5 文件页眉或首页提示条", 2)
add_body("【AI 生成文件】本报价文件由造价通 AI 自动生成，仅供参考，不作为合同附件与结算依据。")

add_heading("7.6 使用说明", 2)
add_body("1. 7.1 用于导出文件末尾固定段落，作为唯一的正式免责条款文本，需与《用户协议》第 X.6 条保持一致表述。")
add_body(
    "2. 7.2 至 7.5 为同一口径的简化表达，可按位置裁剪，"
    "但不得改变“AI 生成、仅供参考、责任由用户自负、造价通不负责”四项核心含义。"
)
add_body("3. 文件中的免责文案与 AI 生成标识属于文件组成部分，用户端导出、打印、转发时应完整保留。")

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(str(OUTPUT))
print(f"saved: {OUTPUT}")
