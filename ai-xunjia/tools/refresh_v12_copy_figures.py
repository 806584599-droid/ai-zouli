"""Refresh the V1.2 figures after removing the 通宝提示语 copy entry (2026-10-16).

Only replaces existing image blobs; no paragraph is inserted or removed, so the
script is idempotent for the figure set already present in the document.
"""

from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


ROOT = Path("/Users/zouxiaoli/Documents/ChatGPT/ai询价")
DOC_PATH = ROOT / "造价通AI询价需求规格说明书（V1.2）.docx"
BACKUP_PATH = ROOT / ".qa/造价通AI询价需求规格说明书（V1.2）_复制入口截图更新前备份.docx"
SHOTS = ROOT / ".qa/ai_xunjia_doc/pc"
FLOW = ROOT / "ai询价业务流程图.png"

# 图 3-7 对应两张截图：条件不足提示在前，候选不足结果在后。
FIGURES = {
    "图 2.4-1": [FLOW],
    "图 3-1": [SHOTS / "01_home.png"],
    "图 3-2": [SHOTS / "02_workspace_empty.png"],
    "图 3-3": [SHOTS / "12_invalid_input.png"],
    "图 3-4": [SHOTS / "11_incomplete_spec.png"],
    "图 3-5": [SHOTS / "04_thinking_steps.png"],
    "图 3-6": [SHOTS / "05_result_cement.png"],
    "图 3-7": [SHOTS / "21_insufficient_conditions.png", SHOTS / "09_result_insufficient.png"],
    "图 3-8": [SHOTS / "18_phase_question.png"],
    "图 3-9": [SHOTS / "06_result_feedback.png"],
    "图 3-10": [SHOTS / "07_supplier_contact.png"],
    "图 3-11": [SHOTS / "08_supplier_unavailable.png"],
    "图 3-12": [SHOTS / "10_no_data.png"],
    "图 3-13": [SHOTS / "16_manual_inquiry_modal.png"],
    "图 3-14": [SHOTS / "13_preview_seal.png"],
    "图 3-15": [SHOTS / "14_history.png"],
    "图 3-16": [SHOTS / "15_delete_confirm.png"],
}


def paragraph_index(doc: Document, prefix: str) -> int:
    for index, paragraph in enumerate(doc.paragraphs):
        if paragraph.text.strip().startswith(prefix):
            return index
    raise RuntimeError(f"caption not found: {prefix}")


def picture_paragraphs_before(doc: Document, caption_prefix: str) -> list:
    index = paragraph_index(doc, caption_prefix)
    found = []
    offset = 1
    while index - offset >= 0:
        paragraph = doc.paragraphs[index - offset]
        if paragraph._p.findall(".//" + qn("a:blip")):
            found.append(paragraph)
        elif found:
            break
        offset += 1
        if offset > 8:
            break
    return list(reversed(found))


def replace_blob(doc: Document, paragraph, image_path: Path) -> None:
    blip = paragraph._p.findall(".//" + qn("a:blip"))[0]
    rel_id = blip.get(qn("r:embed"))
    doc.part.related_parts[rel_id]._blob = image_path.read_bytes()


def main() -> None:
    BACKUP_PATH.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(DOC_PATH, BACKUP_PATH)
    doc = Document(DOC_PATH)

    for caption_prefix, images in FIGURES.items():
        paragraphs = picture_paragraphs_before(doc, caption_prefix)
        if len(paragraphs) != len(images):
            raise RuntimeError(
                f"{caption_prefix}: found {len(paragraphs)} pictures, expected {len(images)}"
            )
        for paragraph, image_path in zip(paragraphs, images):
            replace_blob(doc, paragraph, image_path)
        print(f"{caption_prefix} <- {', '.join(image.name for image in images)}")

    doc.save(DOC_PATH)
    print(f"Saved: {DOC_PATH}")
    print(f"Backup: {BACKUP_PATH}")


if __name__ == "__main__":
    main()
