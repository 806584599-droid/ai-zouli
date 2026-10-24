"""Refresh figures and captions in the V1.2 requirement document after the second review round."""

from __future__ import annotations

import shutil
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm
from docx.text.paragraph import Paragraph


ROOT = Path("/Users/zouxiaoli/Documents/ChatGPT/ai询价")
DOC_PATH = ROOT / "造价通AI询价需求规格说明书（V1.2）.docx"
BACKUP_PATH = ROOT / ".qa/造价通AI询价需求规格说明书（V1.2）_插图更新前备份.docx"
SHOTS = ROOT / ".qa/ai_xunjia_doc/pc"
FLOW = ROOT / "ai询价业务流程图.png"

FIGURE_FILES = {
    "图 2.4-1": FLOW,
    "图 3-1": SHOTS / "01_home.png",
    "图 3-2": SHOTS / "02_workspace_empty.png",
    "图 3-3": SHOTS / "12_invalid_input.png",
    "图 3-4": SHOTS / "11_incomplete_spec.png",
    "图 3-5": SHOTS / "04_thinking_steps.png",
    "图 3-6": SHOTS / "05_result_cement.png",
    "图 3-7": SHOTS / "09_result_insufficient.png",
    "图 3-8": SHOTS / "18_phase_question.png",
    "图 3-9": SHOTS / "06_result_feedback.png",
    "图 3-10": SHOTS / "07_supplier_contact.png",
    "图 3-11": SHOTS / "08_supplier_unavailable.png",
    "图 3-12": SHOTS / "10_no_data.png",
    "图 3-13": SHOTS / "16_manual_inquiry_modal.png",
    "图 3-14": SHOTS / "13_preview_seal.png",
    "图 3-15": SHOTS / "14_history.png",
    "图 3-16": SHOTS / "15_delete_confirm.png",
}

CAPTION_UPDATES = {
    "图 3-7": "图 3-7 搜索条件不足提示、候选不足 3 条时全部展示截图",
    "图 3-8": "图 3-8 项目阶段提示位于询价建议结果之后（多轮问答确认）截图",
}

BULLET_UPDATES = (
    (
        "询价建议结果与候选不足：图 3-6、图 3-7；",
        "询价建议结果、条件不足提示与候选不足：图 3-6、图 3-7；",
    ),
    (
        "项目阶段多轮问答确认：图 3-8；",
        "项目阶段提示（位于询价建议结果之后）：图 3-8；",
    ),
)


def set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def paragraph_index(doc: Document, prefix: str) -> int:
    for index, paragraph in enumerate(doc.paragraphs):
        if paragraph.text.strip().startswith(prefix):
            return index
    raise RuntimeError(f"Paragraph not found: {prefix}")


def replace_figure_blob(doc: Document, caption_prefix: str, image_path: Path) -> None:
    index = paragraph_index(doc, caption_prefix)
    for offset in range(1, 5):
        if index - offset < 0:
            break
        paragraph = doc.paragraphs[index - offset]
        blips = paragraph._p.findall(".//" + qn("a:blip"))
        if not blips:
            continue
        rel_id = blips[0].get(qn("r:embed"))
        part = doc.part.related_parts[rel_id]
        part._blob = image_path.read_bytes()
        return
    raise RuntimeError(f"Figure image not found for {caption_prefix}")


def insert_extra_figure(doc: Document, caption_prefix: str, image_path: Path) -> None:
    index = paragraph_index(doc, caption_prefix)
    picture = None
    for offset in range(1, 5):
        if index - offset < 0:
            break
        paragraph = doc.paragraphs[index - offset]
        if paragraph._p.findall(".//" + qn("a:blip")):
            picture = paragraph
            break
    if picture is None:
        raise RuntimeError(f"Figure image not found for {caption_prefix}")

    new_xml = deepcopy(picture._p)
    picture._p.addprevious(new_xml)
    new_paragraph = Paragraph(new_xml, picture._parent)
    for run in list(new_paragraph.runs):
        run._element.getparent().remove(run._element)
    new_paragraph.add_run().add_picture(str(image_path), width=Cm(16.0))


def main() -> None:
    BACKUP_PATH.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(DOC_PATH, BACKUP_PATH)
    doc = Document(DOC_PATH)

    for caption_prefix, image_path in FIGURE_FILES.items():
        replace_figure_blob(doc, caption_prefix, image_path)
        print(f"replaced {caption_prefix} -> {image_path.name}")

    insert_extra_figure(doc, "图 3-7", SHOTS / "21_insufficient_conditions.png")
    print("inserted 21_insufficient_conditions.png before 图 3-7 caption")

    for caption_prefix, text in CAPTION_UPDATES.items():
        paragraph = doc.paragraphs[paragraph_index(doc, caption_prefix)]
        set_paragraph_text(paragraph, text)
        print(f"caption updated: {text}")

    for old, new in BULLET_UPDATES:
        paragraph = doc.paragraphs[paragraph_index(doc, old)]
        set_paragraph_text(paragraph, new)
        print(f"bullet updated: {new}")

    doc.save(DOC_PATH)
    print(f"Saved: {DOC_PATH}")
    print(f"Backup: {BACKUP_PATH}")


if __name__ == "__main__":
    main()
