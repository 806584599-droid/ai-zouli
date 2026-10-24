"""Append the latest V1.2 changes to the document change log."""

from __future__ import annotations

import shutil
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph


ROOT = Path("/Users/zouxiaoli/Documents/ChatGPT/ai询价")
DOC_PATH = ROOT / "造价通AI询价需求规格说明书（V1.2）.docx"
BACKUP_PATH = ROOT / ".qa/造价通AI询价需求规格说明书（V1.2）_履历更新前备份.docx"

CHANGE_TEXT = (
    "补充 AI 询价评测与回归 harness 要求：新增规则、模型、检索排序、编排接口、前端交互与日志审计分层评测，"
    "明确固定数据与模型输出回放、版本追溯、测试集隔离、执行节奏和发布门禁，并新增验收项 A30–A34；"
    "统一项目阶段采集范围：新建会话首次进入正常结果生成流程时，在通宝思考结束后、结果生成前采集一次，"
    "同一会话修改后再次提交或重新查询时沿用已选阶段，无数据、规格不足或非材料询价场景不采集；"
    "明确建议价金额按档位格式化：低于 1000 元保留两位小数，1000 元（含）至 10000 元（含）保留整数，"
    "高于 10000 元按十位取整。"
)


def set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def append_cell_paragraph(cell, text: str) -> None:
    template = cell.paragraphs[-1]
    new_p = deepcopy(template._p)
    cell._tc.append(new_p)
    paragraph = Paragraph(new_p, cell)
    set_paragraph_text(paragraph, text)


def main() -> None:
    BACKUP_PATH.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(DOC_PATH, BACKUP_PATH)
    doc = Document(DOC_PATH)

    changed = False
    for table in doc.tables[:2]:
        for row in table.rows:
            cells = row.cells
            if len(cells) < 4:
                continue
            if cells[1].text.strip() != "V1.2":
                continue
            description = cells[2]
            if "评测与回归 harness" in description.text:
                changed = True
                break
            append_cell_paragraph(description, CHANGE_TEXT)
            changed = True
            break
        if changed:
            break

    if not changed:
        raise RuntimeError("V1.2 change-log row not found")

    doc.save(DOC_PATH)
    print(f"Saved: {DOC_PATH}")
    print(f"Backup: {BACKUP_PATH}")


if __name__ == "__main__":
    main()
