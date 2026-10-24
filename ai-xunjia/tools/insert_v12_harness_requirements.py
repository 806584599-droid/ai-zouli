"""Insert AI inquiry evaluation-harness requirements into the V1.2 document."""

from __future__ import annotations

import shutil
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph


ROOT = Path("/Users/zouxiaoli/Documents/ChatGPT/ai询价")
DOC_PATH = ROOT / "造价通AI询价需求规格说明书（V1.2）.docx"
BACKUP_PATH = ROOT / ".qa/造价通AI询价需求规格说明书（V1.2）_harness写入前备份.docx"


def set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def set_cell_text(cell, text: str) -> None:
    paragraph = cell.paragraphs[0]
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)
    for extra in list(cell.paragraphs[1:]):
        extra._element.getparent().remove(extra._element)


def find_table_by_headers(doc: Document, headers: list[str]) -> Table:
    for table in doc.tables:
        if not table.rows:
            continue
        actual = [cell.text.strip() for cell in table.rows[0].cells]
        if actual[: len(headers)] == headers:
            return table
    raise RuntimeError(f"Table not found by headers: {headers}")


def insert_row_after(table: Table, row_index: int, values: list[str]) -> None:
    source_tr = table.rows[row_index]._tr
    template = deepcopy(source_tr)
    source_tr.addnext(template)
    row = table.rows[row_index + 1]
    for cell, value in zip(row.cells, values):
        set_cell_text(cell, value)


def insert_after(paragraph: Paragraph, text: str, style: str | None = None) -> Paragraph:
    new_p = deepcopy(paragraph._p)
    paragraph._p.addnext(new_p)
    inserted = Paragraph(new_p, paragraph._parent)
    set_paragraph_text(inserted, text)
    if style:
        inserted.style = style
    return inserted


def build_harness_table(doc: Document, anchor: Paragraph) -> None:
    template = find_table_by_headers(doc, ["编号", "验收项", "验收标准"])
    new_tbl = deepcopy(template._tbl)
    header_tr = deepcopy(template.rows[0]._tr)
    body_template = deepcopy(template.rows[1]._tr)
    for tr in list(new_tbl.tr_lst):
        new_tbl.remove(tr)
    new_tbl.append(header_tr)
    rows = [
        ["评测层级", "主要覆盖范围", "验收要求"],
        [
            "规则层",
            "三级排序、五级系数匹配、cid=32 默认系数、历史价倒推、项目阶段系数及建议价金额格式。",
            "固定数据下业务规则结果必须确定；P0 用例通过率 100%。",
        ],
        [
            "模型层",
            "材料名称、规格、地区、目标时间、数量、单位与项目阶段拆解，材料标准化及失败兜底。",
            "固定或录制模型输出可回放；核心字段识别不低于已确认基线，正式发布基线不低于 95%。",
        ],
        [
            "检索与排序",
            "ES 候选召回、地区偏好、DeepSeek 相关性筛选、Top 3 最终排序与候选不足场景。",
            "固定数据集下 Top 3 询价 ID 及顺序核对准确；隐藏回归集不得出现新增退化。",
        ],
        [
            "编排与接口",
            "意图拆解 → 材料标准化 → 系数映射 → ES 检索 → 候选筛选 → 建议价计算 → 前端返回。",
            "无数据、规格不足、非材料询价等分支均可回放；requestId 可串联中间结果。",
        ],
        [
            "前端交互",
            "思考前置、首次阶段采集、同会话阶段复用、结果操作、PDF 导出、转人工与异常提示。",
            "P0 Playwright 流程全部通过；重新查询不得重复采集项目阶段。",
        ],
        [
            "日志与审计",
            "查询、结果、重新查询、扣次、点赞、点踩、导出、转人工、价格来源提问与会话删除。",
            "事件名称、关键属性和关联标识完整率 100%，可按 conversationId、requestId 还原操作。",
        ],
    ]

    for _ in rows[1:]:
        new_tbl.append(deepcopy(body_template))

    table = Table(new_tbl, doc._body)
    for row, values in zip(table.rows, rows):
        for cell, value in zip(row.cells, values):
            set_cell_text(cell, value)
    anchor._p.addprevious(new_tbl)


def main() -> None:
    BACKUP_PATH.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(DOC_PATH, BACKUP_PATH)
    doc = Document(DOC_PATH)

    heading = next(
        paragraph
        for paragraph in doc.paragraphs
        if paragraph.text.strip() == "5 其他需求" and paragraph.style.name == "Heading 1"
    )
    heading.insert_paragraph_before("4.6 AI 询价评测与回归要求", style="Heading 2")
    heading.insert_paragraph_before(
        "系统需建立可重复执行的 AI 询价评测与回归机制，作为模型、Prompt、检索策略、系数表、规则配置与前后端联调变更的发布门槛。"
        "该机制用于在模型或数据存在非确定性时，仍能验证业务规则、候选结果、页面交互与行为日志未发生回归。"
        "具体测试框架、回放文件格式与 CI 实现由技术设计与测试方案确定，本需求约束评测能力、覆盖范围、版本追溯与验收指标。",
        style="Normal",
    )
    build_harness_table(doc, heading)
    for text in [
        "可重复评测：固定历史询价数据、系数映射、市场价、项目阶段、系统时间与模型输出；同一版本在相同 fixture 下重复执行，结构化结果应一致。",
        "版本追溯：每次评测报告记录模型、Prompt、系数表、历史价数据与策略版本；失败报告需定位到意图、检索、排序、计价、UI 或埋点层级。",
        "测试集隔离：研发可见用例与隐藏回归集分开维护，不得针对固定 Golden Case 硬编码业务结果或绕过正常处理链路。",
        "执行节奏：提交阶段执行快速确定性用例与关键页面流程；每日执行完整回放；模型、Prompt、检索策略或系数表变更后必须执行完整回归。",
        "发布门槛：P0 业务规则、Top 3 排序、金额格式、项目阶段复用、异常兜底和埋点断言必须全部通过；任一 P0 回归修复前不得发布。",
    ]:
        heading.insert_paragraph_before(text, style="List Bullet")

    toc_risk = next(paragraph for paragraph in doc.paragraphs if paragraph.text.strip() == "4.5 风险边界")
    insert_after(toc_risk, "4.6 AI 询价评测与回归要求")

    terms = find_table_by_headers(doc, ["术语", "说明"])
    if not any(row.cells[0].text.strip() == "评测 harness" for row in terms.rows):
        insert_row_after(
            terms,
            len(terms.rows) - 1,
            [
                "评测 harness",
                "通过固定输入、历史询价、系数表、市场价、时间基准与模型回放结果，对 AI 询价各处理阶段重复执行并输出结构化结果的评测机制。",
            ],
        )

    functions = find_table_by_headers(doc, ["模块", "功能点", "需求描述", "优先级"])
    if not any(row.cells[1].text.strip() == "AI 询价评测与回归 harness" for row in functions.rows):
        insert_row_after(
            functions,
            len(functions.rows) - 1,
            [
                "质量保障",
                "AI 询价评测与回归 harness",
                "建立分层、可重复的评测与回归机制，覆盖规则、模型、检索排序、编排接口、前端交互及日志审计；支持固定数据与模型输出回放、版本追溯和发布门禁。具体测试框架与 CI 实现由技术设计与测试方案确定。",
                "P0",
            ],
        )

    acceptance = find_table_by_headers(doc, ["编号", "验收项", "验收标准"])
    existing = {row.cells[0].text.strip() for row in acceptance.rows}
    for values in [
        [
            "A30",
            "评测可重复性",
            "固定输入、历史询价、系数表、时间基准与模型输出连续执行 20 次，结构化结果完全一致，不出现候选顺序、阶段状态或金额格式波动。",
        ],
        [
            "A31",
            "评测分层覆盖",
            "规则、模型、检索排序、编排接口、前端交互与日志审计六层均配置 P0 用例；P0 用例覆盖率与通过率均为 100%。",
        ],
        [
            "A32",
            "模型回归",
            "模型或 Prompt 变更后，封印评测集上的核心字段识别准确率不得低于已确认基线，正式发布基线不低于 95%；材料名称、规格与目标时间字段准确率不低于 98%。",
        ],
        [
            "A33",
            "评测版本追溯",
            "每次评测报告记录模型、Prompt、系数表、历史价数据与策略版本；可通过 caseId、requestId 还原各处理阶段的结构化中间结果与失败原因。",
        ],
        [
            "A34",
            "发布门禁",
            "P0 业务规则、Top 3 排序、金额格式、项目阶段复用、异常兜底与埋点断言全部通过；性能按 4.4 验证；任一 P0 回归修复前不得发布。",
        ],
    ]:
        if values[0] not in existing:
            insert_row_after(acceptance, len(acceptance.rows) - 1, values)

    doc.save(DOC_PATH)
    print(f"Saved: {DOC_PATH}")
    print(f"Backup: {BACKUP_PATH}")


if __name__ == "__main__":
    main()
