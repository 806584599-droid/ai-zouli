"""Clarify the project-phase collection scope in the current V1.2 document."""

from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph


ROOT = Path("/Users/zouxiaoli/Documents/ChatGPT/ai询价")
DOC_PATH = ROOT / "造价通AI询价需求规格说明书（V1.2）.docx"
BACKUP_PATH = ROOT / ".qa/造价通AI询价需求规格说明书（V1.2）_阶段口径澄清前备份.docx"


def set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def set_cell_text(cell, text: str) -> None:
    paragraph = cell.paragraphs[0]
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)
    for extra in list(cell.paragraphs[1:]):
        extra._element.getparent().remove(extra._element)


def replace_by_prefix(doc: Document, prefix: str, new_text: str) -> None:
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            set_paragraph_text(paragraph, new_text)
            return
    raise RuntimeError(f"Paragraph not found by prefix: {prefix}")


def find_table_by_headers(doc: Document, headers: list[str]):
    for table in doc.tables:
        if not table.rows:
            continue
        actual = [cell.text.strip() for cell in table.rows[0].cells]
        if actual[: len(headers)] == headers:
            return table
    raise RuntimeError(f"Table not found by headers: {headers}")


def main() -> None:
    BACKUP_PATH.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(DOC_PATH, BACKUP_PATH)
    doc = Document(DOC_PATH)

    replace_by_prefix(
        doc,
        "每次新建询价会话首次提交材料询价条件后，通宝按",
        "每次新建询价会话首次提交且进入正常结果生成流程时，通宝按“接收需求 — 理解条件 — 匹配历史询价与市场价”优先展示思考过程；思考结束后、结果生成前，通过多轮对话采集一次本次询价所属的项目阶段，随后进入询价结果展示。同一会话内修改条件后再次提交或点击重新查询时，沿用已选项目阶段，不重复采集；仅新建会话、重置会话或用户明确指定其他阶段时重新确定。无数据、规格不足或非材料询价场景不采集项目阶段。思考态由通宝形象与步骤文案组成，通宝形象保持轻微上下浮动、左右停顿的动画，用于表达思考状态；动画节奏需平稳，不出现明显跳动或闪烁。",
    )
    replace_by_prefix(
        doc,
        "项目阶段通过多轮对话采集：每次新建询价会话首次提交材料询价条件后",
        "项目阶段通过多轮对话采集：每次新建询价会话首次提交且进入正常结果生成流程时，在通宝思考结束、结果生成前询问一次本次询价所属的项目阶段，选项为施工图预算、投资估算、设计概算与工程竣工结算；用户未填写或跳过时默认按施工图预算阶段处理。同一会话内修改条件后再次提交或点击重新查询时，沿用已选项目阶段，不重复弹出阶段采集；仅新建会话、重置会话或用户明确指定其他阶段时重新确定。无数据、规格不足或非材料询价场景不采集项目阶段。后端按“价格基础 × 时间系数 × 项目阶段系数”计算建议价，项目阶段系数仅用于后端计算，对话页与结果区不展示系数、换算公式与提示语。系数表如下：",
    )

    function_table = find_table_by_headers(doc, ["模块", "功能点", "需求描述", "优先级"])
    for row in function_table.rows:
        if row.cells[1].text.strip() == "项目阶段采集":
            set_cell_text(
                row.cells[2],
                "每次新建询价会话首次提交且进入正常结果生成流程时，在通宝思考结束、结果生成前采集一次项目阶段；同一会话内修改条件后再次提交或重新查询时沿用已选阶段，不重复采集。无数据、规格不足或非材料询价场景不采集。选项为施工图预算、投资估算、设计概算与工程竣工结算，用户未填写或跳过时默认施工图预算。",
            )
            break

    prompt_table = find_table_by_headers(doc, ["场景", "图标或控件", "文案", "说明"])
    for row in prompt_table.rows:
        if row.cells[0].text.strip() == "项目阶段确认":
            set_cell_text(
                row.cells[3],
                "每次新建会话首次提交且进入正常结果流程后采集一次；同一会话内修改后再次提交或重新查询沿用已选阶段，不重复弹出。无数据、规格不足或非材料询价不采集。未填写或跳过默认施工图预算，不展示阶段系数与换算公式。",
            )
            break

    acceptance_table = find_table_by_headers(doc, ["编号", "验收项", "验收标准"])
    for row in acceptance_table.rows:
        if row.cells[0].text.strip() == "A12":
            set_cell_text(
                row.cells[2],
                "每次新建会话首次提交且进入正常结果生成流程时采集一次项目阶段，未填写或跳过时默认施工图预算；同一会话内修改后再次提交或重新查询沿用已选阶段，不重复采集；无数据、规格不足或非材料询价不采集。后端按对应阶段系数计算建议价，对话页与结果区不展示系数、换算公式与提示语。",
            )
            break

    doc.save(DOC_PATH)
    print(f"Saved: {DOC_PATH}")
    print(f"Backup: {BACKUP_PATH}")


if __name__ == "__main__":
    main()
